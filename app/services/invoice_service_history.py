from __future__ import annotations

import csv
import hashlib
import io
import re
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.models.documents import Document, VehicleDocumentRecordTag
from app.models.invoice_service_history import (
    InvoiceServiceEvent,
    InvoiceServiceEventRevision,
    InvoiceServiceImportBatch,
)
from app.models.vehicles import Vehicle
from app.models.workshop import WorkshopProcess
from app.services.invoice_service_classifier import classify_invoice_service_text
from app.services.spreadsheets import normalize_header

EVENT_STATUSES = {"auto_extracted", "pending_validation", "validated", "rejected", "blocked"}
EVENT_SNAPSHOT_FIELDS = (
    "vehicle_id",
    "document_id",
    "workshop_process_id",
    "service_date",
    "odometer_km",
    "service_code",
    "service_label",
    "classification",
    "axle",
    "position",
    "supplier_name",
    "work_order_reference",
    "work_order_confidence",
    "amount",
    "currency",
    "status",
    "evidence_text",
    "evidence_json",
    "active",
    "validated_by_id",
    "validated_at",
    "last_batch_id",
)
STATUS_LABELS = {
    "auto_extracted": "Extraído automaticamente",
    "pending_validation": "Por validar",
    "validated": "Validado",
    "rejected": "Rejeitado/excluído",
    "blocked": "Bloqueado",
}
REQUIRED_COLUMNS = {"document_id", "stable_key"}
ALIASES = {
    "stable_key": {"stable_key", "event_key", "line_key", "chave_estavel", "chave_evento"},
    "document_id": {"document_id", "documento_id"},
    "service_text": {"service_text", "service", "servico", "descricao"},
    "service_code": {"service_code", "codigo_servico", "classificacao"},
    "service_date": {"service_date", "data", "data_servico", "data_fatura"},
    "odometer_km": {"odometer_km", "km", "quilometros"},
    "axle": {"axle", "eixo"},
    "position": {"position", "posicao", "localizacao"},
    "supplier_name": {"supplier_name", "fornecedor"},
    "work_order_reference": {"work_order_reference", "folha_obra", "fo"},
    "work_order_confidence": {"work_order_confidence", "confianca_fo"},
    "amount": {"amount", "valor"},
    "status": {"status", "estado"},
    "evidence": {"evidence", "evidencia"},
    "source_line_ids": {"source_line_ids", "linhas_origem", "source_lines"},
}


class InvoiceServiceImportError(ValueError):
    pass


def _canonical_header(value: Any) -> str:
    normalized = normalize_header(value)
    for canonical, aliases in ALIASES.items():
        if normalized in {normalize_header(alias) for alias in aliases}:
            return canonical
    return normalized


def _read_rows(content: bytes, filename: str) -> list[dict[str, Any]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        text = content.decode("utf-8-sig")
        sample = text[:4096]
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t") if sample.strip() else csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        return [{_canonical_header(k): v for k, v in row.items()} for row in reader]
    if suffix not in {".xlsx", ".xlsm"}:
        raise InvoiceServiceImportError("Formato não suportado; usa XLSX ou CSV.")
    with NamedTemporaryFile(suffix=suffix) as temporary:
        temporary.write(content)
        temporary.flush()
        workbook = load_workbook(temporary.name, data_only=True, read_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            values = sheet.iter_rows(values_only=True)
            headers = [_canonical_header(value) for value in next(values, ())]
            return [
                dict(zip(headers, row, strict=False))
                for row in values
                if any(v not in (None, "") for v in row)
            ]
        finally:
            workbook.close()


def _text(value: Any) -> str:
    return str(value or "").strip()


def _integer(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(float(str(value).replace(" ", "").replace(",", ".")))
    except ValueError:
        return None


def _decimal(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value).replace(" ", "").replace(",", "."))
    except InvalidOperation:
        return None


def _date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _event_values(row: dict[str, Any], document: Document) -> tuple[dict[str, Any], list[str]]:
    errors: list[str] = []
    service_text = _text(row.get("service_text"))
    proposal = classify_invoice_service_text(service_text)
    proposed_primary = next(
        (code for code in proposal.taxonomy_codes if code.startswith("MAINT.")),
        proposal.taxonomy_codes[0] if proposal.taxonomy_codes else "",
    )
    service_code = _text(row.get("service_code")) or proposed_primary
    status = _text(row.get("status")).lower() or (
        "blocked" if proposal.blocked_reasons else "auto_extracted"
    )
    status_aliases = {
        "extraido automaticamente": "auto_extracted",
        "extraído automaticamente": "auto_extracted",
        "por validar": "pending_validation",
        "validado": "validated",
        "rejeitado": "rejected",
        "excluido": "rejected",
        "excluído": "rejected",
        "bloqueado": "blocked",
    }
    status = status_aliases.get(status, status)
    if not service_code:
        errors.append("Serviço/classificação em falta ou não reconhecido.")
    if status not in EVENT_STATUSES:
        errors.append("Estado inválido.")
    source_line_ids = [
        item.strip()
        for item in re.split(r"[,;|]+", _text(row.get("source_line_ids")))
        if item.strip()
    ]
    evidence_json = proposal.as_dict()
    evidence_json["import_metadata"] = {"source_line_ids": source_line_ids}
    return {
        "vehicle_id": document.vehicle_id,
        "document_id": document.id,
        "service_date": _date(row.get("service_date")) or document.document_date,
        "odometer_km": _integer(row.get("odometer_km")),
        "service_code": service_code,
        "service_label": service_text or service_code,
        "classification": service_code.split(".", 1)[0] if service_code else None,
        "axle": _text(row.get("axle")) or None,
        "position": _text(row.get("position")) or None,
        "supplier_name": _text(row.get("supplier_name")) or document.supplier_name,
        "work_order_reference": _text(row.get("work_order_reference")) or None,
        "work_order_confidence": _decimal(row.get("work_order_confidence")),
        "amount": _decimal(row.get("amount")),
        "currency": "EUR",
        "status": status,
        "evidence_text": _text(row.get("evidence")) or service_text or None,
        "evidence_json": evidence_json,
        "active": status != "rejected",
    }, errors


def _document_projection(event: InvoiceServiceEvent) -> tuple[str, str | None, str | None]:
    code = event.service_code or ""
    axle = event.axle if event.axle in {"front", "rear", "both"} else "undefined"
    if code == "MAINT.PLAN":
        return "maintenance", "revision", None
    if code == "MAINT.OIL_INTERIM":
        return "maintenance", "degradation", None
    if code == "BRAKE.PAD":
        return "pads", axle, None
    if code == "BRAKE.DISC":
        return "discs", axle, None
    if code == "TYRE.REPAIR":
        return "tyres", "puncture", None
    if code.startswith("TYRE."):
        return "tyres", axle, None
    if code == "INSPECTION.PERIODIC":
        return "ipo", "yes", None
    return "other", None, event.service_label or code


def _sync_document_projection(db: Session, event: InvoiceServiceEvent) -> None:
    if not event.document_id or not event.vehicle_id:
        raise InvoiceServiceImportError("Evento sem vínculo documental e de viatura.")
    document = db.get(Document, event.document_id)
    if not document or document.vehicle_id != event.vehicle_id:
        raise InvoiceServiceImportError("Vínculo documental do evento é inválido.")
    source_kind = f"invoice_service:{event.id}"
    db.execute(
        delete(VehicleDocumentRecordTag).where(
            VehicleDocumentRecordTag.document_id == event.document_id,
            VehicleDocumentRecordTag.source_kind == source_kind,
        )
    )
    if not event.active or event.status in {"rejected", "blocked"}:
        return
    category, value, free_text = _document_projection(event)
    db.add(
        VehicleDocumentRecordTag(
            vehicle_id=event.vehicle_id,
            document_id=event.document_id,
            category=category,
            value=value,
            free_text=free_text,
            source_kind=source_kind,
        )
    )


def _canonical_snapshot(value: Any) -> Any:
    if isinstance(value, dict):
        return {key: _canonical_snapshot(item) for key, item in sorted(value.items())}
    if isinstance(value, list):
        return [_canonical_snapshot(item) for item in value]
    if isinstance(value, Decimal):
        return str(value)
    return value


def _snapshot_matches(current: dict[str, Any], expected: dict[str, Any]) -> bool:
    return all(
        _canonical_snapshot(current.get(key)) == _canonical_snapshot(value)
        for key, value in expected.items()
    )


def preview_invoice_service_import(
    db: Session,
    content: bytes,
    filename: str,
    *,
    version: str,
    source: str,
    allowed_vehicle_id: int | None = None,
) -> dict[str, Any]:
    rows = _read_rows(content, filename)
    file_hash = hashlib.sha256(content).hexdigest()
    result: list[dict[str, Any]] = []
    counts = {
        key: 0 for key in ("create", "update", "keep", "block", "reject", "error", "conflict")
    }
    seen: set[str] = set()
    for number, row in enumerate(rows, start=2):
        errors: list[str] = []
        stable_key = _text(row.get("stable_key"))
        document_id = _integer(row.get("document_id"))
        if not stable_key:
            errors.append("stable_key em falta.")
        if stable_key in seen:
            errors.append("stable_key duplicada no ficheiro.")
        seen.add(stable_key)
        document = db.get(Document, document_id) if document_id else None
        if not document:
            errors.append("document_id inexistente.")
        elif not document.vehicle_id or not db.get(Vehicle, document.vehicle_id):
            errors.append("Documento sem viatura válida.")
        elif allowed_vehicle_id is not None and document.vehicle_id != allowed_vehicle_id:
            errors.append("Documento fora do âmbito da viatura selecionada.")
        values, value_errors = _event_values(row, document) if document else ({}, [])
        errors.extend(value_errors)
        if (
            document
            and source.startswith("invoice_service_remediation")
            and not values.get("evidence_json", {})
            .get("import_metadata", {})
            .get("source_line_ids")
        ):
            errors.append("Linhas de origem em falta para remediação documental.")
        existing = (
            db.scalar(
                select(InvoiceServiceEvent).where(
                    InvoiceServiceEvent.source == source,
                    InvoiceServiceEvent.stable_key == stable_key,
                )
            )
            if stable_key
            else None
        )
        changed = existing is not None and any(
            getattr(existing, key) != value for key, value in values.items()
        )
        action = (
            "error" if errors else "create" if not existing else "update" if changed else "keep"
        )
        if not errors and values.get("status") == "blocked":
            action = "block"
        elif not errors and values.get("status") == "rejected":
            action = "reject"
        if existing and existing.status == "validated" and changed:
            action = "conflict"
            errors.append("Registo validado: atualização exige decisão explícita.")
        counts[action] += 1
        result.append(
            {
                "row": number,
                "stable_key": stable_key,
                "action": action,
                "errors": errors,
                "values": values,
                "event_id": existing.id if existing else None,
            }
        )
    return {
        "schema": "carfast.invoice-service-import.v1",
        "version": version,
        "source": source,
        "filename": filename,
        "file_hash": file_hash,
        "rows": result,
        "counts": counts,
        "can_apply": counts["error"] == 0 and counts["conflict"] == 0,
    }


def _snapshot(event: InvoiceServiceEvent) -> dict[str, Any]:
    return {
        key: (
            str(value)
            if isinstance(value, Decimal)
            else value.isoformat()
            if isinstance(value, (date, datetime))
            else value
        )
        for key in EVENT_SNAPSHOT_FIELDS
        for value in (getattr(event, key),)
    }


def _restore_snapshot(event: InvoiceServiceEvent, snapshot: dict[str, Any]) -> None:
    """Restore every field known by the stored snapshot, including explicit nulls.

    Older revisions omitted null fields, so an absent key remains intentionally unknown.
    New revisions contain the complete mutable state and are fully symmetric.
    """

    for key in EVENT_SNAPSHOT_FIELDS:
        if key not in snapshot:
            continue
        value = snapshot[key]
        if key == "service_date":
            value = _date(value)
        elif key in {"amount", "work_order_confidence"}:
            value = _decimal(value)
        elif key == "validated_at" and isinstance(value, str):
            value = datetime.fromisoformat(value)
        setattr(event, key, value)


def _json_safe(value: Any) -> Any:
    if isinstance(value, (date, datetime, Decimal)):
        return value.isoformat() if not isinstance(value, Decimal) else str(value)
    if isinstance(value, dict):
        return {key: _json_safe(item) for key, item in value.items()}
    if isinstance(value, list):
        return [_json_safe(item) for item in value]
    return value


def apply_invoice_service_import(
    db: Session,
    preview: dict[str, Any],
    *,
    actor_id: int | None,
    allow_validated_updates: bool = False,
) -> InvoiceServiceImportBatch:
    if preview.get("schema") != "carfast.invoice-service-import.v1":
        raise InvoiceServiceImportError("Pré-visualização inválida.")
    if not preview.get("can_apply") and not allow_validated_updates:
        raise InvoiceServiceImportError("O dry-run contém erros ou conflitos.")
    previous_batch = db.scalar(
        select(InvoiceServiceImportBatch).where(
            InvoiceServiceImportBatch.file_hash == preview["file_hash"],
            InvoiceServiceImportBatch.version == preview["version"],
        )
    )
    if previous_batch and all(row["action"] == "keep" for row in preview["rows"]):
        return previous_batch
    if previous_batch:
        raise InvoiceServiceImportError("Esta versão e hash de lote já foram aplicados.")
    batch = InvoiceServiceImportBatch(
        version=preview["version"],
        file_hash=preview["file_hash"],
        filename=preview["filename"],
        source=preview["source"],
        status="applied",
        author_id=actor_id,
        applied_at=datetime.now(UTC),
        summary_json=preview["counts"],
        differences_json=_json_safe(preview["rows"]),
    )
    db.add(batch)
    db.flush()
    for row in preview["rows"]:
        if row["action"] in {"error", "keep"}:
            continue
        if row["action"] == "conflict" and not allow_validated_updates:
            raise InvoiceServiceImportError("Atualização de validado sem decisão explícita.")
        event = db.scalar(
            select(InvoiceServiceEvent).where(
                InvoiceServiceEvent.source == preview["source"],
                InvoiceServiceEvent.stable_key == row["stable_key"],
            )
        )
        before = _snapshot(event) if event else None
        if not event:
            event = InvoiceServiceEvent(
                source=preview["source"], stable_key=row["stable_key"], **row["values"]
            )
            db.add(event)
            db.flush()
            action = "create"
        else:
            for key, value in row["values"].items():
                setattr(event, key, value)
            action = "update"
        event.last_batch_id = batch.id
        db.flush()
        _sync_document_projection(db, event)
        db.add(
            InvoiceServiceEventRevision(
                event_id=event.id,
                batch_id=batch.id,
                action=action,
                before_json=before,
                after_json=_snapshot(event),
                actor_id=actor_id,
                reason="validated_override" if row["action"] == "conflict" else None,
            )
        )
    db.flush()
    return batch


def rollback_invoice_service_batch(
    db: Session,
    batch_id: int,
    *,
    actor_id: int | None,
    expected_vehicle_id: int | None = None,
) -> InvoiceServiceImportBatch:
    batch = db.get(InvoiceServiceImportBatch, batch_id)
    if not batch or batch.status not in {"applied", "rolled_back"}:
        raise InvoiceServiceImportError("Lote inexistente ou não aplicável.")
    revisions = list(
        db.scalars(
            select(InvoiceServiceEventRevision)
            .where(
                InvoiceServiceEventRevision.batch_id == batch.id,
                InvoiceServiceEventRevision.action.in_({"create", "update"}),
            )
            .order_by(InvoiceServiceEventRevision.id.desc())
        )
    )
    if not revisions:
        raise InvoiceServiceImportError("Lote sem revisões de aplicação; rollback recusado.")
    events = [db.get(InvoiceServiceEvent, revision.event_id) for revision in revisions]
    if any(event is None for event in events):
        raise InvoiceServiceImportError("Evento do lote inexistente; rollback recusado.")
    if expected_vehicle_id is not None and any(
        event.vehicle_id != expected_vehicle_id for event in events if event is not None
    ):
        raise InvoiceServiceImportError("Lote fora do âmbito da viatura selecionada.")
    if batch.status == "rolled_back":
        rollback_revisions = list(
            db.scalars(
                select(InvoiceServiceEventRevision).where(
                    InvoiceServiceEventRevision.batch_id == batch.id,
                    InvoiceServiceEventRevision.action == "rollback",
                )
            )
        )
        rollback_by_event = {revision.event_id: revision for revision in rollback_revisions}
        if any(
            event is None
            or event.id not in rollback_by_event
            or not _snapshot_matches(
                _snapshot(event), rollback_by_event[event.id].after_json or {}
            )
            for event in events
        ):
            raise InvoiceServiceImportError(
                "Estado do rollback inconsistente; operação idempotente recusada."
            )
        return batch

    if len({revision.event_id for revision in revisions}) != len(revisions):
        raise InvoiceServiceImportError(
            "Lote com revisões de aplicação duplicadas; rollback recusado."
        )

    # Fail closed before mutating any event. Callers should not have to rely on a
    # transaction rollback to undo a partially processed batch.
    for revision, event in zip(revisions, events, strict=True):
        assert event is not None
        current = _snapshot(event)
        expected = revision.after_json or {}
        if event.last_batch_id != batch.id or not _snapshot_matches(current, expected):
            raise InvoiceServiceImportError(
                "Evento alterado após o lote; rollback recusado para preservar auditoria."
            )

    for revision, event in zip(revisions, events, strict=True):
        assert event is not None
        current = _snapshot(event)
        if revision.before_json is None:
            event.active = False
            event.status = "rejected"
        else:
            _restore_snapshot(event, revision.before_json)
        _sync_document_projection(db, event)
        db.add(
            InvoiceServiceEventRevision(
                event_id=event.id,
                batch_id=batch.id,
                action="rollback",
                before_json=current,
                after_json=_snapshot(event),
                actor_id=actor_id,
                reason="logical_batch_rollback",
            )
        )
    batch.status = "rolled_back"
    batch.rolled_back_at = datetime.now(UTC)
    db.flush()
    return batch


def vehicle_service_history(
    db: Session, vehicle_id: int, *, service_code: str = "", include_pending: bool = True
) -> dict[str, Any]:
    statement = select(InvoiceServiceEvent).where(
        InvoiceServiceEvent.vehicle_id == vehicle_id, InvoiceServiceEvent.active.is_(True)
    )
    if service_code:
        statement = statement.where(InvoiceServiceEvent.service_code.startswith(service_code))
    if not include_pending:
        statement = statement.where(InvoiceServiceEvent.status == "validated")
    events = list(
        db.scalars(
            statement.order_by(InvoiceServiceEvent.service_date.asc(), InvoiceServiceEvent.id.asc())
        )
    )
    confirmed = [event for event in events if event.status == "validated"]
    return {
        "events": events,
        "confirmed": confirmed,
        "pending": [event for event in events if event.status != "validated"],
        "first_service_date": confirmed[0].service_date if confirmed else None,
        "first_service_km": confirmed[0].odometer_km if confirmed else None,
    }


def link_event_to_work_order(
    db: Session,
    event: InvoiceServiceEvent,
    process: WorkshopProcess,
    *,
    confidence: Decimal | None,
    actor_id: int | None,
) -> None:
    if process.vehicle_id != event.vehicle_id:
        raise InvoiceServiceImportError("A folha de obra pertence a outra viatura.")
    before = _snapshot(event)
    event.workshop_process_id = process.id
    event.work_order_reference = str(process.id)
    event.work_order_confidence = confidence
    db.flush()
    db.add(
        InvoiceServiceEventRevision(
            event_id=event.id,
            action="link_work_order",
            before_json=before,
            after_json=_snapshot(event),
            actor_id=actor_id,
            reason="human_validation",
        )
    )


def decide_invoice_service_event(
    db: Session,
    event: InvoiceServiceEvent,
    *,
    status: str,
    service_code: str,
    axle: str | None,
    position: str | None,
    actor_id: int | None,
    reason: str,
) -> None:
    if status not in EVENT_STATUSES:
        raise InvoiceServiceImportError("Estado inválido.")
    if status in {"validated", "rejected", "blocked"} and not reason.strip():
        raise InvoiceServiceImportError("A decisão exige justificação.")
    before = _snapshot(event)
    event.status = status
    event.service_code = service_code.strip() or event.service_code
    event.classification = event.service_code.split(".", 1)[0]
    event.axle = axle.strip() if axle else None
    event.position = position.strip() if position else None
    event.active = status != "rejected"
    if status == "validated":
        event.validated_by_id = actor_id
        event.validated_at = datetime.now(UTC)
    db.flush()
    _sync_document_projection(db, event)
    db.add(
        InvoiceServiceEventRevision(
            event_id=event.id,
            action="decision",
            before_json=before,
            after_json=_snapshot(event),
            actor_id=actor_id,
            reason=reason.strip() or None,
        )
    )


def validate_document_invoice_service_events(
    db: Session,
    *,
    document_id: int,
    vehicle_id: int,
    decisions: list[dict[str, Any]],
    actor_id: int | None,
) -> None:
    """Validate or reject every active imported proposal shown on an invoice.

    The exact-set check prevents a stale browser form from silently omitting a
    proposal added since the invoice was opened. All validation happens before
    the first mutation so the command fails closed.
    """

    document = db.get(Document, document_id)
    if not document or document.vehicle_id != vehicle_id:
        raise InvoiceServiceImportError("Fatura ou associação de viatura inválida.")
    events = list(
        db.scalars(
            select(InvoiceServiceEvent)
            .where(
                InvoiceServiceEvent.document_id == document_id,
                InvoiceServiceEvent.vehicle_id == vehicle_id,
                InvoiceServiceEvent.active.is_(True),
            )
            .order_by(InvoiceServiceEvent.id)
        )
    )
    if not events:
        return
    if len({item.get("event_id") for item in decisions}) != len(decisions):
        raise InvoiceServiceImportError("Decisões de serviços duplicadas ou inválidas.")
    by_id = {item.get("event_id"): item for item in decisions}
    if set(by_id) != {event.id for event in events}:
        raise InvoiceServiceImportError("A lista de serviços mudou; reabre a fatura.")

    for event in events:
        item = by_id[event.id]
        decision = str(item.get("decision") or "").strip()
        reason = str(item.get("reason") or "").strip()
        if decision not in {"accept", "reject"}:
            raise InvoiceServiceImportError("Todos os serviços exigem uma decisão.")
        if decision == "reject" and not reason:
            raise InvoiceServiceImportError("A exclusão de um serviço exige motivo.")
        if decision == "accept" and event.status == "blocked":
            raise InvoiceServiceImportError("Um serviço bloqueado não pode ser publicado.")

    for event in events:
        item = by_id[event.id]
        decision = str(item["decision"])
        if decision == "accept" and event.status == "validated":
            continue
        decide_invoice_service_event(
            db,
            event,
            status="validated" if decision == "accept" else "rejected",
            service_code=event.service_code,
            axle=event.axle,
            position=event.position,
            actor_id=actor_id,
            reason=(
                "validated_from_invoice"
                if decision == "accept"
                else str(item.get("reason") or "").strip()
            ),
        )
