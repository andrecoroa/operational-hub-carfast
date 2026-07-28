import hashlib
import json
import zipfile
from datetime import date, datetime
from pathlib import Path
from io import BytesIO

import fitz
from openpyxl import Workbook
from sqlalchemy import select

import app.web.router as web_router
from app.core.config import settings
from app.models.documents import (
    DiagnosticDocument,
    Document,
    DocumentEvent,
    DocumentLink,
    VehicleDocumentAuditField,
    VehicleDocumentRecord,
    VehicleDocumentRecordTag,
)
from app.models.vehicles import Vehicle, VehicleIdentifier, VehicleManualField
from app.services.vehicle_document_history import (
    DOCUMENT_HISTORY_QUICK_CLASSIFICATIONS,
    document_center_module_context,
    preclassify_invoices_and_work_orders,
    vehicle_document_module_context,
)
from app.web.router import (
    _batch_document_type,
    _batch_document_vehicle,
    _batch_invoice_filinto_stacked_lines,
    _batch_invoice_payload,
    _batch_invoice_total_from_lines,
    _invoice_payload_was_extracted,
    _historical_report_date,
    _historical_report_match_metadata,
    _historical_report_payloads,
    _historical_report_vehicle,
    _document_latest_ocr_metadata,
    local_document_storage_folder,
)


def test_invoice_payload_accepts_structured_scanned_ocr_without_raw_preview():
    assert _invoice_payload_was_extracted(
        {
            "ocr_status": "extracted",
            "text_source": "pdf_image_ocr",
            "raw_text_preview": "",
            "document_number": "FS0002124",
            "invoice_lines": [{"description": "Serviço"}],
        }
    )
    assert not _invoice_payload_was_extracted(
        {
            "ocr_status": "not_extracted",
            "document_number": "FS0002124",
            "invoice_lines": [],
        }
    )


def _make_workbook(headers: list[str], rows: list[list[object]]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def _make_rentway_export_workbook(title: str, headers: list[str], rows: list[list[object]]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([title])
    sheet.append([""])
    sheet.append([f"{len(rows)} resultados"])
    sheet.append([None for _ in headers])
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def _make_zip(files: dict[str, bytes]) -> BytesIO:
    stream = BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    stream.seek(0)
    return stream


def _make_pdf(text: str) -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    content = document.tobytes()
    document.close()
    return content


def test_document_inbox_type_detection_is_conservative():
    assert _batch_document_type("Faturas/Fatura_123.pdf") == ("workshop_supplier_invoice", "workshop")
    assert _batch_document_type("Relatorios/Diagnostico_Autel.pdf") == ("workshop_report", "workshop")
    assert _batch_document_type("S_IM_VR3EDYHTXRJ968622_260512_1612.pdf") == ("workshop_report", "workshop")
    assert _batch_document_type("Folha de obra/FO_123.pdf") == ("workshop_work_order", "workshop")
    assert _batch_document_type("documento_recebido_001.pdf") == ("workshop_other", "workshop")


def test_legacy_stellantis_report_is_counted_as_vehicle_diagnostic(db_session):
    vehicle = _create_vehicle(db_session)
    document = Document(
        title="S_IM_VR3EDYHTXRJ968622_260512_1612",
        document_type="unknown_document",
        classification="triage",
        source="v2_clean_manual",
        entry_channel="document_inbox",
        original_name="S_IM_VR3EDYHTXRJ968622_260512_1612.pdf",
        file_name="S_IM_VR3EDYHTXRJ968622_260512_1612.pdf",
        storage_provider="local",
        storage_path="Frota/_POR_ASSOCIAR/00_Caixa_Entrada/S_IM.pdf",
        folder_path="Frota/_POR_ASSOCIAR/00_Caixa_Entrada",
        status="pending_triage",
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
        document_date=date(2026, 5, 12),
    )
    db_session.add(document)
    db_session.commit()

    module_ctx = vehicle_document_module_context(db_session, vehicle)

    assert module_ctx["group_counts"]["diagnostics"] == 1
    assert any(row["archive_group"] == "diagnostics" for row in module_ctx["archive_rows"])


def test_historical_report_payloads_accept_files_and_safe_zip_entries():
    archive = _make_zip(
        {
            "AQ-41-XJ/manutencao.pdf": b"report-a",
            "../escape.pdf": b"blocked",
            "notas.txt": b"ignored",
        }
    ).getvalue()

    payloads, error = _historical_report_payloads(
        [("lote.zip", archive), ("diagnostico.pdf", b"report-b")]
    )

    assert error is None
    assert [payload["original_name"] for payload in payloads] == [
        "manutencao.pdf",
        "diagnostico.pdf",
    ]


def test_historical_report_batch_defers_deep_ocr(
    authenticated_client,
    db_session,
    monkeypatch,
    tmp_path,
):
    monkeypatch.setattr(settings, "document_archive_root", str(tmp_path / "documents"))

    def fail_if_deep_ocr_runs(*_args, **_kwargs):
        raise AssertionError("batch import must defer deep OCR")

    monkeypatch.setattr(
        web_router,
        "extract_workshop_report_values_from_bytes",
        fail_if_deep_ocr_runs,
    )
    monkeypatch.setattr(
        web_router,
        "classify_workshop_report_from_bytes",
        fail_if_deep_ocr_runs,
    )
    monkeypatch.setattr(web_router, "extract_diagnostic_pdf", fail_if_deep_ocr_runs)

    archive = BytesIO()
    with zipfile.ZipFile(archive, "w") as bundle:
        for index in (1, 2):
            pdf = fitz.open()
            page = pdf.new_page()
            page.insert_text((72, 72), f"Relatorio de diagnostico {index}")
            bundle.writestr(f"relatorio_{index}.pdf", pdf.tobytes())
            pdf.close()
    archive.seek(0)

    response = authenticated_client.post(
        "/v2-clean/documents/import/historical-reports",
        files={"files": ("relatorios.zip", archive, "application/zip")},
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert "report_imported=2" in response.headers["location"]
    documents = db_session.scalars(
        select(Document).where(Document.entry_channel == "historical_report_import")
    ).all()
    assert len(documents) == 2
    assert {document.status for document in documents} == {"received"}
    profiles = db_session.scalars(select(DiagnosticDocument)).all()
    assert len(profiles) == 2
    assert {profile.ocr_status for profile in profiles} == {"pending"}


def test_historical_report_import_get_redirects_to_document_center(authenticated_client):
    response = authenticated_client.get(
        "/v2-clean/documents/import/historical-reports",
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert response.headers["location"] == "/v2-clean/documents#imports"


def test_historical_report_metadata_uses_batch_filename_without_opening_pdf():
    metadata = web_router._historical_report_metadata_from_filename(
        "VR3EDYHT9RJ968630/A_LD_VR3EDYHT9RJ968630_260728_1425.pdf"
    )

    assert metadata["vehicle_identifier"] == "VR3EDYHT9RJ968630"
    assert metadata["vin_candidates"] == ["VR3EDYHT9RJ968630"]
    assert metadata["report_code"] == "fault_reading"
    assert metadata["report_date"] == "28/07/2026"
    assert metadata["report_time"] == "14:25"


def test_historical_report_vehicle_prefers_extracted_vin():
    by_plate_vehicle = Vehicle(id=1, plate="AQ-41-XJ", vin="VIN-PLATE")
    by_vin_vehicle = Vehicle(id=2, plate="BB-69-TE", vin="VR3EDYHT9RJ968630")

    matched = _historical_report_vehicle(
        {
            "vin_candidates": ["VR3EDYHT9RJ968630"],
            "plate_candidates": ["AQ-41-XJ"],
        },
        "AQ-41-XJ.pdf",
        {"AQ41XJ": by_plate_vehicle},
        {"VR3EDYHT9RJ968630": by_vin_vehicle},
    )

    assert matched is by_vin_vehicle
    assert _historical_report_date("27/07/2026") == date(2026, 7, 27)


def test_historical_report_match_metadata_includes_extracted_vehicle_identity():
    metadata = _historical_report_match_metadata(
        {"report_code": "fault_reading"},
        {"registration": "AQ-41-XJ", "chassis": "VR3EDYHT9RJ968630"},
    )

    assert metadata["plate_candidates"] == ["AQ-41-XJ"]
    assert metadata["vin_candidates"] == ["VR3EDYHT9RJ968630"]
    assert metadata["vehicle_identifier"] == "VR3EDYHT9RJ968630"


def test_historical_report_extracted_values_are_visible_to_ocr_validation():
    event = DocumentEvent(
        action="historical_report.imported",
        new_value=json.dumps(
            {
                "report_code": "fault_reading",
                "extracted_values": {
                    "registration": "AQ-41-XJ",
                    "odometer_km": "42127",
                    "faults": ["P0011"],
                },
            }
        ),
    )

    metadata = _document_latest_ocr_metadata([event])

    assert metadata["report_code"] == "fault_reading"
    assert metadata["registration"] == "AQ-41-XJ"
    assert metadata["odometer_km"] == "42127"
    assert metadata["faults"] == ["P0011"]


def _create_vehicle(db_session):
    vehicle = Vehicle(
        plate="CC-11-AA",
        vin="VINCC11AA123456789",
        brand="PEUGEOT",
        model="2008",
        version="ALLURE",
        rentway_unit_nr="911",
        lifecycle_status="active",
        operational_status="free",
        active=True,
    )
    db_session.add(vehicle)
    db_session.commit()
    db_session.refresh(vehicle)
    return vehicle


def test_clean_document_storage_uses_configured_archive_root(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "document_archive_root", str(tmp_path))

    storage_folder = local_document_storage_folder(
        "Frota/BB-69-TE_VR7EFYHT2PJ697244/00_Importacoes_Estruturadas",
        plate="BB-69-TE",
        vin="VR7EFYHT2PJ697244",
    )

    assert storage_folder == tmp_path / "Frota" / "BB-69-TE_VR7EFYHT2PJ697244" / "00_Importacoes_Estruturadas"


def test_clean_document_batch_zip_associates_pending_and_deduplicates(
    authenticated_client,
    db_session,
    monkeypatch,
    tmp_path,
):
    monkeypatch.setattr(settings, "document_archive_root", str(tmp_path))
    vehicle = _create_vehicle(db_session)
    pdf = _make_pdf("Fatura FT-4458\nData 15/05/2026\nOleo motor 5W30 45,00\nFiltro de oleo 12,00")
    batch = _make_zip(
        {
            "CC-11-AA/Faturas/fatura_2026-05-15.pdf": pdf,
            "CC-11-AA/Faturas/copia_fatura.pdf": pdf,
            "Sem matricula/diagnostico.png": b"sample-image",
            "ignorar.txt": b"not a document",
        }
    )

    response = authenticated_client.post(
        "/v2-clean/documents/import/archive-batch",
        files={"file": ("documentos.zip", batch.getvalue(), "application/zip")},
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert "batch_imported=2" in response.headers["location"]
    assert "batch_matched=1" in response.headers["location"]
    assert "batch_pending=1" in response.headers["location"]
    assert "batch_duplicates=1" in response.headers["location"]
    assert "batch_reprocessed=1" in response.headers["location"]
    documents = db_session.scalars(
        select(Document).where(Document.entry_channel == "v2_clean_batch").order_by(Document.id)
    ).all()
    assert len(documents) == 2
    matched = next(item for item in documents if item.vehicle_id == vehicle.id)
    pending = next(item for item in documents if item.vehicle_id is None)
    assert matched.document_type == "workshop_supplier_invoice"
    assert matched.document_date == date(2026, 5, 15)
    assert matched.folder_path.endswith("01_Documentacao_Financeira/Faturas")
    assert Path(matched.storage_path).exists()
    ocr_event = db_session.scalar(
        select(DocumentEvent).where(
            DocumentEvent.document_id == matched.id,
            DocumentEvent.action == "invoice.ocr.extracted",
        )
    )
    assert ocr_event is not None
    payload = json.loads(ocr_event.new_value)
    assert payload["ocr_status"] == "extracted"
    assert payload["ocr_extractor_version"] == "invoice-ocr-2026-07-24-v4"
    assert any("Oleo motor" in row["description"] for row in payload["invoice_lines"])
    reprocessed_event = db_session.scalar(
        select(DocumentEvent).where(
            DocumentEvent.document_id == matched.id,
            DocumentEvent.action == "invoice.ocr.reprocessed",
        )
    )
    assert reprocessed_event is not None
    assert pending.folder_path == "Frota/_POR_ASSOCIAR/99_Pendentes_Classificar"
    assert Path(pending.storage_path).exists()

    matched.document_type = "workshop_other"
    matched.classification = "other"
    matched.status = "classified"
    db_session.commit()

    repeated = authenticated_client.post(
        "/v2-clean/documents/import/archive-batch",
        files={"file": ("documentos.zip", batch.getvalue(), "application/zip")},
        follow_redirects=False,
    )

    assert repeated.status_code == 303
    assert "batch_reprocessed=2" in repeated.headers["location"]
    db_session.refresh(matched)
    assert matched.document_type == "workshop_supplier_invoice"
    assert matched.status == "classified"


def test_clean_document_batch_zip_explicitly_reprocesses_every_existing_file(
    authenticated_client,
    db_session,
    monkeypatch,
    tmp_path,
):
    monkeypatch.setattr(settings, "document_archive_root", str(tmp_path))
    pdf = _make_pdf("Conteudo sem marcadores suficientes para deteção automática")
    batch = _make_zip({"arquivo/desconhecido.pdf": pdf})

    imported = authenticated_client.post(
        "/v2-clean/documents/import/archive-batch",
        files={"file": ("documentos.zip", batch.getvalue(), "application/zip")},
        follow_redirects=False,
    )
    assert imported.status_code == 303
    document = db_session.scalar(
        select(Document).where(Document.entry_channel == "v2_clean_batch")
    )
    assert document is not None
    document.document_type = "workshop_other"
    document.classification = "other"
    db_session.commit()

    repeated = authenticated_client.post(
        "/v2-clean/documents/import/archive-batch",
        files={"file": ("documentos.zip", batch.getvalue(), "application/zip")},
        data={"operation": "reprocess_invoices"},
        follow_redirects=False,
    )

    assert repeated.status_code == 303
    location = repeated.headers["location"]
    assert "batch_duplicates=1" in location
    assert "batch_reprocessed=1" in location
    assert "batch_reprocess_failed=0" in location
    assert "batch_operation=reprocess_invoices" in location
    db_session.refresh(document)
    assert document.document_type == "workshop_supplier_invoice"
    assert db_session.scalar(
        select(DocumentEvent).where(
            DocumentEvent.document_id == document.id,
            DocumentEvent.action == "invoice.ocr.reprocessed",
        )
    ) is not None


def test_invoice_ocr_manifest_updates_existing_document_by_hash(
    authenticated_client,
    db_session,
):
    digest = hashlib.sha256(b"same-invoice").hexdigest()
    document = Document(
        title="Fatura antiga",
        document_type="workshop_supplier_invoice",
        classification="invoice",
        source="v2_clean_manual",
        entry_channel="v2_clean_batch",
        original_name="fatura.pdf",
        file_name="fatura.pdf",
        file_type="pdf",
        storage_provider="local",
        storage_path="missing.pdf",
        file_hash=digest,
        status="classified",
        archived=True,
    )
    db_session.add(document)
    db_session.commit()
    db_session.refresh(document)
    manifest = {
        "schema": "carfast.invoice-ocr-manifest.v1",
        "documents": [
            {
                "file_hash": digest,
                "status": "extracted",
                "payload": {
                    "ocr_status": "extracted",
                    "text_source": "pdf_image_ocr",
                    "document_number": "FS0002124",
                    "supplier_name": "Cruz & Allen",
                    "document_date": "2025-05-27",
                    "invoice_lines": [{"description": "Serviço", "amount": "25,00"}],
                },
            }
        ],
    }

    response = authenticated_client.post(
        "/v2-clean/documents/import/invoice-ocr-manifest",
        files={
            "file": (
                "resultados.json",
                json.dumps(manifest).encode("utf-8"),
                "application/json",
            )
        },
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert "manifest_updated=1" in response.headers["location"]
    assert "manifest_lines=1" in response.headers["location"]
    db_session.refresh(document)
    assert document.contract_number == "FS0002124"
    assert document.supplier_name == "Cruz & Allen"
    assert document.document_date == date(2025, 5, 27)
    assert document.status == "classified"
    actions = db_session.scalars(
        select(DocumentEvent.action)
        .where(DocumentEvent.document_id == document.id)
        .order_by(DocumentEvent.id)
    ).all()
    assert actions == ["invoice.ocr.extracted", "invoice.ocr.manifest_imported"]


def test_clean_document_batch_vehicle_match_falls_back_to_vin(db_session):
    vehicle = _create_vehicle(db_session)
    vehicles_by_plate = {"OTHERPLATE": vehicle}
    vehicles_by_vin = {"VINCC11AA123456789": vehicle}

    matched = _batch_document_vehicle(
        "FACTURA\nChassis: VINCC11AA123456789\nData: 12/06/2026",
        vehicles_by_plate,
        vehicles_by_vin,
    )

    assert matched == vehicle


def test_clean_document_reprocess_invoice_ocr(authenticated_client, db_session, tmp_path):
    vehicle = _create_vehicle(db_session)
    invoice_path = tmp_path / "fatura_2026-05-15.pdf"
    invoice_path.write_bytes(
        _make_pdf("Fatura FT-4458\nData 15/05/2026\nOleo motor 5W30 45,00\nFiltro de oleo 12,00")
    )
    invoice = Document(
        title="Fatura oficina",
        document_type="workshop_supplier_invoice",
        classification="fleet",
        source="v2_clean_manual",
        entry_channel="v2_clean_batch",
        original_name=invoice_path.name,
        file_name=invoice_path.name,
        file_type="pdf",
        file_size=invoice_path.stat().st_size,
        storage_provider="local",
        storage_path=str(invoice_path),
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
        status="received",
        archived=True,
    )
    db_session.add(invoice)
    db_session.commit()
    db_session.refresh(invoice)

    response = authenticated_client.post(
        f"/v2-clean/documents/{invoice.id}/reprocess-ocr",
        data={"return_url": f"/v2-clean/fleet/{vehicle.id}/documents?main_group=invoices"},
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert "ocr_reprocessed=1" in response.headers["location"]
    db_session.refresh(invoice)
    assert invoice.status == "extracted"
    events = db_session.scalars(
        select(DocumentEvent)
        .where(DocumentEvent.document_id == invoice.id)
        .order_by(DocumentEvent.id)
    ).all()
    assert [event.action for event in events] == [
        "invoice.ocr.extracted",
        "invoice.ocr.reprocessed",
    ]
    payload = json.loads(events[0].new_value)
    assert payload["ocr_status"] == "extracted"
    assert payload["ocr_extractor_version"] == "invoice-ocr-2026-07-24-v4"
    assert payload["document_number"] == "4458"
    assert any("Oleo motor" in row["description"] for row in payload["invoice_lines"])


def test_clean_document_reprocess_invoice_ocr_batch_and_replace_old_metadata(
    authenticated_client,
    db_session,
    tmp_path,
):
    vehicle = _create_vehicle(db_session)
    batch_label = "ZIP faturas-filinto.zip [20260724-010203]"
    documents = []
    for index, number in enumerate(("4458", "4459"), start=1):
        invoice_path = tmp_path / f"fatura_{number}.pdf"
        invoice_path.write_bytes(
            _make_pdf(
                f"Fatura FT-{number}\nData 15/05/2026\n"
                f"Matrícula: {vehicle.plate}\nOleo motor 5W30 {40 + index},00"
            )
        )
        invoice = Document(
            title=f"Fatura {number}",
            document_type="workshop_supplier_invoice",
            classification="workshop",
            source="v2_clean_manual",
            entry_channel="v2_clean_batch",
            source_subject=f"{batch_label}: lote/{invoice_path.name}",
            original_name=invoice_path.name,
            file_name=invoice_path.name,
            file_type="pdf",
            file_size=invoice_path.stat().st_size,
            storage_provider="local",
            storage_path=str(invoice_path),
            vehicle_id=vehicle.id,
            plate=vehicle.plate,
            contract_number="DOCUMENTO",
            status="received",
            archived=True,
        )
        db_session.add(invoice)
        documents.append(invoice)
    db_session.commit()

    response = authenticated_client.post(
        "/v2-clean/documents/reprocess-ocr-batch",
        data={
            "batch_label": batch_label,
            "return_url": "/v2-clean/documents?main_group=invoices",
        },
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert "ocr_batch_processed=2" in response.headers["location"]
    assert "ocr_batch_failed=0" in response.headers["location"]
    for document, number in zip(documents, ("4458", "4459"), strict=True):
        db_session.refresh(document)
        assert document.contract_number == number
        assert document.status == "extracted"
        actions = db_session.scalars(
            select(DocumentEvent.action)
            .where(DocumentEvent.document_id == document.id)
            .order_by(DocumentEvent.id)
        ).all()
        assert actions == ["invoice.ocr.extracted", "invoice.ocr.reprocessed"]
    center = authenticated_client.get("/v2-clean/documents?main_group=invoices")
    assert center.status_code == 200
    assert 'action="/v2-clean/documents/reprocess-ocr-batch"' in center.text
    assert 'type="hidden" name="batch_label"' in center.text
    assert "Reprocessar" in center.text
    assert batch_label in center.text


def test_clean_document_detail_previews_invoice_and_identifies_batch(
    authenticated_client,
    db_session,
    tmp_path,
):
    invoice_path = tmp_path / "fatura-preview.pdf"
    pdf_content = _make_pdf("Fatura FT-8899\nData 15/05/2026\nServico 25,00")
    invoice_path.write_bytes(pdf_content)
    invoice = Document(
        title="Fatura com preview",
        document_type="workshop_supplier_invoice",
        classification="workshop",
        source="v2_clean_manual",
        entry_channel="v2_clean_batch",
        source_subject="ZIP maio.zip [20260724-020304]: faturas/fatura-preview.pdf",
        original_name=invoice_path.name,
        file_name=invoice_path.name,
        file_type="pdf",
        file_size=invoice_path.stat().st_size,
        storage_provider="local",
        storage_path=str(invoice_path),
        status="extracted",
        archived=True,
    )
    db_session.add(invoice)
    db_session.flush()
    db_session.add(
        DocumentEvent(
            document_id=invoice.id,
            action="invoice.ocr.extracted",
            new_value=json.dumps(
                {
                    "document_number": "8899",
                    "invoice_lines": [{"description": "Servico", "amount": "25,00"}],
                }
            ),
        )
    )
    db_session.commit()

    page = authenticated_client.get(f"/v2-clean/documents/{invoice.id}")

    assert page.status_code == 200
    assert "Pré-visualizar" in page.text
    assert f'/v2-clean/documents/{invoice.id}/file?inline=1' in page.text
    assert f'<iframe src="/v2-clean/documents/{invoice.id}/file?inline=1' not in page.text
    assert "ZIP maio.zip [20260724-020304]" in page.text
    assert "Reprocessar lote" in page.text

    file_response = authenticated_client.get(f"/v2-clean/documents/{invoice.id}/file?inline=1")
    assert file_response.status_code == 200
    assert file_response.content == pdf_content
    assert file_response.headers["content-type"].startswith("application/pdf")
    assert file_response.headers["content-disposition"].startswith("inline;")


def test_batch_invoice_payload_extracts_vertical_filinto_invoice():
    text = """
Sede: Filinto Mota Sucessores S.A. Rua Pinto Bessa, 550
Doc.Nº
Data
Conta
NIF
Vendedor
Marca :
Cor :
Modelo :
Descrição
Valor Total
16002610
26/07/2022
227010
509285970
João Sousa
Matrícula:
AS-92-ET
Chassis :
VR7BBYHZBNE038504
CIAL_FACVN 2 022 /
FACTURA
ORIGINAL
NIF: 500 115 966
R-Imposto Único de Circulação
147,21
E
0,00
147,21
0,00
Observações :
Total do documento
147,21
"""

    payload = _batch_invoice_payload(b"", ".pdf", "filinto.pdf", existing_text=text)

    assert payload["document_number"] == "16002610"
    assert payload["document_date"] == "2022-07-26"
    assert payload["supplier_name"] == "Filinto Mota Sucessores S.A. (NIF 500115966)"
    assert payload["plate"] == "AS-92-ET"
    assert payload["vin"] == "VR7BBYHZBNE038504"
    assert payload["invoice_lines"] == [
        {
            "reference": "",
            "description": "R-Imposto Único de Circulação",
            "quantity": "1",
            "unit": "",
            "unit_price": "147,21",
            "tax": "E",
            "amount": "147,21",
            "service": "Outro",
        }
    ]


def test_batch_invoice_payload_extracts_filinto_vnc_iuc_invoice():
    text = """
Filinto Mota Sucessores S.A.
Estr. Exterior da Circunvalação, 10686
4460-281 Srª da Hora - Matosinhos
NIF: 500 115 966
Exmos Senhores Carfast - Rent-A-Car, Lda
FACTURA Rua das Indústrias, 220
4785-625 TROFA
Doc.Nº CIAL_FACVN2 025 /3822
Conta 227010 Data 31/10/2025 NIF 509285970
Vendedor João Sousa
Marca : FIAT Modelo : 600 Hybrid Série 2 600 Hybrid 1.2 100cv DCT Matrícula: BZ-73-SC
Chassis : ZFA5FBAT9SJ078967 Nº Motor : Cilindrada : 1199 Nº Stock : 136 727
Cor : VERMELHO Interior : TECIDO COM FIO RECICLADO E MONOGRAM Combustível : Gasolina
PASSIONE
Descrição Valor Total
R-Imposto Único de Circulação 111,46
.
Observações : ALD n.º 2025.057598.00. IUC - Não Sujeito
Base de incidência de I.V.A. Síntese de Pagamento
Cod Taxa % Base IVA Valor IVA
Total Líquido 111,46
E 0,00 111,46 0,00 Total IVA 0,00
Total do documento 111,46
ATCUD: JJWNNZZ4-3822
"""

    payload = _batch_invoice_payload(b"", ".pdf", "VNC.22 (Original)-29204-28.pdf", existing_text=text)

    assert payload["ocr_template"] == "filinto_mota_venda_iuc"
    assert payload["document_number"] == "CIAL_FACVN2 025 /3822"
    assert payload["document_date"] == "2025-10-31"
    assert payload["supplier_name"] == "Filinto Mota Sucessores S.A. (NIF 500115966)"
    assert payload["supplier_nif"] == "500115966"
    assert payload["client_nif"] == "509285970"
    assert payload["plate"] == "BZ-73-SC"
    assert payload["vin"] == "ZFA5FBAT9SJ078967"
    assert payload["total_with_vat"] == "111,46"
    assert payload["vat_amount"] == "0,00"
    assert payload["atcud"] == "JJWNNZZ4-3822"
    assert payload["model"] == "600 Hybrid Série 2 600 Hybrid 1.2 100cv DCT"
    assert payload["invoice_lines"] == [
        {
            "reference": "",
            "description": "R-Imposto Único de Circulação",
            "quantity": "1",
            "unit": "",
            "unit_price": "111,46",
            "tax": "E",
            "amount": "111,46",
            "service": "Outro",
        }
    ]


def test_batch_invoice_payload_extracts_filinto_vnc_finance_invoice():
    text = """
Filinto Mota Sucessores S.A.
NIF: 500 115 966
Exmos Senhores Carfast - Rent-A-Car, Lda
FACTURA Rua das Indústrias, 220
Doc.Nº CIAL_FACVN2 025 /1195
Conta 227010 Data 01/04/2025 NIF 509285970
Marca : PEUGEOT Modelo : Partner Longa 1.5 BlueHDi Matrícula: BS-61-FU
Chassis : VR3EDYHTXRJ968622 Combustível : Diesel
Descrição Valor Total
Isento IVA ao Abrigo Art.º 16 - N.º 6, Alinea C
R-Despesas de Locação Financeira 68,00
R-Serviços Prestados Loc. Financeira 37,00
Observações : Viatura Vendida com financiamento
Total Líquido 105,00
Total IVA 8,51
Total do documento 113,51
ATCUD: JJWNNZZ4-1195
"""

    payload = _batch_invoice_payload(b"", ".pdf", "filinto_1195.pdf", existing_text=text)

    assert payload["ocr_template"] == "filinto_mota_venda_financeira"
    assert payload["document_number"] == "CIAL_FACVN2 025 /1195"
    assert payload["supplier_nif"] == "500115966"
    assert payload["client_nif"] == "509285970"
    assert payload["plate"] == "BS-61-FU"
    assert payload["vin"] == "VR3EDYHTXRJ968622"
    assert payload["total_with_vat"] == "113,51"
    assert [(line["description"], line["amount"]) for line in payload["invoice_lines"]] == [
        ("R-Despesas de Locação Financeira", "68,00"),
        ("R-Serviços Prestados Loc. Financeira", "37,00"),
    ]


def test_batch_invoice_payload_extracts_filinto_automoveis_invoice_without_vertical_noise():
    text = """
788
581
105
€ 000.573
laicoS
.paC
Filinto Mota - Automóveis, Lda
Estr. Exterior da Circunvalação, 10686
4460-281 Srª da Hora - Matosinhos
NIF: 501 185 887
Exmos Senhores Carfast - Rent-A-Car, Lda
FACTURA Rua das Indústrias, 220
4785-625 TROFA
Doc.Nº CIAL_FACVO2 025 /585
Conta 227010 Data 29/05/2025 NIF509285970
Vendedor First Rent
Marca : CITROEN Modelo : Jumper FF 33 L3H2 2.2 BlueHDi Matrícula: AS-01-ZG
Chassis : VF7YBBPFCNG007451 Nº Motor : Cilindrada : 2179 Nº Stock : 701 002
Descrição Valor Total
Despesas Recondicionamento 1 953,00
.
Observações :
Base de incidência de I.V.A. Síntese de Pagamento
Cod Taxa % Base IVA Valor IVA
Total Líquido 1 953,00
V 23,00 1 953,00 449,19 Total IVA 449,19
Total do documento 2 402,19
ATCUD: JJWJNKPD-585
"""

    payload = _batch_invoice_payload(b"", ".pdf", "Fatura-585-0.pdf", existing_text=text)

    assert payload["ocr_template"] == "filinto_mota_automoveis_cial_facvo2"
    assert payload["supplier_name"] == "Filinto Mota - Automóveis, Lda (NIF 501185887)"
    assert payload["supplier_nif"] == "501185887"
    assert payload["client_nif"] == "509285970"
    assert payload["document_number"] == "CIAL_FACVO2 025/585"
    assert payload["document_date"] == "2025-05-29"
    assert payload["plate"] == "AS-01-ZG"
    assert payload["vin"] == "VF7YBBPFCNG007451"
    assert payload["km"] == ""
    assert payload["work_order_reference"] == ""
    assert payload["subtotal_without_vat"] == "1953,00"
    assert payload["vat_amount"] == "449,19"
    assert payload["total_with_vat"] == "2402,19"
    assert payload["atcud"] == "JJWJNKPD-585"
    assert payload["invoice_lines"] == [
        {
            "reference": "",
            "description": "Despesas Recondicionamento",
            "quantity": "",
            "unit": "",
            "unit_price": "",
            "tax": "",
            "amount": "1953,00",
            "service": "Por classificar",
        }
    ]


def test_batch_invoice_payload_does_not_map_carfast_nif_as_filinto_supplier():
    text = """
CLIENTE : 227010 NIF :509285970
Exmos Senhores Carfast - Rent-A-Car, Lda
Filinto Mota Sucessores S.A.
NIF: 500 115 966
DOCUMENTO : TAL_FAC 2025/11170001
Descrição Referência P.V.Unit Desc P.Liq.Unit Tmp/Qt Total Liq. CT
FILTRO OLEO 1109AL 20,65 20,00 16,52 1 16,52B
CÓDIGO/DESCRIÇÃO I.V.A. TAXA I.V.A. BASE INCIDÊNCIA VALOR I.V.A.
"""

    payload = _batch_invoice_payload(b"", ".pdf", "filinto.pdf", existing_text=text)

    assert payload["supplier_name"] == "Filinto Mota Sucessores S.A. (NIF 500115966)"
    assert payload["supplier_nif"] == "500115966"
    assert payload["client_number"] == "227010"
    assert payload["client_nif"] == "509285970"


def test_batch_invoice_payload_does_not_invent_supplier_from_carfast_client_block():
    text = """
Fornecedor sem NIF visível
CLIENTE : 227010 NIF :509285970
Fatura FT-7788
Data 15/05/2026
Serviço 25,00
"""

    payload = _batch_invoice_payload(b"", ".pdf", "fatura.pdf", existing_text=text)

    assert payload["supplier_nif"] == ""
    assert payload["supplier_name"] == ""
    assert payload["client_nif"] == "509285970"


def test_batch_invoice_payload_extracts_eugenio_sage_invoice_with_work_order():
    text = """
Fatura FAC 020/17027
Data Emissão 26/03/2024
NIF: PT510464157
V/VIATURA AS-65-ZG 63163,00
EQUILIBRAGEM DE JANTES
1
5,6911 EUR
5,69 EUR
REPARAÇAO DE FURO COMERCIAL
1
8,1301 EUR
8,13 EUR
REQ. 1136
Observações:
11,23 EUR
2,58 EUR
13,81 EUR
Data/Hora Carga
© Sage licenciado a: EUGENIO & JORGE PEREIRA LDA /510464157
"""

    payload = _batch_invoice_payload(b"", ".pdf", "Fatura_020_17027.pdf", existing_text=text)

    assert payload["document_number"] == "FAC 020/17027"
    assert payload["document_date"] == "2024-03-26"
    assert payload["supplier_name"] == "Eugenio & Jorge Pereira Lda"
    assert payload["supplier_nif"] == "510464157"
    assert payload["plate"] == "AS-65-ZG"
    assert payload["km"] == "63163"
    assert payload["total_with_vat"] == "13,81"
    assert payload["work_order_reference"] == "1136"
    assert payload["ocr_template"] == "eugenio_jorge_sage_fac"
    assert payload["invoice_lines"] == [
        {
            "reference": "",
            "description": "EQUILIBRAGEM DE JANTES",
            "quantity": "1",
            "unit": "",
            "unit_price": "5,6911",
            "tax": "",
            "amount": "5,69",
            "service": "Pneus",
            "service_detail": "",
        },
        {
            "reference": "",
            "description": "REPARAÇAO DE FURO COMERCIAL",
            "quantity": "1",
            "unit": "",
            "unit_price": "8,1301",
            "tax": "",
            "amount": "8,13",
            "service": "Pneus",
            "service_detail": "Furo",
        },
    ]


def test_batch_invoice_payload_uses_work_order_when_eugenio_vehicle_block_is_missing():
    text = """
Factura FAC MT25/89
Data Emissão 23/01/2025
NIF: PT510464157
V/VIATURA 0 TRANSFERÊNCIA PARA OFICINA
PNEU 215/65R16C BESTDRIVE
2
70,0000 EUR
140,00 EUR
EQUILIBRAGEM DE JANTES
2
5,6911 EUR
11,38 EUR
REQ. Nº 2300
Observações:
© Sage licenciado a: EUGENIO & JORGE PEREIRA LDA /510464157
"""

    payload = _batch_invoice_payload(b"", ".pdf", "Fatura_MT25_89.pdf", existing_text=text)

    assert payload["document_number"] == "FAC MT25/89"
    assert payload["document_date"] == "2025-01-23"
    assert payload["supplier_nif"] == "510464157"
    assert payload["plate"] == ""
    assert payload["work_order_reference"] == "2300"
    assert [line["service"] for line in payload["invoice_lines"]] == ["Pneus", "Pneus"]


def test_batch_invoice_payload_extracts_only_eugenio_invoice_items():
    text = """
Fatura
FAC 020/15834
2024-01-11
Data Vencimento
V/VIATURA
2024-01-11
Data de emissão
EUGENIO & JORGE PEREIRA LDA Rua Principal NIF: PT510464157
NIF nº: 510464157
Marca
Matricula
Kms
PEUGEOT PARTNER
AO-52-ZX
63985,00
Descrição
Qt.
P.Venda S/Iva
Desc.
Ecovalor
Valor Liquido
Iva
ATCUD: JF98898F-15834
195/65R15 HANKOOK K435 91T
A
C
69db 2
59,1125 EUR
2,1000 EUR
118,23 EUR
23%
EQUILIBRAGEM ESPECIAL OU COM ARO LIGEIRO
2
4,8780 EUR
0,0000 EUR
9,76 EUR
23%
LAMPADA H7
1
6,5000 EUR
0,0000 EUR
6,50 EUR
23%
Observações:
136,59 EUR
31,42 EUR
168,01 EUR
© Sage licenciado a: EUGENIO & JORGE PEREIRA LDA /510464157
"""

    payload = _batch_invoice_payload(b"", ".pdf", "Fatura_020_15834.pdf", existing_text=text)

    assert payload["document_number"] == "FAC 020/15834"
    assert payload["total_with_vat"] == "168,01"
    assert [line["description"] for line in payload["invoice_lines"]] == [
        "195/65R15 HANKOOK K435 91T",
        "EQUILIBRAGEM ESPECIAL OU COM ARO LIGEIRO",
        "LAMPADA H7",
    ]
    assert payload["invoice_lines"][0]["quantity"] == "2"
    assert payload["invoice_lines"][0]["unit_price"] == "59,1125"
    assert payload["invoice_lines"][0]["amount"] == "118,23"
    assert payload["invoice_lines"][1]["quantity"] == "2"
    assert payload["invoice_lines"][2]["quantity"] == "1"


def test_batch_invoice_payload_extracts_filinto_tal_invoice():
    text = """
ORIGINAL
Filinto Mota Sucessores S.A.
NIF: 500 115 966
DOCUMENTO : TAL_FAC 2023/11156583
DE : 22/12/2023 O.R. : 321803 OFICINA : (081) MMarca Reparação VENCIMENTO : 22/12/2023
MAT : BB-32-FT MARCA : PEUGEOT MODELO : 2008 1.2PureTech Allure
KMS : 23506 RECEPCIONISTA : Gil Teixeira CHASSIS : VR3USHNSSPJ666227
Descrição Referência P.V.Unit Desc P.Liq.Unit Tmp/Qt Total Liq. CT
REVISÃO A
- OPERAÇÕES SISTEMÁTICAS DE MANUTENÇÃO - 95N48A 49,00 25,00 36,75 1 36,75B
OLEO TOTAL INEO XTRA FIRST 0W20 LT QINEOXF 40,00 39,00 24,40 3,50 85,40B
Folha de Obra nº 378
CÓDIGO/DESCRIÇÃO I.V.A. TAXA I.V.A. BASE INCIDÊNCIA VALOR I.V.A. TOTAL LÍQUIDO TOTAL I.V.A. TOTAL A PAGAR
B APV TX NORMAL 23,00 175,18 40,29 175,18 40,29 215,47
"""

    payload = _batch_invoice_payload(b"", ".pdf", "tal.pdf", existing_text=text)

    assert payload["document_kind"] == "invoice"
    assert payload["document_number"] == "TAL_FAC 2023/11156583"
    assert payload["document_date"] == "2023-12-22"
    assert payload["supplier_name"] == "Filinto Mota Sucessores S.A. (NIF 500115966)"
    assert payload["supplier_nif"] == "500115966"
    assert payload["plate"] == "BB-32-FT"
    assert payload["vin"] == "VR3USHNSSPJ666227"
    assert payload["km"] == "23506"
    assert payload["total_with_vat"] == "215,47"
    assert payload["work_order_reference"] == "378"
    assert payload["repair_order_reference"] == "321803"
    assert payload["ocr_template"] == "filinto_mota_tal"
    assert payload["invoice_lines"][0]["service"] == "Manutenção"
    assert payload["invoice_lines"][0]["amount"] == "36,75"


def test_batch_invoice_payload_extracts_filinto_tal_11156038_complete_table():
    text = """
ORIGINAL
Filinto Mota Sucessores S.A.
NIF: 500 115 966
FACTURA
PRONTO PAGAMENTO
Modo de Pagamento : Pronto Pagamento
CLIENTE : 227010 NIF :509285970
ATCUD: JFS4WFN7-11156038
DOCUMENTO : TAL_FAC 2023/11156038
DE : 05/12/2023 O.R. : 320340 OFICINA : (001) CIT Reparação VENCIMENTO : 05/12/2023
MAT : AN-48-RA MARCA : CITROEN MODELO : C3 1.5 BlueHDi 100 S&S CVM Shine Pack
KMS : 39643 RECEPCIONISTA : Gil Teixeira CHASSIS : VF7SXYHTUNT506748
Descrição Referência P.V.Unit Desc P.Liq.Unit Tmp/Qt Total Liq. CT
VER TRAVÕES E
SUBSTITUICAO DISCO TRAVAO AF (2) 25680910 57,00 25,00 42,75 1,10 47,03B
E:2 DISCOS DE TRAVÃO DIA PSA 1686717180 110,95 20,00 88,76 1 88,76B
CONJ 4 PASTILHAS TRAVÕES FR 1647863780 98,38 20,00 78,70 1 78,70B
SPRAY LIMPEZA TRAVOES 500ML D1636264180 4,17 20,00 3,34 1 3,34B
Subtotal Peças (170,80)
Nova Intervenção 217,83
OFERTA DE LAVAGEM G
Lavagem Automática com Aspiração LAVACA 12,50 99,99 1 B
Nova Intervenção
CÓDIGO/DESCRIÇÃO I.V.A. TAXA I.V.A. BASE INCIDÊNCIA VALOR I.V.A. TOTAL LÍQUIDO TOTAL I.V.A. TOTAL A PAGAR
B APV TX NORMAL 23,00 217,83 50,10 217,83 50,10 267,93
"""

    payload = _batch_invoice_payload(b"", ".pdf", "1-TAL-2-11156038-227010-1300984.pdf", existing_text=text)

    assert payload["document_number"] == "TAL_FAC 2023/11156038"
    assert payload["document_date"] == "2023-12-05"
    assert payload["supplier_nif"] == "500115966"
    assert payload["client_nif"] == "509285970"
    assert payload["plate"] == "AN-48-RA"
    assert payload["vin"] == "VF7SXYHTUNT506748"
    assert payload["km"] == "39643"
    assert payload["repair_order_reference"] == "320340"
    assert payload["total_with_vat"] == "267,93"
    assert payload["invoice_lines"] == [
        {
            "reference": "25680910",
            "description": "SUBSTITUICAO DISCO TRAVAO AF (2)",
            "quantity": "1,10",
            "unit": "",
            "unit_price": "42,75",
            "list_price": "57,00",
            "discount_percent": "25,00",
            "tax": "",
            "amount": "47,03",
            "service": "Discos",
        },
        {
            "reference": "1686717180",
            "description": "E:2 DISCOS DE TRAVÃO DIA PSA",
            "quantity": "1",
            "unit": "",
            "unit_price": "88,76",
            "list_price": "110,95",
            "discount_percent": "20,00",
            "tax": "",
            "amount": "88,76",
            "service": "Discos",
        },
        {
            "reference": "1647863780",
            "description": "CONJ 4 PASTILHAS TRAVÕES FR",
            "quantity": "1",
            "unit": "",
            "unit_price": "78,70",
            "list_price": "98,38",
            "discount_percent": "20,00",
            "tax": "",
            "amount": "78,70",
            "service": "Calços",
        },
        {
            "reference": "D1636264180",
            "description": "SPRAY LIMPEZA TRAVOES 500ML",
            "quantity": "1",
            "unit": "",
            "unit_price": "3,34",
            "list_price": "4,17",
            "discount_percent": "20,00",
            "tax": "",
            "amount": "3,34",
            "service": "Por classificar",
        },
        {
            "reference": "LAVACA",
            "description": "Lavagem Automática com Aspiração",
            "quantity": "1",
            "unit": "",
            "unit_price": "0,00",
            "list_price": "12,50",
            "discount_percent": "99,99",
            "tax": "",
            "amount": "0,00",
            "service": "Por classificar",
        },
    ]


def test_batch_invoice_payload_extracts_only_filinto_tal_table_rows():
    text = """
ORIGINAL
Filinto Mota Sucessores S.A.
NIF: 500 115 966
DOCUMENTO : TAL_FAC 2022/11146106
DE : 16/12/2022 O.R. : 305504 OFICINA : (001) CIT Reparação VENCIMENTO : 16/12/2022
MAT : AO-26-GP MARCA : CITROEN MODELO : Jumper
KMS : 26720 RECEPCIONISTA : Gil Teixeira CHASSIS : VF7VAYHVMMZ106533
Descrição Referência P.V.Unit Desc P.Liq.Unit Tmp/Qt Total Liq. CT
1ª REVISAO A
- OPERAÇÕES SISTEMÁTICAS DE MANUTENÇÃO - 93830000 49,00 25,00 36,75 1 36,75B
LÍQUIDO DE LAVA-VIDROS CONCENTRADO 1611908680 1,31 12,00 1,15 1 1,15B
JUNTA TAMPÃO BLOCO MOTOR 016488 2,19 8,00 2,01 1 2,01B
OLEO TOTAL INEO XTRA FIRST 0W20 LT QINEOXF 40,00 39,00 24,40 5,75 140,30B
FILTRO OLEO 1680682480 14,65 20,00 11,72 1 11,72B
SIGOU - Ecolub 1L 0,08 0,08 5,75 0,46B
Subtotal Peças (155,64)
Nova Intervenção 192,39
OFERTA DE LAVAGEM B Lavagem Simples LAVOFE 10,00 99,99 1 B
Nova Intervenção
CÓDIGO/DESCRIÇÃO I.V.A. TAXA I.V.A. BASE INCIDÊNCIA VALOR I.V.A. TOTAL LÍQUIDO TOTAL I.V.A. TOTAL A PAGAR
B APV TX NORMAL 23,00 192,39 44,25 192,39 44,25 236,64
"""

    payload = _batch_invoice_payload(b"", ".pdf", "tal.pdf", existing_text=text)
    descriptions = [line["description"] for line in payload["invoice_lines"]]

    assert len(payload["invoice_lines"]) == 7
    assert "1ª REVISAO" not in descriptions
    assert "Subtotal Peças" not in descriptions
    assert "Nova Intervenção" not in descriptions
    assert payload["invoice_lines"][-1]["description"] == "Lavagem Simples"
    assert payload["invoice_lines"][-1]["reference"] == "LAVOFE"
    assert payload["invoice_lines"][-1]["quantity"] == "1"


def test_batch_invoice_payload_keeps_filinto_package_components_without_individual_prices():
    text = """
ORIGINAL
Filinto Mota Sucessores S.A.
NIF: 500 115 966
DOCUMENTO : TAL_FAC 2025/11173485
DE : 18/06/2025 O.R. : 348678 OFICINA : (001) CIT Reparação VENCIMENTO : 18/06/2025
MAT : BB-69-TE MARCA : CITROEN MODELO : Berlingo Van XL 1.5 BlueHDi 100 S&S CVM6
KMS : 63170 RECEPCIONISTA : Gil Teixeira CHASSIS : VR7EFYHT2PJ697244
Descrição Referência P.V.Unit Desc P.Liq.Unit Tmp/Qt Total Liq. CT
INDICAÇÃO DE MANUTENÇÃO A
Mudança Óleo dinâmica diesel (5,3L) 2 135,43 135,43 135,43 B
Mão de Obra Mecânica MMC001 1
OLEO TOTAL QUARTZ INEO RCP 5W30 LT QINEORCP 5,30
FILTRO OLEO 1680682480 1
SIGOU - Ecolub 1L 5,30
JUNTA DO BUJAO 016488 1
Nova Intervenção 135,43
MUDANÇA DE OLEO DE TRAVÕES B
MUDANCA DE OLEO CIRCUITO DE TRAVAGEM MANUTENCAO 25022707 66,00 31,06 45,50 0,40 18,20 B
ÓLEO DOS TRAVÕES 1610725580 12,56 20,00 10,05 2 20,10 B
SIGOU - Ecolub 0,5L 0,04 0,04 2 0,08 B
Subtotal Peças (20,18)
Nova Intervenção 38,38
OFERTA DE LAVAGEM D
Lavagem Automática com Aspiração LAVACA 14,00 99,99 1 B
Nova Intervenção
Folha de Obra nº 580
GRUPO FILINTO MOTA - VANTAGEM CLIENTE: 27,22€ DE DESCONTO + I.V.A.
CÓDIGO/DESCRIÇÃO I.V.A. TAXA I.V.A. BASE INCIDÊNCIA VALOR I.V.A. TOTAL LÍQUIDO TOTAL I.V.A. TOTAL A PAGAR
B APV TX NORMAL 23,00 173,81 39,98 173,81 39,98 213,79
"""

    payload = _batch_invoice_payload(b"", ".pdf", "11173485.pdf", existing_text=text)
    lines = payload["invoice_lines"]

    assert payload["document_number"] == "TAL_FAC 2025/11173485"
    assert payload["work_order_reference"] == "580"
    assert payload["km"] == "63170"
    assert payload["total_with_vat"] == "213,79"
    assert len(lines) == 10
    assert lines[0]["description"] == "Mudança Óleo dinâmica diesel (5,3L)"
    assert lines[0]["amount"] == "135,43"
    assert lines[1]["reference"] == "MMC001"
    assert lines[1]["description"] == "Mão de Obra Mecânica"
    assert lines[2]["reference"] == "QINEORCP"
    assert lines[2]["quantity"] == "5,30"
    assert lines[3]["reference"] == "1680682480"
    assert lines[4]["description"] == "SIGOU - Ecolub 1L"
    assert lines[5]["reference"] == "016488"
    assert lines[6]["reference"] == "25022707"
    assert lines[6]["amount"] == "18,20"
    assert lines[7]["reference"] == "1610725580"
    assert lines[8]["description"] == "SIGOU - Ecolub 0,5L"
    assert lines[9]["reference"] == "LAVACA"
    assert lines[9]["amount"] == "0,00"


def test_batch_invoice_filinto_stacked_lines_keep_invoice_columns():
    lines = [
        "ORIGINAL",
        "Filinto Mota Sucessores S.A.",
        "NIF: 500 115 966",
        "DOCUMENTO :",
        "TAL_FAC 2022/11146106",
        "DE :",
        "16/12/2022",
        "Descrição",
        "Referência",
        "P.V.Unit Desc P.Liq.Unit Tmp/Qt Total Liq.",
        "CT",
        "1ª REVISAO",
        "A",
        "- OPERAÇÕES SISTEMÁTICAS DE MANUTENÇÃO -",
        "93830000",
        "49,00",
        "25,00",
        "36,75",
        "1",
        "36,75 B",
        "LÍQUIDO DE LAVA-VIDROS CONCENTRADO",
        "1611908680",
        "1,31",
        "12,00",
        "1,15",
        "1",
        "1,15 B",
        "JUNTA TAMPÃO BLOCO MOTOR",
        "016488",
        "2,19",
        "8,00",
        "2,01",
        "1",
        "2,01 B",
        "OLEO TOTAL INEO XTRA FIRST 0W20 LT",
        "QINEOXF",
        "40,00",
        "39,00",
        "24,40",
        "5,75",
        "140,30 B",
        "FILTRO OLEO",
        "1680682480",
        "14,65",
        "20,00",
        "11,72",
        "1",
        "11,72 B",
        "SIGOU - Ecolub 1L",
        "0,08",
        "0,08",
        "5,75",
        "0,46 B",
        "Subtotal Peças (155,64)",
        "Nova Intervenção 192,39",
        "OFERTA DE LAVAGEM",
        "B",
        "Lavagem Simples",
        "LAVOFE",
        "10,00",
        "99,99",
        "1",
        "B",
        "Nova Intervenção",
        "OBSERVAÇÕES:",
    ]

    invoice_lines = _batch_invoice_filinto_stacked_lines(lines)

    assert len(invoice_lines) == 7
    assert invoice_lines[0] == {
        "reference": "93830000",
        "description": "OPERAÇÕES SISTEMÁTICAS DE MANUTENÇÃO",
        "quantity": "1",
        "unit": "",
        "unit_price": "36,75",
        "list_price": "49,00",
        "discount_percent": "25,00",
        "tax": "",
        "amount": "36,75",
        "service": "Manutenção",
    }
    assert invoice_lines[3]["reference"] == "QINEOXF"
    assert invoice_lines[3]["quantity"] == "5,75"
    assert invoice_lines[3]["amount"] == "140,30"
    assert invoice_lines[-1]["reference"] == "LAVOFE"
    assert invoice_lines[-1]["amount"] == "0,00"


def test_batch_invoice_payload_extracts_filinto_columnar_pdf_order():
    text = """
Cap. Social 1.250.000 € - Matriculada na CRC Porto
nº 500 115 966 Sede: Filinto Mota Sucessores S.A. Rua Pinto Bessa, 550
TAL_FAC
19/12/2024
339843
O.R. :
42608
VXKUPHMHDP4099734
BB-96-DU
Descrição
KMS :
2024/11167597
DOCUMENTO :
FACTURA
Exmos Senhores Carfast - Rent-A-Car, Lda
NIF: 500 115 966
ORIGINAL
A
REVISÃO 40.000 KM
B
43,75
1
43,75
30,00
62,50
93830000
OPERAÇÕES SISTEMÁTICAS DE MANUTENÇÃO----
B
4,38
0,10
43,75
30,00
62,50
06250907
VELAS (JOGO)-SUBSTITUICAO-MANUTENCAO
Subtotal Mão Obra (56,89)
B
71,50
3,25
22,00
45,00
40,00
QINEOXF
OLEO TOTAL INEO XTRA FIRST 0W20 B712010
B
0,20
3,25
0,06
0,06
SIGOU - Ecolub 1L
TOTAL A PAGAR
332,19
TOTAL I.V.A.
62,12
TOTAL LÍQUIDO
270,07
"""

    payload = _batch_invoice_payload(b"", ".pdf", "Fatura-11167597-0.pdf", existing_text=text)
    descriptions = [line["description"] for line in payload["invoice_lines"]]

    assert payload["document_number"] == "TAL_FAC 2024/11167597"
    assert payload["total_with_vat"] == "332,19"
    assert "ORIGINAL" not in descriptions
    assert "REVISÃO 40.000 KM" not in descriptions
    assert payload["invoice_lines"][0]["reference"] == "93830000"
    assert payload["invoice_lines"][0]["description"] == "OPERAÇÕES SISTEMÁTICAS DE MANUTENÇÃO"
    assert payload["invoice_lines"][0]["quantity"] == "1"
    assert payload["invoice_lines"][0]["unit_price"] == "43,75"
    assert payload["invoice_lines"][0]["list_price"] == "62,50"
    assert payload["invoice_lines"][0]["discount_percent"] == "30,00"
    assert payload["invoice_lines"][0]["amount"] == "43,75"
    assert payload["invoice_lines"][-1]["description"] == "SIGOU - Ecolub 1L"
    assert len(payload["invoice_lines"]) == 4


def test_batch_invoice_payload_extracts_filinto_inline_pdf_order_without_column_shift():
    text = """
Cap. Social 1.250.000 € - Matriculada na CRC Porto
nº 500 115 966 Sede: Filinto Mota Sucessores S.A. Rua Pinto Bessa, 550
FACTURA
Modo de Pagamento : Pronto Pagamento
CLIENTE : 227010 NIF :509285970
DOCUMENTO : TAL_FAC 2022/11144696
DE : 28/10/2022 O.R. : 303476 OFICINA : (81) MMarca Manutenção VENCIMENTO : 28/10/2022
MAT : AQ-41-XJ MARCA : PEUGEOT MODELO : 308
KMS : 11015 RECEPCIONISTA : GTEIXEIRA CHASSIS : VR3FRHNSLNY544049
Descrição Referência P.V.Unit Desc P.Liq.Unit Tmp/Qt Total Liq. CT
SUBSTITUIR FAROLIM A
FAROLIM DE TRÁS DIREITO - SUBSTITUIÇÃO MMC001 49,00 25,00 36,75 0,50 18,38 B
FAROLIM TRAS 9835300580 177,00 20,00 141,60 1 141,60 B
Nova Intervenção 159,98
CÓDIGO/DESCRIÇÃO I.V.A. TAXA I.V.A. BASE INCIDÊNCIA VALOR I.V.A. TOTAL LÍQUIDO TOTAL I.V.A. TOTAL A PAGAR
B APV TX NORMAL 23,00 159,98 36,80 159,98 36,80 196,78
"""

    payload = _batch_invoice_payload(b"", ".pdf", "1-TAL-2-11144696-227010-1211374.pdf", existing_text=text)

    assert payload["document_number"] == "TAL_FAC 2022/11144696"
    assert payload["document_date"] == "2022-10-28"
    assert payload["total_with_vat"] == "196,78"
    assert payload["repair_order_reference"] == "303476"
    assert len(payload["invoice_lines"]) == 2
    assert payload["invoice_lines"][0] == {
        "reference": "MMC001",
        "description": "FAROLIM DE TRÁS DIREITO - SUBSTITUIÇÃO",
        "quantity": "0,50",
        "unit": "",
        "unit_price": "36,75",
        "list_price": "49,00",
        "discount_percent": "25,00",
        "tax": "B",
        "amount": "18,38",
        "service": "Por classificar",
    }
    assert payload["invoice_lines"][1] == {
        "reference": "9835300580",
        "description": "FAROLIM TRAS",
        "quantity": "1",
        "unit": "",
        "unit_price": "141,60",
        "list_price": "177,00",
        "discount_percent": "20,00",
        "tax": "B",
        "amount": "141,60",
        "service": "Por classificar",
    }


def test_batch_invoice_payload_recovers_filinto_stacked_pdf_order_from_column_shift():
    text = """
Mod. PVP.05.2
nº 500 115 966 Sede: Filinto Mota Sucessores S.A. Rua Pinto Bessa, 550
FACTURA
DOCUMENTO :
2022/11144696
DE :
MAT :
KMS :
Descrição
28/10/2022
OFICINA :
(81) MMarca Manutenção
AQ-41-XJ
MARCA : PEUGEOT
MODELO : 308
CHASSIS : VR3FRHNSLNY544049
11015
O.R. : 303476
RECEPCIONISTA : GTEIXEIRA
VENCIMENTO : 28/10/2022
Referência
P.V.Unit Desc P.Liq.Unit Tmp/Qt Total Liq.
CT
Pronto Pagamento
Modo de Pagamento :
TAL_FAC
CLIENTE :
227010
Exmos Senhores Carfast - Rent-A-Car, Lda
ORIGINAL
NIF: 500 115 966
SUBSTITUIR FAROLIM
A
FAROLIM DE TRÁS DIREITO - SUBSTITUIÇÃO
MMC001
49,00 25,00
36,75
0,50
18,38 B
FAROLIM TRAS
9835300580
177,00 20,00
141,60
1
141,60 B
Nova Intervenção 159,98
.
B
23,00
159,98
36,80
CÓDIGO/DESCRIÇÃO I.V.A. TAXA I.V.A. BASE INCIDÊNCIA
APV TX NORMAL
VALOR I.V.A.
TOTAL LÍQUIDO
159,98
TOTAL I.V.A.
36,80
TOTAL A PAGAR
196,78
"""

    payload = _batch_invoice_payload(b"", ".pdf", "1-TAL-2-11144696-227010-1211374.pdf", existing_text=text)

    assert payload["document_number"] == "TAL_FAC 2022/11144696"
    assert payload["total_with_vat"] == "196,78"
    assert payload["vat_amount"] == "36,80"
    assert payload["subtotal_without_vat"] == "159,98"
    assert len(payload["invoice_lines"]) == 2
    assert payload["invoice_lines"][0]["reference"] == "MMC001"
    assert payload["invoice_lines"][0]["description"] == "FAROLIM DE TRÁS DIREITO - SUBSTITUIÇÃO"
    assert payload["invoice_lines"][0]["quantity"] == "0,50"
    assert payload["invoice_lines"][0]["amount"] == "18,38"
    assert payload["invoice_lines"][1]["reference"] == "9835300580"
    assert payload["invoice_lines"][1]["description"] == "FAROLIM TRAS"
    assert payload["invoice_lines"][1]["quantity"] == "1"
    assert payload["invoice_lines"][1]["amount"] == "141,60"


def test_batch_invoice_payload_extracts_filinto_package_invoice_line():
    text = """
Cap. Social 1.250.000 € - Matriculada na CRC Porto
nº 500 115 966 Sede: Filinto Mota Sucessores S.A. Rua Pinto Bessa, 550
TAL_FAC
31/12/2024
340423
O.R. :
14144
VR3UDYHZSPJ925908
BI-73-EL
Descrição
2024/11167923
DOCUMENTO :
FACTURA
NIF: 500 115 966
ORIGINAL
A
INDICAÇÃO DE MANUTENÇÃO (ÓLEO E FILTRO)
Nova Intervenção
B
Mudança Óleo dinâmica diesel (3,95L)
B
101,62
101,62
101,62
1
Mudança Óleo dinâmica diesel (3,95L)
1
1680682480
FILTRO OLEO
1
016488
JUNTA DO BUJAO
3,95
QINEORCP
OLEO TOTAL QUARTZ INEO RCP 5W30 FPW9.55535/03
1
MMC001
Mão de Obra Mecânica
3,95
SIGOU - Ecolub 1L
Nova Intervenção 101,62
C
OFERTA DE LAVAGEM
B
1
99,99
14,00
LAVACA
Lavagem Automática com Aspiração
23,37
101,62
23,00
B
11167923
JFS4WFN7
€ DE DESCONTO + I.V.A.
14,00
124,99
TOTAL A PAGAR
23,37
TOTAL I.V.A.
101,62
TOTAL LÍQUIDO
"""

    payload = _batch_invoice_payload(b"", ".pdf", "Fatura-11167923-0.pdf", existing_text=text)

    assert payload["document_number"] == "TAL_FAC 2024/11167923"
    assert payload["total_with_vat"] == "124,99"
    assert payload["invoice_lines"][0]["description"] == "Mudança Óleo dinâmica diesel (3,95L)"
    assert payload["invoice_lines"][0]["quantity"] == "1"
    assert payload["invoice_lines"][0]["unit_price"] == "101,62"
    assert payload["invoice_lines"][0]["amount"] == "101,62"
    assert payload["invoice_lines"][0]["service"] == "Manutenção"
    assert payload["invoice_lines"][0]["quantity"] != "1680682480"


def test_batch_invoice_payload_extracts_filinto_collision_parts_with_spaced_thousands():
    text = """
Filinto Mota Sucessores S.A.
NIF: 500 115 966
FACTURA
DOCUMENTO : TAL_FAC 2025/11167979
DE : 02/01/2025 O.R. : 340280 OFICINA : (082) MMarca Colisão VENCIMENTO : 02/01/2025
MAT : BH-86-EB MARCA : PEUGEOT MODELO : Partner KMS : 29452
CHASSIS : VR3EFYHT2PJ918611
Descrição Referência P.V.Unit Desc P.Liq.Unit Tmp/Qt Total Liq. CT
SINSITRO FRENTE A
Mão de Obra Chapa MCH001 67,00 15,00 56,95 9,94 566,08 B
Mão de Obra Pintura MPT001 67,00 15,00 56,95 3,51 199,89 B
Ingredientes de Pintura INGPNT 83,92 15,00 71,33 1 71,33 B
Subtotal Mão Obra (837,30)
Orcamento Pecas ORCP 3 522,86 15,00 2 994,43 1 2 994,43 B
Nova Intervenção 3 831,73
B APV TX NORMAL 23,00 3 831,73 881,30 3 831,73 881,30 4 713,03
TOTAL A PAGAR 4 713,03
ATCUD:JFS4WFN7-11167979
"""

    payload = _batch_invoice_payload(b"", ".pdf", "Fatura-11167979-0.pdf", existing_text=text)
    invoice_lines = payload["invoice_lines"]

    assert payload["document_number"] == "TAL_FAC 2025/11167979"
    assert payload["work_order_reference"] == ""
    assert payload["km"] == "29452"
    assert payload["total_with_vat"] == "4713,03"
    assert len(invoice_lines) == 4
    assert invoice_lines[-1] == {
        "reference": "ORCP",
        "description": "Orcamento Pecas",
        "quantity": "1",
        "unit": "",
        "unit_price": "2994,43",
        "list_price": "3522,86",
        "discount_percent": "15,00",
        "tax": "B",
        "amount": "2994,43",
        "service": "Por classificar",
    }
    assert _batch_invoice_total_from_lines(invoice_lines) == "3831,73"


def test_batch_invoice_payload_extracts_filinto_package_line_with_shifted_quantity():
    text = """
Filinto Mota Sucessores S.A.
NIF: 500 115 966
TAL_FAC
CT
Total Liq.
Tmp/Qt
P.Liq.Unit
Desc
P.V.Unit
Referência
26/02/2025
VENCIMENTO :
Gil Teixeira
RECEPCIONISTA :
343355
O.R. :
11903
VR3EDYHT4RJ747033
CHASSIS :
Partner Asphalt Standard 1.5 BlueHDi 100cv CVM5
MODELO :
PEUGEOT
MARCA :
BN-35-MO
(081) MMarca Reparação
OFICINA :
26/02/2025
Descrição
KMS :
MAT :
DE :
2025/11169922
DOCUMENTO :
FACTURA
Exmos Senhores Carfast - Rent-A-Car, Lda
CLIENTE :
NIF : 509285970
ORIGINAL
A
INDICAÇÃO DE MANUTENÇÃO
B
135,43
135,43
135,43
2
Mudança Óleo dinâmica diesel (5,3L)
1
MMC001
Mão de Obra Mecânica
5,30
QINEORCP
OLEO TOTAL QUARTZ INEO RCP 5W30 LT
1
1680682480
FILTRO OLEO
5,30
SIGOU - Ecolub 1L
1
016488
JUNTA DO BUJAO
Nova Intervenção 135,43
B
OFERTA DE LAVAGEM
B
1
99,99
14,00
LAVACA
Lavagem Automática com Aspiração
Nova Intervenção
Folha de Obra nº 329
VALOR I.V.A.
CÓDIGO/DESCRIÇÃO I.V.A.
31,15
135,43
23,00
B
11169922
JFS4WFN7
€ DE DESCONTO + I.V.A.
14,00
166,58
TOTAL A PAGAR
31,15
TOTAL I.V.A.
135,43
TOTAL LÍQUIDO
"""

    payload = _batch_invoice_payload(b"", ".pdf", "Fatura-11169922-0.pdf", existing_text=text)

    assert payload["document_number"] == "TAL_FAC 2025/11169922"
    assert payload["total_with_vat"] == "166,58"
    assert len(payload["invoice_lines"]) == 1
    assert payload["invoice_lines"][0]["description"] == "Mudança Óleo dinâmica diesel (5,3L)"
    assert payload["invoice_lines"][0]["quantity"] == "1"
    assert payload["invoice_lines"][0]["unit_price"] == "135,43"
    assert payload["invoice_lines"][0]["amount"] == "135,43"


def test_batch_invoice_payload_extracts_filinto_inline_package_without_tal_fallback():
    text = """
Filinto Mota Sucessores S.A.
NIF: 500 115 966
DOCUMENTO : TAL_FAC 2025/11179616
DE : 31/12/2025 O.R. : 356423 OFICINA : (001) CIT Reparação VENCIMENTO : 31/01/2026
MAT : BC-98-FA MARCA : CITROEN MODELO : Berlingo Van XL 1.5 BlueHDi 100 S&S CVM6
KMS : 103153 RECEPCIONISTA : Gil Teixeira CHASSIS : VR7EFYHT2PJ721791
Descrição Referência P.V.Unit Desc P.Liq.Unit Tmp/Qt Total Liq. CT
INDICAÇÃO DE DEFEITO DO MOTOR C F
Nova Intervenção
INDICAÇÃO DE MANUTENÇÃO EM 300 KM, FEZ TROCA DE ÓELO D F
AOS 95.841 KM
Mudança Óleo dinâmica diesel (5,3L) 2 135,43 135,43 135,43 F
OLEO TOTAL QUARTZ INEO RCP 5W30 FPW9.55535/03 QINEORCP 5,30 B
JUNTA DO BUJAO 016488 1 B
FILTRO OLEO 1680682480 1 B
Mão de Obra Mecânica MMC001 1 F
SIGOU - Ecolub 1L 5,30 F
Nova Intervenção 135,43
CONTROLO DE LUZES E F
LAMPADA DE CHAPA DE MATRICULA-SUBSTITUICAO-NO 52890910 66,00 31,06 45,50 0,20 9,10 F
VEICULO
ABRA#ADEIRA 7588E8 0,16 12,00 0,14 1 0,14B
LAMPADA 12V-W5W 6216A1 2,70 20,00 2,16 2 4,32B
Subtotal Peças (4,46)
Nova Intervenção 13,56
LAVAGEM DE OFERTA F F
Lavagem Automática com Aspiração LAVACA 15,00 99,99 1 F
CÓDIGO/DESCRIÇÃO I.V.A. TAXA I.V.A.BASE INCIDÊNCIA VALOR I.V.A. TOTAL LÍQUIDO TOTAL I.V.A. TOTAL A PAGAR
B APV TX NORMAL 23,00 112,53 25,88 148,99 34,27 183,26
ATCUD:JFS4WFN7-11179616
"""

    payload = _batch_invoice_payload(b"", ".pdf", "Fatura-11179616-0.pdf", existing_text=text)

    assert payload["document_number"] == "TAL_FAC 2025/11179616"
    assert payload["document_date"] == "2025-12-31"
    assert payload["km"] == "103153"
    assert payload["total_with_vat"] == "183,26"
    assert payload["repair_order_reference"] == "356423"
    assert len(payload["invoice_lines"]) == 4
    assert [line["amount"] for line in payload["invoice_lines"]] == ["135,43", "9,10", "0,14", "4,32"]
    assert payload["invoice_lines"][0]["description"] == "Mudança Óleo dinâmica diesel (5,3L)"
    assert payload["invoice_lines"][0]["quantity"] == "1"
    assert payload["invoice_lines"][1]["reference"] == "52890910"
    assert payload["invoice_lines"][1]["description"] == "LAMPADA DE CHAPA DE MATRICULA-SUBSTITUICAO-NO"
    assert payload["invoice_lines"][2]["reference"] == "7588E8"
    assert payload["invoice_lines"][2]["unit_price"] == "0,14"
    assert payload["invoice_lines"][3]["reference"] == "6216A1"
    assert all("AOS 95" not in line["description"] for line in payload["invoice_lines"])
    assert all("LAVAGEM" not in line["description"] for line in payload["invoice_lines"])


def test_batch_invoice_payload_extracts_filinto_package_line_without_explicit_service_quantity():
    text = """
Filinto Mota Sucessores S.A.
NIF: 500 115 966
TAL_FAC
CT
Total Liq.
Tmp/Qt
P.Liq.Unit
Desc
P.V.Unit
Referência
26/02/2025
KMS :
DE :
2025/11168657
DOCUMENTO :
FACTURA
Descrição
ORIGINAL
A
INDICAÇÃO DE MANUTENÇÃO
Nova Intervenção
B
Mudança Óleo dinâmica diesel (5,3L)
B
135,43
135,43
135,43
2
Mudança Óleo dinâmica diesel (5,3L)
5,30
QINEORCP
OLEO TOTAL QUARTZ INEO RCP 5W30 FPW9.55535/03
1
016488
JUNTA DO BUJAO
1
1680682480
FILTRO OLEO
1
MMC001
Mão de Obra Mecânica
5,30
SIGOU - Ecolub 1L
GRUPO FILINTO MOTA - VANTAGEM CLIENTE:
APV TX NORMAL
BASE INCIDÊNCIA
TAXA I.V.A.
CÓDIGO/DESCRIÇÃO I.V.A.
31,15
135,43
23,00
B
11168657
JFS4WFN7
€ DE DESCONTO + I.V.A.
14,00
166,58
TOTAL A PAGAR
31,15
TOTAL I.V.A.
135,43
TOTAL LÍQUIDO
"""

    payload = _batch_invoice_payload(b"", ".pdf", "Fatura-11168657-0.pdf", existing_text=text)

    assert payload["document_number"] == "TAL_FAC 2025/11168657"
    assert len(payload["invoice_lines"]) == 1
    assert payload["invoice_lines"][0]["description"] == "Mudança Óleo dinâmica diesel (5,3L)"
    assert payload["invoice_lines"][0]["quantity"] == "1"
    assert payload["invoice_lines"][0]["amount"] == "135,43"


def test_batch_invoice_payload_extracts_filinto_franchise_without_header_contamination():
    text = """
Filinto Mota Sucessores S.A.
Rua Pinto Bessa, 546
4300-428 Porto
NIF: 500 115 966
ORIGINAL
2ª VIA
402,40
Franquia referente ao veículo com a matricula AQ-91-AA, apolice
Nº 0004475646000324 da companhia de seguros Tranquilidade.
VALOR I.V.A.
GRUPO FILINTO MOTA - VANTAGEM CLIENTE:
APV TX NORMAL
BASE INCIDÊNCIA
TAXA I.V.A.
CÓDIGO/DESCRIÇÃO I.V.A.
92,55
402,40
23,00
B
11158298
JFS4WFN7
€ DE DESCONTO + I.V.A.
0,00
494,95
TOTAL A PAGAR
92,55
TOTAL I.V.A.
402,40
TOTAL LÍQUIDO
"""

    payload = _batch_invoice_payload(b"", ".pdf", "11158298.pdf", existing_text=text)

    assert len(payload["invoice_lines"]) == 1
    assert payload["invoice_lines"][0]["reference"] == "FRANQUIA"
    assert payload["invoice_lines"][0]["description"].startswith("Franquia referente")
    assert payload["invoice_lines"][0]["amount"] == "402,40"
    assert "ORIGINAL" not in {line["description"] for line in payload["invoice_lines"]}


def test_batch_invoice_payload_treats_filinto_operation_duration_as_time():
    text = """
Filinto Mota Sucessores S.A.
NIF: 500 115 966
DOCUMENTO : TAL_FAC 2025/11170799
Descrição Referência P.V.Unit Desc P.Liq.Unit Tmp/Qt Total Liq. CT
(22514510) 2 RODAS-EQUILIBRAGEM-NO VEICULO 0,60
CÓDIGO/DESCRIÇÃO I.V.A. TAXA I.V.A. BASE INCIDÊNCIA VALOR I.V.A.
"""

    payload = _batch_invoice_payload(b"", ".pdf", "tal.pdf", existing_text=text)

    assert payload["invoice_lines"] == [
        {
            "reference": "22514510",
            "description": "2 RODAS-EQUILIBRAGEM-NO VEICULO",
            "quantity": "0,60",
            "unit": "",
            "unit_price": "",
            "list_price": "",
            "discount_percent": "",
            "tax": "",
            "amount": "",
            "service": "Pneus",
        }
    ]


def test_batch_invoice_payload_extracts_cruz_allen_invoice_layout():
    text = """
FATURA
VIA NÚMERO DATA
Original FS0003721 15/10/25
Exmo(s) Senhor(es):
CARFAST, RENT-A-CAR, LDA
R. DAS INDUSTRIAS, 220
4785-625-TROFA
Cliente nº NIF Cliente Nº OR Página
000310 509285970 OR0008491 1 / 1
Forma Paga.: PAGAM A 60 DIAS Recepcionista: CARLOS OLIVEIRA
Modelo: JUMPER III Chassis: VF7YBBPFCPG030533
Matrícula: BG-84-TN Data de Garantia: 05/04/24
Kilómetros: 58669 Data de Entrega: 15/10/2025
Referência Descrição Qtd. Preço Unitário Dto. IVA Valor
- Mao Obra
25210910 SUBSTITUIR PLACAS DE TRAVÃO FTE ,70 41,90 € 15,00 23% 24,93 €
- Peças
1617273980 E:4 PASTI TR D 1,00 89,37 € 25,00 23% 67,03 €
- Outros Débitos
ENSAIO SIMPLES-OFERTA 1,00 16,00 € 99,99 23%
LIMPEZA DA VIATURA-OFERTA 1,00 18,00 € 99,99 23%
P & D PICKUP & DELIVERY 1,00 0,01 € 99,99 23%
OBS: V/FOLHA OBRA Nº793 1,00 0,01 € 99,99 23%
Observações:
Total Mão de Obra: 29,33 €
Total Peças: 89,37 €
Outros Déb. + T. Ext: 34,02 €
Total Descontos 60,76 €
Total Líquido: 91,96 €
I.V.A: 23 % 21,15 €
Total Fatura: 113,11 €
CRUZ & ALLEN - B.B.C OFICINA DE AUTOMÓVEIS, LDA
CONTRIBUINTE Nº 504 104 250
"""

    payload = _batch_invoice_payload(b"", ".pdf", "20B0CA~1.PDF", existing_text=text)

    assert payload["ocr_template"] == "cruz_allen_bbc_invoice"
    assert payload["supplier_name"] == "Cruz & Allen - B.B.C Oficina de Automóveis, Lda"
    assert payload["supplier_nif"] == "504104250"
    assert payload["document_number"] == "FS0003721"
    assert payload["document_date"] == "2025-10-15"
    assert payload["client_name"] == "CARFAST, RENT-A-CAR, LDA R. DAS INDUSTRIAS, 220 4785-625-TROFA"
    assert payload["client_number"] == "000310"
    assert payload["client_nif"] == "509285970"
    assert payload["repair_order_reference"] == "OR0008491"
    assert payload["work_order_reference"] == "793"
    assert payload["payment_method"] == "PAGAM A 60 DIAS"
    assert payload["receptionist"] == "CARLOS OLIVEIRA"
    assert payload["vehicle_model"] == "JUMPER III"
    assert payload["vin"] == "VF7YBBPFCPG030533"
    assert payload["plate"] == "BG-84-TN"
    assert payload["km"] == "58669"
    assert payload["warranty_date"] == "2024-04-05"
    assert payload["delivery_date"] == "2025-10-15"
    assert payload["labor_total"] == "29,33"
    assert payload["materials_total"] == "89,37"
    assert payload["misc_total"] == "34,02"
    assert payload["discount_without_vat"] == "60,76"
    assert payload["subtotal_without_vat"] == "91,96"
    assert payload["vat_amount"] == "21,15"
    assert payload["total_with_vat"] == "113,11"
    assert len(payload["invoice_lines"]) == 6
    assert payload["invoice_lines"][0] == {
        "line_type": "MO",
        "section": "Mão de obra",
        "reference": "25210910",
        "description": "SUBSTITUIR PLACAS DE TRAVÃO FTE",
        "quantity": "0,70",
        "unit": "",
        "unit_price": "41,90",
        "discount_percent": "15,00",
        "tax": "23%",
        "amount": "24,93",
        "service": "Calços",
    }
    assert payload["invoice_lines"][1]["reference"] == "1617273980"
    assert payload["invoice_lines"][1]["amount"] == "67,03"
    assert payload["invoice_lines"][-1]["description"] == "OBS: V/FOLHA OBRA Nº793"


def test_batch_invoice_payload_extracts_cruz_allen_scanned_invoice_ocr_text():
    text = """
FATURA
VIA NÚMERO DATA
Original FS0002239 11/07/24
Exmo(s) Senhor(es):
CARFAST, RENT-A-CAR, LDA
Rua das Industrias, 220
4785-625 TROFA
Cliente n° NIF Cliente N°OR Página
000310 509285970 ORO006547 Ue à
Forma Pag PAGAM. A 30 DIAS Recepcionista: SONIA SILVA
Modelo: JUMPER III Chassis: VFTYBBPFCPGO30534
Matricula: BD-87-LZ Data de Garantia: 01/09/23
Kilómetros: 41808 Data de Entrega: 11/07/2024
Referência Descrição Qtd. Preço Unitário Dto. IVA Valor
= Mao Obra
00000510 CONTROLO VEICULO 410 39,00€ 10,00 23% 351€
01022815 MUDAR OLEO MOTOR E FILTRO OLEO - +50 39,00€ 10,00 23% 1755 €
OPERAÇÃO SERVIÇO
= Peças
9809532380 FILTRO OLEO 1,00 20,83€ 10,00 23% 1875 €
0w30 OLEOQUARTZINEO FIRST 0w30 6,60 35,62 € 20,00 23% 188,07 €
OW30:GAU ECOLUB 1,00 0,53 € 23% 0,53€
- Outros Débitos
016488 JUNTA BUJÃO 1,00 2,39€ 23% 2,39€
- Outros Débitos
ENSAIO SIMPLES - OFERTA 1,00 16,00 € 99,99
P&D PICK UP & DELIVERY 1,00 1,00 € 99,99
ALERTA VIATURA PRECISA DE VIR OFICINA EFETUAR A 1,00 1,00€ 99,99
MANUTENÇÃO DOS 50.000KMS
Observações:
Total Mão de Obra: 23,40 €
Total Peças: 256,45 €
Outros Déb. + T. Ext: 20,39 €
Total Descontos 6944 €
Total Líquido: 23080 €
I.V.A: 23 % 5308 €
Total Fatura: 28388 €
CRUZ & ALLEN -B.B.C OFICINA DE AUTOMÓVEIS, LDA
MATRÍCULA Nº 504 104 250
Id. Único: Fact FS/0002239
"""

    payload = _batch_invoice_payload(b"", ".pdf", "FATURA FS2239 CARFAST.pdf", existing_text=text)

    assert payload["ocr_template"] == "cruz_allen_bbc_invoice"
    assert payload["text_source"] == "pdf_text"
    assert payload["supplier_nif"] == "504104250"
    assert payload["document_number"] == "FS0002239"
    assert payload["document_date"] == "2024-07-11"
    assert payload["client_number"] == "000310"
    assert payload["client_nif"] == "509285970"
    assert payload["repair_order_reference"] == "OR0006547"
    assert payload["payment_method"] == "PAGAM. A 30 DIAS"
    assert payload["receptionist"] == "SONIA SILVA"
    assert payload["vehicle_model"] == "JUMPER III"
    assert payload["vin"] == "VF7YBBPFCPG030534"
    assert payload["plate"] == "BD-87-LZ"
    assert payload["km"] == "41808"
    assert payload["warranty_date"] == "2023-09-01"
    assert payload["delivery_date"] == "2024-07-11"
    assert payload["labor_total"] == "23,40"
    assert payload["materials_total"] == "256,45"
    assert payload["misc_total"] == "20,39"
    assert payload["discount_without_vat"] == "69,44"
    assert payload["subtotal_without_vat"] == "230,80"
    assert payload["vat_amount"] == "53,08"
    assert payload["total_with_vat"] == "283,88"
    assert len(payload["invoice_lines"]) == 9
    assert payload["invoice_lines"][0]["quantity"] == "0,10"
    assert payload["invoice_lines"][0]["amount"] == "3,51"
    assert payload["invoice_lines"][1]["quantity"] == "0,50"
    assert payload["invoice_lines"][1]["amount"] == "17,55"
    assert payload["invoice_lines"][4]["reference"] == "OW30:GAU"
    assert payload["invoice_lines"][4]["description"] == "ECOLUB"
    assert payload["invoice_lines"][4]["amount"] == "0,53"
    assert payload["invoice_lines"][6]["description"] == "ENSAIO SIMPLES - OFERTA"
    assert payload["invoice_lines"][6]["amount"] == "0,00"
    assert payload["invoice_lines"][7]["description"] == "P&D PICK UP & DELIVERY"
    assert payload["invoice_lines"][7]["amount"] == "0,00"
    assert payload["invoice_lines"][8]["description"].endswith("MANUTENÇÃO DOS 50.000KMS")
    assert payload["invoice_lines"][8]["amount"] == "0,00"


def test_batch_invoice_payload_does_not_treat_cruz_allen_customer_as_supplier():
    text = """
FERREIRA DA COSTA & IRMAO LDA
RUA NOVAIS DA CUNHA, 973
NIF: 500114013
DOC NR: FR F E1331001FAC1E/22037832
DATA/HORA: 12-02-25 16:20
ARTIGO QT. P.Un. VALOR
GASOLEO 1 42,39 1,764 74,78
TOTAL EUROS 74,78
N.I.F.: 504104250
Nome: CRUZ ALLEN BBC
Morada: GONDOMAR
"""

    payload = _batch_invoice_payload(b"", ".pdf", "Scan2025-02-13_100117.pdf", existing_text=text)

    assert payload["ocr_template"] != "cruz_allen_bbc_invoice"
    assert payload["supplier_nif"] != "504104250"
    assert payload["supplier_name"] != "Cruz & Allen - B.B.C Oficina de Automóveis, Lda"


def test_batch_invoice_payload_normalizes_cruz_allen_scanned_header_km_and_order():
    text = """
FATURA
VIA NÚMERO DATA
FS 0002342 17/09/24
Exmo(s) Senhor(es):
CARFAST, RENT-A-CAR, LDA
Cliente nº NIF Cliente Nº OR Página
000310 509285970 OR0O06861 1 / 1
Modelo: JUMPER III Chassis: VF7YAAPFBPG032533
Matrícula: BD-57-MC
Kms. 28822 Data de Entrega: 17/09/2024
Referência Descrição Qtd. Preço Unitário Dto. IVA Valor
- Mao Obra
93830000 REVISAO 1,00 41,90 € 15,00 23% 35,61 €
Observações:
Total Fatura: 43,80 €
CRUZ & ALLEN - B.B.C OFICINA DE AUTOMÓVEIS, LDA
CONTRIBUINTE Nº 504 104 250
"""

    payload = _batch_invoice_payload(b"", ".pdf", "Scan2024-09-17_105229.pdf", existing_text=text)

    assert payload["ocr_template"] == "cruz_allen_bbc_invoice"
    assert payload["document_number"] == "FS0002342"
    assert payload["document_date"] == "2024-09-17"
    assert payload["repair_order_reference"] == "OR0006861"
    assert payload["km"] == "28822"


def test_batch_invoice_payload_uses_canonical_cruz_allen_filename_when_ocr_misreads_number():
    text = """
FATURA
VIA NÚMERO DATA
Original FS0002322 05/09/24
Exmo(s) Senhor(es):
CARFAST, RENT-A-CAR, LDA
Cliente nº NIF Cliente Nº OR Página
000310 509285970 OR0006801 1 / 1
Modelo: BERLINGO Chassis: VR7EFYHT2PJ721795
Matrícula: BC-92-EE
Kilómetros: 25380
Referência Descrição Qtd. Preço Unitário Dto. IVA Valor
93830000 REVISAO 1,00 41,90 € 15,00 23% 35,61 €
Total Fatura: 43,80 €
CRUZ & ALLEN - B.B.C OFICINA DE AUTOMÓVEIS, LDA
CONTRIBUINTE Nº 504 104 250
"""

    payload = _batch_invoice_payload(b"", ".pdf", "FS-0002320.pdf", existing_text=text)

    assert payload["ocr_template"] == "cruz_allen_bbc_invoice"
    assert payload["document_number"] == "FS0002320"


def test_batch_invoice_payload_handles_cruz_allen_degraded_header_and_split_totals():
    text = """
FATURA
CITROEN VIA NÚMERO, DATA
Original FS0002235 10/07/24
Cruz & Allen CARFAST, RENT-A-CAR, LDA
Cliente n° NIF Cliente Página
000310 509285970 OR 1/1
Modelo: BERLINGO VU (K9) FUR Chassis: VR7EFYHT2PJ721795
Matricula: BC-92-EE
Kilómetros: 25380 Data de Entrega: 10/07/2024
Referência Descrição ASS Qtd. Preço Unitário Dto. IVA Valor
- Mao Obra
00000510 CONTROLO VEICULO ,10 39,00€ 10,00 23% 3,51€
01022815 MUDANCA DE OLEO E FILTRO MOTOR 40 39,00€ 10,00 23% 14,04 €
- Peças
1680682480 FILTRO OLEO 1,00 15,93€ 10,00 23% 14,34 €
OW20 OLEOQUARTZINEO FIRST OW20 5,30 39,90 € 20,00 23% 169,18 €
OW20:GAU ECOLUB 1,00 0,42 € 23% 0,42 €
- Outros Débitos
016488 JUNTA BUJÃO 1,00 2,39€ 23% 2,39€
P&D PICK UP & DELIVERY 1,00 1,00€ 99,99 23%
Observações:
Total Mão de Obra:
Total Peças:
Outros Déb. + T. Ext:
Total Descontos
Total Líquido:
IVA: 23 %
Total Fatura:
19,50 €
227,82 €
3,39 €
46,83 €
203,38 €
46,89 €
250,77 €
CRUZ & ALLEN - B.B.C OFICINA DE AUTOMÓVEIS, LDA
CONTRIBUINTE Nº 504 104 250
"""

    payload = _batch_invoice_payload(b"", ".pdf", "FS-0002235.pdf", existing_text=text)

    assert payload["ocr_template"] == "cruz_allen_bbc_invoice"
    assert payload["km"] == "25380"
    assert payload["total_with_vat"] == "250,77"
    assert payload["labor_total"] == "19,50"
    assert payload["materials_total"] == "227,82"
    assert len(payload["invoice_lines"]) == 7
    assert payload["invoice_lines"][1]["quantity"] == "0,40"
    assert all("Total " not in line["description"] for line in payload["invoice_lines"])


def test_batch_invoice_payload_rebuilds_cruz_allen_stacked_two_page_table():
    text = """
FATURA
VIA NUMERO DATA
Original FS0002242 11/07/24
Cruz & Allen CARFAST, RENT-A-CAR, LDA
Cliente nº NIF Cliente Nº OR Página
000310 509285970 OR0006545 1/2
Modelo: C4 (C41) DVSRC/UE63
Matrícula: AX-30-EF
Kilómetros: 18261
Chassis: VR7BBYHZBNE063309
Referência Descrição
- Mao Obra
93830000 OPERAÇÕES SISTEMATICAS MANUTENÇÃO
49450907 SUBSTITUICAO FILTRO DE POLEN
- Peças
1680682480 FILTRO OLEO
9833351080 FILTRO CARVÃO
5W30 RCP OLEO QUARTZ INEO RCP 5W30
- Outros Débitos
016488 JUNTA BUJAO
1611908680 LIQUIDO L-VIDROS
- Mao Obra
CHAP REPARAR FRISO / CALHA VIDRO PORTA TRAS
DIREITA ( PARTIDO )
- Outros Débitos
ENSAIO SIMPLES - OFERTA
- Outros Débitos
LIMPEZA DA VIATURA - OFERTA
P&D PICK UP & DELIVERY
CRUZ & ALLEN - B.B.C OFICINA DE AUTOMÓVEIS, LDA
Qtd. Preço Unitário Dto. IVA Valor
1,00 39,00€ 10,00 23% 35,10 €
,10 39,00 € 10,00 23% 3,51€
1,00 15,93 € 10,00 23% 14,34 €
1,00 43,13 € 10,00 23% 38,82 €
4,00 45,84 € 20,00 23% 146,69 €
1,00 2,39€ 23% 2,39€
1,00 1,37 € 23% 1,37 €
1,20 39,00 € 23% 46,80 €
1,00 16,00 € 99,99
1,00 15,00 € 99,99
1,00 1,00 € 99,99
Continua...
FATURA
Original FS0002242 11/07/24
Referência Descrição Qtd. Preço Unitário Dto. IVA Valor
ALERTA MANUTENÇÃO DESTA VIATURA 30.000KMS / 1ANO - 1,00 1,00€ 99,99
Observações:
Total Fatura: 355,49 €
CONTRIBUINTE Nº 504 104 250
"""

    payload = _batch_invoice_payload(b"", ".pdf", "FS-0002242.pdf", existing_text=text)

    assert payload["ocr_template"] == "cruz_allen_bbc_invoice"
    assert payload["document_number"] == "FS0002242"
    assert payload["km"] == "18261"
    assert len(payload["invoice_lines"]) == 11
    assert payload["invoice_lines"][0]["reference"] == "93830000"
    assert payload["invoice_lines"][0]["amount"] == "35,10"
    assert payload["invoice_lines"][7]["description"].endswith("DIREITA ( PARTIDO )")
    assert payload["invoice_lines"][-1]["description"] == "PICK UP & DELIVERY"
    assert all("CAPITAL SOCIAL" not in line["description"] for line in payload["invoice_lines"])


def test_batch_invoice_payload_handles_cruz_allen_abbreviated_header_and_garbled_km_label():
    text = """
FATURA
Original FS0002485 27/11/24
Cruz & Allen CARFAST, RENT-A-CAR, LDA
Modelo: BERLINGO
Matrícula: BN-91-MN
pi sltâmetras: 35880 Data de Entrega: 27/11/2024
Referência Descrição Qtd. Preço Uni Dto. IVA Valor
- Mao Obra
01022815 MUDANCA DE OLEO E FILTRO MOTOR 40 39,00€ 10,00 23% 14,04 €
- Peças
1680682480 FILTRO OLEO 1,00 15,93€ 10,00 23% 14,34 €
- Outros Débitos
P&D PICK UP & DELIVERY 1,00 1,00€ 99,99 23%
Continua
CRUZ & ALLEN - B.B.C OFICINA DE AUTOMÓVEIS, LDA
Telefone 229 000 000
Capital Social 100000 EUR
Observações:
Total Fatura: 461,61 €
CONTRIBUINTE Nº 504 104 250
"""

    payload = _batch_invoice_payload(b"", ".pdf", "FS-0002485.pdf", existing_text=text)

    assert payload["ocr_template"] == "cruz_allen_bbc_invoice"
    assert payload["km"] == "35880"
    assert payload["total_with_vat"] == "461,61"
    assert len(payload["invoice_lines"]) == 3
    assert payload["invoice_lines"][0]["quantity"] == "0,40"
    assert payload["invoice_lines"][-1]["description"] == "P&D PICK UP & DELIVERY"
    assert all("Telefone" not in line["description"] for line in payload["invoice_lines"])
    assert all("Capital Social" not in line["description"] for line in payload["invoice_lines"])


def test_batch_invoice_payload_handles_cruz_allen_nd_service_invoice():
    text = """
FATURA
Cruz & Allen VIA NÚMERO DATA
Original ND 0000548 19-08-2025
CARFAST RENT-A-CAR, LDA
Descrição Tx. IVA Qtd. Preço Dto. Valor
Serviços Prestados
Valor E
23% 1,00 206,78 0,00 206,78 €
Observações
Total 206,78 €
Total Desconto 0,00 €
Total IVA 47,56 €
Total Documento 254,34 €
CRUZ & ALLEN - B.B.C OFICINA DE AUTOMÓVEIS, LDA
CONTRIBUINTE Nº 504 104 250
"""

    payload = _batch_invoice_payload(b"", ".pdf", "ND-0000548.pdf", existing_text=text)

    assert payload["ocr_template"] == "cruz_allen_bbc_invoice"
    assert payload["document_number"] == "ND0000548"
    assert payload["document_date"] == "2025-08-19"
    assert payload["subtotal_without_vat"] == "206,78"
    assert payload["vat_amount"] == "47,56"
    assert payload["total_with_vat"] == "254,34"
    assert len(payload["invoice_lines"]) == 1
    assert payload["invoice_lines"][0]["description"] == "Serviços Prestados Valor E"
    assert payload["invoice_lines"][0]["amount"] == "206,78"


def test_batch_invoice_payload_handles_cruz_allen_wrapped_service_amount():
    text = """
FATURA
Cruz & Allen VIA NUMERO DATA
Original ND 0000410 09-09-2024
CARFAST, RENT-A-CAR, LDA
Descrição Tx. IVA Qtd. Preço Dto. Valor
Serviços Prestados Referente a 23% 1,00 663,16 0,00
Comissões Processos
Seguradoras de junho a agosto
663,16 €
Total 663,16 €
Total IVA 152,53 €
Total Documento 815,69 €
CRUZ & ALLEN - B.B.C OFICINA DE AUTOMÓVEIS, LDA
CONTRIBUINTE Nº 504 104 250
"""

    payload = _batch_invoice_payload(b"", ".pdf", "ND-0000410.pdf", existing_text=text)

    assert payload["document_number"] == "ND0000410"
    assert payload["total_with_vat"] == "815,69"
    assert len(payload["invoice_lines"]) == 1
    assert payload["invoice_lines"][0]["amount"] == "663,16"
    assert "Comissões Processos Seguradoras" in payload["invoice_lines"][0]["description"]


def test_batch_invoice_payload_handles_cruz_allen_parts_receipt():
    text = """
FATURA/RECIBO
VIA NUMERO DATA
Original V10000341 23/05/24
Cruz & Allen CARFAST, RENT-A-CAR, LDA
Marca Referência Descrição Qtd Preço Unitário Dto. Iva% Valor
cl 9844895280 FAROLIM TRAS 1,00 203,75 € 10,00 23,00 183,37 €
Total Bruto: 203,75 €
Descontos 20,38 €
Subtotal 183,37 €
IVA 23 % 42,18€
Total Factura: 225,55 €
CRUZ & ALLEN - B.B.C OFICINA DE AUTOMÓVEIS, LDA
CONTRIBUINTE Nº 504 104 250
"""

    payload = _batch_invoice_payload(b"", ".pdf", "V1-0000341.pdf", existing_text=text)

    assert payload["document_number"] == "V10000341"
    assert payload["total_with_vat"] == "225,55"
    assert len(payload["invoice_lines"]) == 1
    assert payload["invoice_lines"][0]["description"].endswith("FAROLIM TRAS")
    assert payload["invoice_lines"][0]["tax"] == "23,00"
    assert payload["invoice_lines"][0]["amount"] == "183,37"


def test_batch_invoice_payload_extracts_cruz_allen_fs0004003():
    text = """
FATURA
VIA NÚMERO DATA
Original FS0004003 22/12/25
Exmo(s) Senhor(es):
CARFAST, RENT-A-CAR, LDA
R. DAS INDUSTRIAS, 220
4785-625-TROFA
Cliente nº NIF Cliente Nº OR Página
000310 509285970 OR0008902 1 / 1
Forma Paga.: PAGAM A 60 DIAS Recepcionista: CARLOS OLIVEIRA
Modelo: FABIA Chassis: TMBJG8NX0PY154577
Matrícula: BC-73-RJ Data de Garantia: 26/07/23
Kilómetros: 51366 Data de Entrega: 22/12/2025
Referência Descrição Qtd. Preço Unitário Dto. IVA Valor
- Mao Obra
151BRM07 ENCHIMENTO DEPOSITO DE LIQUIDO ADBLUE ,10 41,90 € 15,00 23% 3,56 €
- Peças
1618891880 UREIA 10,00 1,80 € 23% 18,00 €
Observações: Total Mão de Obra: 4,19 €
Total Peças:
18,00 €
Outros Déb. + T. Ext:
Total Descontos
0,63 €
Total Líquido:
21,56 €
I.V.A: 23 %
4,96 €
Total Fatura:
26,52 €
CRUZ & ALLEN - B.B.C OFICINA DE AUTOMÓVEIS, LDA
CONTRIBUINTE Nº 504 104 250
"""

    payload = _batch_invoice_payload(b"", ".pdf", "2025-12-22_BC-73-RJ_FORN_Fatura_VIA_DOC.pdf", existing_text=text)

    assert payload["ocr_template"] == "cruz_allen_bbc_invoice"
    assert payload["supplier_nif"] == "504104250"
    assert payload["document_number"] == "FS0004003"
    assert payload["document_date"] == "2025-12-22"
    assert payload["client_number"] == "000310"
    assert payload["client_nif"] == "509285970"
    assert payload["repair_order_reference"] == "OR0008902"
    assert payload["payment_method"] == "PAGAM A 60 DIAS"
    assert payload["receptionist"] == "CARLOS OLIVEIRA"
    assert payload["vehicle_model"] == "FABIA"
    assert payload["vin"] == "TMBJG8NX0PY154577"
    assert payload["plate"] == "BC-73-RJ"
    assert payload["km"] == "51366"
    assert payload["warranty_date"] == "2023-07-26"
    assert payload["delivery_date"] == "2025-12-22"
    assert payload["labor_total"] == "4,19"
    assert payload["materials_total"] == "18,00"
    assert payload["discount_without_vat"] == "0,63"
    assert payload["subtotal_without_vat"] == "21,56"
    assert payload["vat_amount"] == "4,96"
    assert payload["total_with_vat"] == "26,52"
    assert payload["invoice_lines"] == [
        {
            "line_type": "MO",
            "section": "Mão de obra",
            "reference": "151BRM07",
            "description": "ENCHIMENTO DEPOSITO DE LIQUIDO ADBLUE",
            "quantity": "0,10",
            "unit": "",
            "unit_price": "41,90",
            "discount_percent": "15,00",
            "tax": "23%",
            "amount": "3,56",
            "service": "Por classificar",
        },
        {
            "line_type": "MAT",
            "section": "Peças",
            "reference": "1618891880",
            "description": "UREIA",
            "quantity": "10,00",
            "unit": "",
            "unit_price": "1,80",
            "discount_percent": "",
            "tax": "23%",
            "amount": "18,00",
            "service": "Por classificar",
        },
    ]


def test_batch_invoice_payload_extracts_caetano_gamobar_3003():
    text = """
Ident ÚnicoDoc Valor Com IVA Data Vencim.
HFO/3003/2025
JF68PM28-3003
ATCUD
26-08-2025
27-10-2025
Data Vencim.
Carfast Rent-A-Car Lda
NIF Cliente
509285970
OR
HOJ/4949/2025
Núm. Autorização
765/2025
Forma de Pagamento
CREDITO
Matrícula / VIN
BQ96IP / VR3EDYHT9RJ968630
Modelo
PEUGEOT PARTNER ASPHALT LONGA 1.5 B
Kms.
9233,00
Data Abertura
25-08-2025
Valor Com IVA
Cliente queixa-se de luz de aviso de revisão e luz de motor acesas
Observações:
Viatura apresenta taxa elevada de carbono no óleo e elevada taxa de diluição - necessário substituir óleo e filtro de óleo e efetuar telecarregamento calculador do motor.
Carfast Rent-A-Car
Tipo
23,00
13,2192
40,00
73,4400
Horas
0,30
SUBSTITUICAO : FILTRO DE ÓLEO (NO VEICULO)
01E3YA
MO
23,00
17,6256
40,00
73,4400
Horas
0,40
MUDANÇA DE ÓLEO
15510A
MO
23,00
30,8448
40,00
73,4400
Horas
0,70
TELECARREGAMENTO DO CALCULADOR DO MOTOR
19009B
MO
23,00
93,6000
35,00
36,0000
UDS
4,00
ÓLEO CASTROL MAGNATEC 5W30 P
CST-1612B1
MAT
23,00
12,4575
25,00
16,6100
UDS
1,00
FILTRO OLEO
PSA-1680682480
MAT
23,00
2,1450
25,00
2,8600
UDS
1,00
JUNTA
PSA-1682801480
MAT
23,00
30,4200
35,00
36,0000
UDS
1,30
ÓLEO CASTROL MAGNATEC 5W30 P
CST-1612B1
MAT
23,00
2,4676
2,4676
UDS
1,00
Pequenos materiais
PM
VAR
23,00
0,2747
Ecolub 1L
23,00
0,1000
Ecolub 1L
Os serviços constantes
Imposto
46,7255
23,00
203,1544
249,88
46,7255
112,77
315,93
0,37
2,47
61,69
138,62
TOTAL
"""

    payload = _batch_invoice_payload(b"", ".pdf", "gamobar_3003.pdf", existing_text=text)

    assert payload["ocr_template"] == "caetano_gamobar_hfo"
    assert payload["supplier_name"] == "Caetano Gamobar Motores, S.A."
    assert payload["supplier_nif"] == "500112967"
    assert payload["document_number"] == "HFO/3003/2025"
    assert payload["document_date"] == "2025-08-26"
    assert payload["due_date"] == "2025-10-27"
    assert payload["atcud"] == "JF68PM28-3003"
    assert payload["client_name"] == "Carfast Rent-A-Car Lda"
    assert payload["client_nif"] == "509285970"
    assert payload["repair_order_reference"] == "HOJ/4949/2025"
    assert payload["authorization_number"] == "765/2025"
    assert payload["payment_method"] == "CREDITO"
    assert payload["plate"] == "BQ96IP"
    assert payload["vin"] == "VR3EDYHT9RJ968630"
    assert payload["vehicle_model"] == "PEUGEOT PARTNER ASPHALT LONGA 1.5 B"
    assert payload["km"] == "9233"
    assert payload["opening_date"] == "25-08-2025"
    assert "luz de aviso de revisão" in payload["complaint"]
    assert "taxa elevada de carbono" in payload["technical_observations"]
    assert payload["materials_total"] == "138,62"
    assert payload["labor_total"] == "61,69"
    assert payload["misc_total"] == "2,47"
    assert payload["taxes_total"] == "0,37"
    assert payload["gross_without_vat"] == "315,93"
    assert payload["discount_without_vat"] == "112,77"
    assert payload["taxable_base"] == "203,1544"
    assert payload["vat_amount"] == "46,7255"
    assert payload["total_with_vat"] == "249,88"
    assert payload["ecolub_total"] == "0,3747"
    assert payload["ocr_total_check"] == "ok"
    assert len(payload["invoice_lines"]) == 8
    assert payload["invoice_lines"][0]["line_type"] == "MO"
    assert payload["invoice_lines"][0]["reference"] == "01E3YA"
    assert payload["invoice_lines"][0]["quantity"] == "0,30"
    assert payload["invoice_lines"][0]["unit_price"] == "73,4400"
    assert payload["invoice_lines"][0]["discount_percent"] == "40,00"
    assert payload["invoice_lines"][0]["amount"] == "13,2192"
    assert payload["invoice_lines"][3]["reference"] == "CST-1612B1"
    assert payload["invoice_lines"][3]["amount"] == "93,6000"
    assert payload["invoice_lines"][-1]["line_type"] == "VAR"
    assert payload["invoice_lines"][-1]["amount"] == "2,4676"


def test_batch_invoice_payload_extracts_caetano_gamobar_xfo_696():
    text = """
FATURA
Ident. Único Doc: TAT XFO/696
2025-09-23
CRÉDITO
Forma de
Pagamento
13197
Kms.
PEUGEOT PARTNER ASPHALT LONGA 1.5 B
BS61FU / VR3EDYHTXRJ968622
Matrícula / VIN
2025-09-16
Data Abertura
XOJ/1986/2025
OR
509285970
NIF Cliente
2025-11-24
Data Vencim.
XFO/696/2025
Doc. Núm.
160,5000
Valor Com IVA
CLIENTE QUEIXA SE DA CHAVE DE SERVICO ANTES DO TEMPO
Tipo
23,00
18,8184
40,00
78,4100
Horas
0,40
MUDANCA DE OLEO E FIL
01E8DA
MO
23,00
1,9800
25,00
2,6400
Uds
1,00
LAMPADA 12V 16W S/CASQ
PSA-6216E0
MAT
23,00
13,5000
25,00
18,0000
Uds
1,00
FILTRO OLEO
PSA-1680682480
MAT
23,00
93,7137
35,00
36,5000
Uds
3,95
ÓLEO MOTOR TOTAL INEO RCP (5W30)
PSA-TPPRCPINEO
MAT
23,00
0,2793
Ecolub 1L
23,00
2,1975
25,00
2,9300
Uds
1,00
JUNTA
PSA-1682801480
MAT
Os serviços constantes deste documento foram concluidos em 2025-09-23
130,4889
IVA
0,2793
Ecolub 1L
160,50
30,01
68,90
199,39
0,28
0,00
18,82
111,39
Total
IVA
Desconto
Bruto
Taxas
Diversos
M. Obra
Materiais
ATCUD:JF6NPXD9-696
"""

    payload = _batch_invoice_payload(b"", ".pdf", "gamobar_696.pdf", existing_text=text)

    assert payload["ocr_template"] == "caetano_gamobar_xfo"
    assert payload["supplier_nif"] == "500112967"
    assert payload["document_number"] == "XFO/696/2025"
    assert payload["plate"] == "BS61FU"
    assert payload["vin"] == "VR3EDYHTXRJ968622"
    assert payload["km"] == "13197"
    assert payload["total_with_vat"] == "160,50"
    assert [line["reference"] for line in payload["invoice_lines"]] == [
        "01E8DA",
        "PSA-6216E0",
        "PSA-1680682480",
        "PSA-TPPRCPINEO",
        "PSA-1682801480",
    ]


def test_batch_invoice_payload_extracts_caetano_gamobar_bfo_632():
    text = """
CAETANO MOTORS
Caetano Gamobar Motors, S.A.
NIPC: 500 112 967
FATURA
Doc. Núm.
BFO/632/2025
Data Doc.
2025-07-31
Data Vencim.
2025-08-30
NIF Cliente
509285970
OR
BOJ/2214/2025
Data Abertura
2025-07-25
Matrícula / VIN
BH86EB / VR3EFYHT2PJ918611
Modelo
PEUGEOT PARTNER LONGA 1.5 BLUEHDI 1
Kms.
49872
Forma de Pagamento
CRÉDITO
Valor Com IVA
243,5900
Ident. Único Doc: TAT BFO/632
Tipo
23,00
47,0460
40,00
78,4100
Horas
1,00
- OPERAÇÕES SISTEMÁTICAS DE MANUTENÇÃO -
95N48A
MO
23,00
125,7425
35,00
36,5000
Uds
5,30
ÓLEO MOTOR TOTAL INEO RCP (5W30)
PSA-TPPRCPINEO
MAT
449,34
Total
"""

    payload = _batch_invoice_payload(b"", ".pdf", "gamobar_bfo_632.pdf", existing_text=text)

    assert payload["ocr_template"] == "caetano_gamobar_bfo"
    assert payload["supplier_name"] == "Caetano Gamobar Motores, S.A."
    assert payload["supplier_nif"] == "500112967"
    assert payload["client_nif"] == "509285970"
    assert payload["document_number"] == "BFO/632/2025"
    assert payload["document_date"] == "2025-07-31"
    assert payload["repair_order_reference"] == "BOJ/2214/2025"
    assert payload["plate"] == "BH86EB"
    assert payload["vin"] == "VR3EFYHT2PJ918611"
    assert payload["km"] == "49872"
def test_batch_invoice_payload_extracts_caetano_gamobar_2945():
    text = """
Ident ÚnicoDoc Valor Com IVA Data Vencim.
HFO/2945/2025
JF68PM28-2945
ATCUD
06-08-2025
05-09-2025
Data Vencim.
Carfast Rent-A-Car Lda
NIF Cliente
509285970
OR
HOJ/4590/2025
Núm. Autorização
Forma de Pagamento
CREDITO
Matrícula / VIN
BQ48IP / VR3EDYHT2RJ968629
Modelo
PEUGEOT PARTNER PRO STANDARD 1.5 B
Kms.
8899,00
Data Abertura
04-08-2025
Valor Com IVA
Cliente queixa-se que surgiu aviso "defeito no motor"
Tipo
23,00
17,6256
40,00
73,4400
Horas
0,40
MUDANÇA DE ÓLEO
95R48A
MO
23,00
12,4575
25,00
16,6100
UDS
1,00
FILTRO OLEO
PSA-1680682480
MAT
23,00
2,1450
25,00
2,8600
UDS
1,00
JUNTA
PSA-1682801480
MAT
23,00
124,0200
35,00
36,0000
UDS
5,30
ÓLEO CASTROL MAGNATEC 5W30 P
CST-1612B1
MAT
23,00
0,7050
0,7050
UDS
1,00
Pequenos materiais
PM
VAR
23,00
0,3747
Ecolub 1L
Os serviços constantes
Imposto
36,1854
23,00
157,3278
193,51
36,1854
83,40
240,73
0,37
0,71
17,63
138,62
TOTAL
"""

    payload = _batch_invoice_payload(b"", ".pdf", "gamobar_2945.pdf", existing_text=text)

    assert payload["ocr_template"] == "caetano_gamobar_hfo"
    assert payload["supplier_nif"] == "500112967"
    assert payload["document_number"] == "HFO/2945/2025"
    assert payload["document_date"] == "2025-08-06"
    assert payload["due_date"] == "2025-09-05"
    assert payload["atcud"] == "JF68PM28-2945"
    assert payload["repair_order_reference"] == "HOJ/4590/2025"
    assert payload.get("authorization_number", "") == ""
    assert payload["payment_method"] == "CREDITO"
    assert payload["plate"] == "BQ48IP"
    assert payload["vin"] == "VR3EDYHT2RJ968629"
    assert payload["vehicle_model"] == "PEUGEOT PARTNER PRO STANDARD 1.5 B"
    assert payload["km"] == "8899"
    assert payload["opening_date"] == "04-08-2025"
    assert payload["complaint"] == 'Cliente queixa-se que surgiu aviso "defeito no motor"'
    assert payload.get("technical_observations", "") == ""
    assert payload["materials_total"] == "138,62"
    assert payload["labor_total"] == "17,63"
    assert payload["misc_total"] == "0,71"
    assert payload["taxes_total"] == "0,37"
    assert payload["gross_without_vat"] == "240,73"
    assert payload["discount_without_vat"] == "83,40"
    assert payload["taxable_base"] == "157,3278"
    assert payload["vat_amount"] == "36,1854"
    assert payload["total_with_vat"] == "193,51"
    assert payload["ecolub_total"] == "0,3747"
    assert payload["ocr_total_check"] == "ok"
    assert len(payload["invoice_lines"]) == 5
    assert payload["invoice_lines"][0]["reference"] == "95R48A"
    assert payload["invoice_lines"][0]["amount"] == "17,6256"
    assert payload["invoice_lines"][3]["reference"] == "CST-1612B1"
    assert payload["invoice_lines"][3]["quantity"] == "5,30"
    assert payload["invoice_lines"][3]["amount"] == "124,0200"
    assert payload["invoice_lines"][-1]["line_type"] == "VAR"
    assert payload["invoice_lines"][-1]["amount"] == "0,7050"


def test_batch_invoice_payload_extracts_caetano_gamobar_hfj_invoice_with_discount():
    text = """
Original
FATURA
Ident. Único Doc: TATC HFJ/12488
2025-08-22
PRONTO PAGAMENTO
Forma de
Pagamento
84161
Kms.
Modelo
PEUGEOT BOXER 330 PREMIUM L2H2 2.2 B
AU86DJ / VF3YABPFB12V76230
Matrícula / VIN
2025-08-20
Data Abertura
HOJ/4894/2025
OR
509285970
NIF Cliente
Data Doc.
HFJ/12488/2025
Doc. Núm.
Observações:
Carfast Rent-A-Car Lda (546)
758/2025
Núm. Autorização
275,3400
Valor Com IVA
Substituir 2 pneus frente
IVA
Valor
Preço
Unid.
Qtd.
Designação
Referência
Tipo
23,00
164,0000
82,0000
Uds
2,00
215/70R15 C 109/107S TBBTIRES ADVENZZA
9PR
PN-TL22400055
MAT
23,00
2,2800
Ecovalor Passageiro/Turismo
23,00
0,7400
0,3700
Uds
2,00
EM:VÁLVULA
PSA-1609060380
MAT
23,00
28,4600
28,4600
Uds
1,00
Substituição 2 Pneus
R502
VAR
23,00
28,3700
28,3700
Uds
1,00
Alinhamento 1 Eixo
R521
VAR
Os serviços constantes deste documento foram concluidos em 2025-08-22
Valor de Imposto
Taxa
B.T.
Imposto
51,4855
23,00
223,8500
IVA
2,2800
Ecovalor
275,34
51,49
0,00
223,85
2,28
56,83
0,00
164,74
Total
IVA
Desconto
Bruto
Taxas
Diversos
M. Obra
Materiais
ATCUD:JF65PXPW-12488
"""

    payload = _batch_invoice_payload(b"", ".pdf", "12488.pdf", existing_text=text)

    assert payload["ocr_template"] == "caetano_gamobar_hfj"
    assert payload["supplier_nif"] == "500112967"
    assert payload["document_number"] == "HFJ/12488/2025"
    assert payload["document_date"] == "2025-08-22"
    assert payload["payment_method"] == "PRONTO PAGAMENTO"
    assert payload["repair_order_reference"] == "HOJ/4894/2025"
    assert payload["authorization_number"] == "758/2025"
    assert payload["plate"] == "AU86DJ"
    assert payload["vin"] == "VF3YABPFB12V76230"
    assert payload["km"] == "84161"
    assert payload["opening_date"] == "2025-08-20"
    assert payload["complaint"] == "Substituir 2 pneus frente"
    assert payload["total_with_vat"] == "275,34"
    assert payload["invoice_lines"][0]["reference"] == "PN-TL22400055"
    assert payload["invoice_lines"][0]["description"] == "215/70R15 C 109/107S TBBTIRES ADVENZZA 9PR"
    assert payload["invoice_lines"][0]["quantity"] == "2,00"
    assert payload["invoice_lines"][0]["unit_price"] == "82,0000"
    assert payload["invoice_lines"][0]["amount"] == "164,0000"
    assert len(payload["invoice_lines"]) == 4


def test_batch_invoice_payload_extracts_caetano_gamobar_hfj_invoice_with_percentage_discount():
    text = """
Original
FATURA
Ident. Único Doc: TATC HFJ/12942
2025-10-02
PRONTO PAGAMENTO
Forma de
Pagamento
68380
Kms.
Modelo
PEUGEOT 208 ACTIVE PACK 1.2 PURETEC
AU99XZ / VR3UPHNEKN5897342
Matrícula / VIN
2025-09-26
Data Abertura
HOJ/5569/2025
OR
509285970
NIF Cliente
Data Doc.
HFJ/12942/2025
Doc. Núm.
Observações:
Carfast Rent-A-Car Lda (546)
838/2025
Núm. Autorização
9,7600
Valor Com IVA
Substituir tapetes (alcatifa)
IVA
Valor
Dto. %
Preço
Unid.
Qtd.
Designação
Referência
Tipo
23,00
7,9380
30,00
11,3400
Uds
1,00
JG TAPETES 208 5 LG 2019 ->
PSA-LPTPJF0998
MAT
23,00
0,0000
0,0000
Uds
1,00
Tapetes
RC03
VAR
Os serviços constantes deste documento foram concluidos em 2025-10-02
Valor de Imposto
Taxa
B.T.
Imposto
1,8257
23,00
7,9380
IVA
9,76 €
1,83
3,40
11,34
0,00
0,00
7,94
Total
IVA
Desconto
Bruto
Diversos
M. Obra
Materiais
ATCUD:JF65PXPW-12942
"""

    payload = _batch_invoice_payload(b"", ".pdf", "12942.pdf", existing_text=text)

    assert payload["ocr_template"] == "caetano_gamobar_hfj"
    assert payload["document_number"] == "HFJ/12942/2025"
    assert payload["plate"] == "AU99XZ"
    assert payload["km"] == "68380"
    assert payload["total_with_vat"] == "9,76"
    assert payload["invoice_lines"][0]["reference"] == "PSA-LPTPJF0998"
    assert payload["invoice_lines"][0]["description"] == "JG TAPETES 208 5 LG 2019 ->"
    assert payload["invoice_lines"][0]["discount_percent"] == "30,00"
    assert payload["invoice_lines"][0]["amount"] == "7,9380"
    assert len(payload["invoice_lines"]) == 2


def test_batch_invoice_payload_extracts_caetano_gamobar_ffj_franchise_invoice():
    text = """
Original
FATURA
Ident. Único Doc: TATC FFJ/2602
2025-08-01
PRONTO PAGAMENTO
Forma de
Pagamento
15332
Kms.
Modelo
PEUGEOT 208 ALLURE HYBRID 100 E-DCS6
BI36EJ / VR3UPHPX5P4369170
Matrícula / VIN
2025-07-28
Data Abertura
FOJ/949/2025
OR
509285970
NIF Cliente
Data Doc.
FFJ/2602/2025
Doc. Núm.
Observações:
Carfast Rent-A-Car Lda (546)
PRVI/9917/2025
Núm. Orçamento
220052383
Núm. Sinistro
Núm. Autorização
769,9400
Valor Com IVA
Débito franquia
IVA
Valor
Preço
Unid.
Qtd.
Designação
Referência
Tipo
23,00
625,9675
625,9675
Uds
1,00
FRANQUIA
FR
VAR
Os serviços constantes deste documento foram concluidos em 2025-08-01
Valor de Imposto
Taxa
B.T.
Imposto
143,9725
23,00
625,9675
IVA
769,94 €
143,97
0,00
625,97
625,97
0,00
0,00
Total
IVA
Desconto
Bruto
Diversos
M. Obra
Materiais
ATCUD:JF6PPMRG-2602
"""

    payload = _batch_invoice_payload(b"", ".pdf", "2602.pdf", existing_text=text)

    assert payload["ocr_template"] == "caetano_gamobar_ffj"
    assert payload["supplier_nif"] == "500112967"
    assert payload["document_number"] == "FFJ/2602/2025"
    assert payload["repair_order_reference"] == "FOJ/949/2025"
    assert payload["authorization_number"] == ""
    assert payload["plate"] == "BI36EJ"
    assert payload["km"] == "15332"
    assert payload["total_with_vat"] == "769,94"
    assert payload["invoice_lines"][0]["reference"] == "FR"
    assert payload["invoice_lines"][0]["description"] == "FRANQUIA"
    assert payload["invoice_lines"][0]["amount"] == "625,9675"
    assert len(payload["invoice_lines"]) == 1


def test_batch_invoice_payload_extracts_caetano_gamobar_ffo_invoice_number():
    text = """
Original
FATURA
Ident. Único Doc: TAT FFO/525
2024-07-30
2024-09-28
Data Vencim.
Data Doc.
FFO/525/2024
Doc. Núm.
42127
Kms.
AS43JQ / VR3UDYHSKNJ684417
Matrícula /
VIN
Carfast Rent-A-Car Lda (546)
NIF Cliente
509285970
Valor Com IVA
REVISÕES : OPERAÇÕES SISTEMÁTICAS
Total Fatura
212,86
"""

    payload = _batch_invoice_payload(b"", ".pdf", "gamobar_ffo_525.pdf", existing_text=text)

    assert payload["ocr_template"] == "caetano_gamobar_ffo"
    assert payload["supplier_nif"] == "500112967"
    assert payload["document_number"] == "FFO/525/2024"
    assert payload["document_date"] == "2024-07-30"
    assert payload["plate"] == "AS43JQ"
    assert payload["km"] == "42127"


def test_batch_invoice_payload_extracts_caetano_gamobar_hfm_parts_header():
    text = """
2ª Via
Fatura
04/09/2025
2025-08-05
509285970
NIF
Carfast Rent-A-Car Lda(546)
HFM/352/2025
Data de vencimento
Nº Fatura
Data Doc.
Armazém Caet. Techn. Porto - Delfim Ferreira
05/08/2025 14:29:46
Data Emissão
CRÉDITO
Forma de Pagamento
Ident. Único Doc: VCRFC HFM/352
Qtd.
Designação
Referência
6,00
RECARGA KIT
9831814080
Total Fatura
140,01
ATCUD:JF6SPMRD-352
"""

    payload = _batch_invoice_payload(b"", ".pdf", "gamobar_hfm_352.pdf", existing_text=text)

    assert payload["ocr_template"] == "caetano_gamobar_hfm"
    assert payload["supplier_nif"] == "500112967"
    assert payload["document_number"] == "HFM/352/2025"
    assert payload["document_date"] == "2025-08-05"
    assert payload["due_date"] == "2025-09-04"


def test_batch_invoice_payload_extracts_caetano_gamobar_ffo_without_financial_rows():
    text = """
Original
FATURA
FFO/889/2025
04/09/2025
REPARAR CONFORME RELATORIO DE PERITAGEM 604062222
935,94
175,01
0,00
760,93
0,00
384,52
376,41
Total
IVA
Desconto
Bruto
Diversos
M. Obra
Materiais
"""

    payload = _batch_invoice_payload(b"", ".pdf", "gamobar_ffo_889.pdf", existing_text=text)

    assert payload["document_number"] == "FFO/889/2025"
    assert payload["total_with_vat"] == "935,94"
    assert payload["taxable_base"] == "760,9300"
    assert payload["vat_amount"] == "175,01"
    assert payload["ocr_total_check"] == "ok"
    assert len(payload["invoice_lines"]) == 1
    assert payload["invoice_lines"][0]["description"] == "REPARAR CONFORME RELATORIO DE PERITAGEM 604062222"
    assert payload["invoice_lines"][0]["amount"] == "760,9300"
    assert all(line["description"] not in {"Total", "IVA", "Bruto"} for line in payload["invoice_lines"])


def test_batch_invoice_payload_extracts_caetano_gamobar_ffo_discounted_total():
    text = """
FATURA
FFO/892/2025
REPARAR CONF RELATÓRIO PERITAGEM 220052383
2548,07
476,47
142,50
2214,10
0,00
1439,11
632,49
Total
IVA
Desconto
Bruto
Diversos
M. Obra
Materiais
"""

    payload = _batch_invoice_payload(b"", ".pdf", "gamobar_ffo_892.pdf", existing_text=text)

    assert payload["total_with_vat"] == "2548,07"
    assert payload["taxable_base"] == "2071,6000"
    assert payload["discount_without_vat"] == "142,50"
    assert len(payload["invoice_lines"]) == 1


def test_batch_invoice_payload_extracts_caetano_gamobar_hfm_single_parts_line():
    text = """
Fatura
04/09/2025
2025-08-05
HFM/352/2025
Data de vencimento
Nº Fatura
Data Doc.
MP
23,00
113,83
10,00
21,08
Uds
6,00
2C10-03
RECARGA KIT
9831814080
PSA
Total Fatura
Total IVA
Dto.
Total B.T.
Total Embalagens
Total Portes
140,01
26,18
12,65
113,83
0,00
0,00
"""

    payload = _batch_invoice_payload(b"", ".pdf", "gamobar_hfm_352.pdf", existing_text=text)

    assert payload["document_number"] == "HFM/352/2025"
    assert payload["total_with_vat"] == "140,01"
    assert payload["taxable_base"] == "113,83"
    assert payload["vat_amount"] == "26,18"
    assert payload["ocr_total_check"] == "ok"
    assert payload["invoice_lines"] == [
        {
            "reference": "9831814080",
            "description": "RECARGA KIT",
            "quantity": "6,00",
            "unit": "Uds",
            "unit_price": "21,08",
            "list_price": "21,08",
            "discount_percent": "10,00",
            "tax": "23,00",
            "amount": "113,83",
            "service": "Por classificar",
            "line_type": "MAT",
        }
    ]


def test_batch_invoice_payload_extracts_filinto_tal_credit_note():
    text = """
ORIGINAL
Filinto Mota Sucessores S.A. (Braga)
NIF: 500 324 174
NOTA DE CRÉDITO
DOCUMENTO : TAL_ABONOFAC 2023/21903953 Fatura Creditada Nr : 21095913
DE : 04/12/2023 O.R. : 307917 OFICINA : (071) PEU Reparação VENCIMENTO : 04/12/2023
MAT : AX-37-FI MARCA : PEUGEOT MODELO : Boxer
KMS : 25065 RECEPCIONISTA : Marcia Costa CHASSIS : VF3YBBPFC12W32778
Descrição Referência P.V.Unit Desc P.Liq.Unit Tmp/Qt Total Liq. CT
MUDAR ÓLEO E FILTRO D
- OPERAÇÕES SISTEMÁTICAS DE MANUTENÇÃO - 95N48A 59,00 25,00 44,25 -1 -44,25B
CÓDIGO/DESCRIÇÃO I.V.A. TAXA I.V.A. BASE INCIDÊNCIA VALOR I.V.A. TOTAL LÍQUIDO TOTAL I.V.A. TOTAL A PAGAR
B APV TX NORMAL 23,00 -61,75 -14,20 -61,75 -14,20 -75,95
"""

    payload = _batch_invoice_payload(b"", ".pdf", "tal_credit.pdf", existing_text=text)

    assert payload["document_kind"] == "credit_note"
    assert payload["document_number"] == "TAL_ABONOFAC 2023/21903953"
    assert payload["supplier_name"] == "Filinto Mota Sucessores S.A. (NIF 500324174)"
    assert payload["supplier_nif"] == "500324174"
    assert payload["plate"] == "AX-37-FI"
    assert payload["km"] == "25065"
    assert payload["total_with_vat"] == "-75,95"
    assert payload["work_order_reference"] == ""
    assert payload["repair_order_reference"] == "307917"
    assert payload["invoice_lines"][0]["service"] == "Manutenção"
    assert payload["invoice_lines"][0]["amount"] == "-44,25"


def test_clean_document_reprocess_invoice_ocr_reports_empty_text(authenticated_client, db_session, tmp_path):
    vehicle = _create_vehicle(db_session)
    invoice_path = tmp_path / "fatura_scanner.jpg"
    invoice_path.write_bytes(b"not an image with readable text")
    invoice = Document(
        title="Fatura scanner",
        document_type="workshop_supplier_invoice",
        classification="fleet",
        source="v2_clean_manual",
        entry_channel="v2_clean_batch",
        original_name=invoice_path.name,
        file_name=invoice_path.name,
        file_type="jpg",
        file_size=invoice_path.stat().st_size,
        storage_provider="local",
        storage_path=str(invoice_path),
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
        status="received",
        archived=True,
    )
    db_session.add(invoice)
    db_session.commit()
    db_session.refresh(invoice)

    response = authenticated_client.post(
        f"/v2-clean/documents/{invoice.id}/reprocess-ocr",
        data={"return_url": f"/v2-clean/fleet/{vehicle.id}/documents?main_group=invoices"},
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert "ocr_error=no_text" in response.headers["location"]
    assert "ocr_lines=0" in response.headers["location"]
    db_session.refresh(invoice)
    assert invoice.status == "ocr_empty"


def test_clean_vehicle_documents_page_renders(authenticated_client, db_session):
    vehicle = _create_vehicle(db_session)

    response = authenticated_client.get(f"/v2-clean/fleet/{vehicle.id}/documents")

    assert response.status_code == 200
    assert ">Faturas<" in response.text
    assert "Documentação estruturada" in response.text
    assert "Timeline documental" in response.text


def test_clean_vehicle_documents_page_renders_with_regressive_km_alert(authenticated_client, db_session):
    vehicle = _create_vehicle(db_session)
    db_session.add_all(
        [
            VehicleDocumentRecord(
                vehicle_id=vehicle.id,
                source_record_type="structured",
                main_group="work_orders",
                title="FO 1",
                plate=vehicle.plate,
                document_date=date(2026, 1, 1),
                km=1000,
            ),
            VehicleDocumentRecord(
                vehicle_id=vehicle.id,
                source_record_type="structured",
                main_group="work_orders",
                title="FO 2",
                plate=vehicle.plate,
                document_date=date(2026, 1, 2),
                km=900,
            ),
        ]
    )
    db_session.commit()

    response = authenticated_client.get(f"/v2-clean/fleet/{vehicle.id}/documents")

    assert response.status_code == 200
    assert "KM regressivo" in response.text


def test_clean_vehicle_summary_hides_legacy_documents(authenticated_client, db_session):
    vehicle = _create_vehicle(db_session)
    legacy_document = Document(
        title="Relatório antigo",
        document_type="workshop_report",
        classification="technical_report",
        source="workshop",
        entry_channel="legacy",
        original_name="legacy.pdf",
        file_name="legacy.pdf",
        storage_provider="local",
        storage_path="/tmp/legacy.pdf",
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
    )
    clean_document = Document(
        title="Relatório v2",
        document_type="workshop_report",
        classification="technical_report",
        source="v2_clean_manual",
        entry_channel="v2_clean",
        original_name="clean.pdf",
        file_name="clean.pdf",
        storage_provider="local",
        storage_path="/tmp/clean.pdf",
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
    )
    db_session.add_all([legacy_document, clean_document])
    db_session.commit()

    response = authenticated_client.get(f"/v2-clean/fleet/{vehicle.id}")

    assert response.status_code == 200
    assert "Relatório v2" in response.text
    assert "Relatório antigo" not in response.text


def test_clean_vehicle_documents_audit_field_syncs_real_start(authenticated_client, db_session):
    vehicle = _create_vehicle(db_session)

    response = authenticated_client.post(
        f"/v2-clean/fleet/{vehicle.id}/documents/audit-field",
        data={
            "field_code": "real_start_date",
            "value": "2024-05-20",
            "audited_on": "2026-07-13",
            "observation": "Validado por documento base",
            "document_basis": "DUA + compra",
        },
        follow_redirects=False,
    )

    assert response.status_code == 303
    audit_field = db_session.scalar(
        select(VehicleDocumentAuditField).where(
            VehicleDocumentAuditField.vehicle_id == vehicle.id,
            VehicleDocumentAuditField.field_code == "real_start_date",
        )
    )
    manual_field = db_session.scalar(
        select(VehicleManualField).where(
            VehicleManualField.vehicle_id == vehicle.id,
            VehicleManualField.field_code == "real_start_date",
        )
    )
    assert audit_field is not None
    assert audit_field.value_json == "2024-05-20"
    assert manual_field is not None
    assert manual_field.value_json == "2024-05-20"


def test_clean_vehicle_documents_import_work_orders(authenticated_client, db_session):
    vehicle = _create_vehicle(db_session)
    workbook = _make_workbook(
        ["Número", "Data", "Matrícula", "Nome fornecedor", "Observações"],
        [["FO-1576", "2026-05-15", "CC-11-AA", "Oficina Porto", "Revisão e pneus dianteiros"]],
    )

    response = authenticated_client.post(
        f"/v2-clean/fleet/{vehicle.id}/documents/import/work-orders",
        files={"file": ("fo.xlsx", workbook.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        follow_redirects=False,
    )

    assert response.status_code == 303
    record = db_session.scalar(
        select(VehicleDocumentRecord).where(
            VehicleDocumentRecord.vehicle_id == vehicle.id,
            VehicleDocumentRecord.main_group == "work_orders",
        )
    )
    assert record is not None
    assert record.title == "FO-1576"
    assert record.supplier_name == "Oficina Porto"
    import_source = db_session.scalar(
        select(Document).where(
            Document.vehicle_id == vehicle.id,
            Document.entry_channel == "structured_import",
            Document.source_subject == "work_orders:1",
        )
    )
    assert import_source is not None
    assert import_source.file_hash
    assert import_source.storage_path.endswith(".xlsx")
    db_session.refresh(record)
    assert record.document_id == import_source.id
    module_ctx = vehicle_document_module_context(db_session, vehicle)
    assert len(module_ctx["structured_rows"]) == 1
    assert len(module_ctx["import_rows"]) == 1
    assert module_ctx["import_rows"][0]["import_label"] == "Folhas de obra"
    assert module_ctx["archive_rows"] == []

    page = authenticated_client.get(f"/v2-clean/fleet/{vehicle.id}/documents")
    assert page.status_code == 200
    assert "Folhas de obra" in page.text
    assert "FO-1576" in page.text
    assert "Documentação estruturada" in page.text
    assert "Fontes importadas" not in page.text
    assert "Manutenção" in page.text
    assert "Calços" in page.text
    assert "Discos" in page.text
    assert f"/v2-clean/fleet/{vehicle.id}/documents/classify-row" in page.text

    fleet_page = authenticated_client.get(f"/v2-clean/fleet/{vehicle.id}")
    assert fleet_page.status_code == 200
    assert "Folhas de obra" in fleet_page.text
    assert f"/v2-clean/fleet/{vehicle.id}/documents?main_group=work_orders" in fleet_page.text
    assert "<strong>1</strong>" in fleet_page.text


def test_clean_vehicle_documents_save_row_classification(authenticated_client, db_session):
    vehicle = _create_vehicle(db_session)
    record = VehicleDocumentRecord(
        vehicle_id=vehicle.id,
        source_record_type="structured",
        main_group="work_orders",
        title="FO-1608",
        external_reference="1608",
        plate=vehicle.plate,
        supplier_name="Oficina Porto",
        raw_description="Calços atrás gastos",
        document_date=date(2026, 5, 25),
        status="structured",
    )
    db_session.add(record)
    db_session.commit()
    db_session.refresh(record)

    response = authenticated_client.post(
        f"/v2-clean/fleet/{vehicle.id}/documents/classify-row",
        data={
            "record_id": str(record.id),
            "return_group": "work_orders",
            "maintenance": "",
            "pads": "rear",
            "discs": "",
            "tyres": "",
            "ipo": "",
            "other": "",
            "manual_note": "Confirmado em oficina",
        },
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert response.headers["location"].endswith("?classified=1&main_group=work_orders")
    db_session.refresh(record)
    assert record.status == "pending_validation"
    assert record.comparison_state == "por_validar"
    assert record.metadata_json["manual_note"] == "Confirmado em oficina"
    tag = db_session.scalar(
        select(VehicleDocumentRecordTag).where(
            VehicleDocumentRecordTag.vehicle_id == vehicle.id,
            VehicleDocumentRecordTag.record_id == record.id,
            VehicleDocumentRecordTag.category == "pads",
        )
    )
    assert tag is not None
    assert tag.value == "rear"

    page = authenticated_client.get(f"/v2-clean/fleet/{vehicle.id}/documents")
    assert page.status_code == 200
    assert '<option value="rear" selected>TR</option>' in page.text
    assert "Por validar" in page.text
    assert "Confirmado em oficina" in page.text
    assert "Guardar classificação" in page.text


def test_clean_vehicle_documents_service_choices_match_operational_vocabulary(
    authenticated_client,
    db_session,
):
    vehicle = _create_vehicle(db_session)
    record = VehicleDocumentRecord(
        vehicle_id=vehicle.id,
        source_record_type="structured",
        main_group="work_orders",
        title="FO-1701",
        external_reference="1701",
        plate=vehicle.plate,
        raw_description="Telecarregamento e filtro habitáculo",
        status="structured",
    )
    db_session.add(record)
    db_session.commit()

    maintenance = dict(DOCUMENT_HISTORY_QUICK_CLASSIFICATIONS["maintenance"])
    other = dict(DOCUMENT_HISTORY_QUICK_CLASSIFICATIONS["other"])
    assert "undefined" not in maintenance
    assert maintenance["telecharge"] == "Telecarregamento"
    assert maintenance["cabin_filter"] == "Filtro de habitáculo"
    assert other == {"claim": "Sinistro", "glass": "Vidros"}

    page = authenticated_client.get(
        f"/v2-clean/fleet/{vehicle.id}/documents?main_group=work_orders"
    )
    assert page.status_code == 200
    assert "Telecarregamento" in page.text
    assert "Filtro de habitáculo" in page.text
    assert "Sinistro" in page.text
    assert "Vidros" in page.text

    context = vehicle_document_module_context(db_session, vehicle)
    assert "telecharge" in context["structured_rows"][0]["service_suggestion_codes"]["maintenance"]
    assert "cabin_filter" in context["structured_rows"][0]["service_suggestion_codes"]["maintenance"]


def test_preclassify_invoices_and_work_orders_persists_suggestions_without_validation(
    db_session,
):
    vehicle = _create_vehicle(db_session)
    invoice = Document(
        title="Fatura sem classificação",
        document_type="workshop_supplier_invoice",
        classification="invoice",
        source="v2_clean_manual",
        original_name="invoice.pdf",
        file_name="invoice.pdf",
        storage_provider="local",
        storage_path="invoice.pdf",
        status="extracted",
        vehicle_id=vehicle.id,
    )
    work_order = VehicleDocumentRecord(
        vehicle_id=vehicle.id,
        source_record_type="structured",
        main_group="work_orders",
        title="FO-1800",
        raw_description="Substituição de pneus por furo",
        status="structured",
    )
    db_session.add_all([invoice, work_order])
    db_session.flush()
    db_session.add(
        DocumentEvent(
            document_id=invoice.id,
            action="invoice.ocr.extracted",
            old_value=None,
            new_value=json.dumps(
                {
                    "invoice_lines": [
                        {"description": "Mudança de óleo e filtro do habitáculo"},
                        {"description": "Telecarregamento"},
                    ]
                }
            ),
        )
    )
    db_session.commit()

    result = preclassify_invoices_and_work_orders(db_session)
    db_session.commit()

    db_session.refresh(invoice)
    db_session.refresh(work_order)
    assert invoice.status == "pending_validation"
    assert work_order.status == "pending_validation"
    assert work_order.comparison_state == "por_validar"
    assert result["invoice_documents"] == 1
    assert result["work_orders"] == 1
    invoice_tags = db_session.scalars(
        select(VehicleDocumentRecordTag).where(
            VehicleDocumentRecordTag.document_id == invoice.id
        )
    ).all()
    assert {(tag.category, tag.value) for tag in invoice_tags} == {
        ("maintenance", "telecharge"),
        ("maintenance", "cabin_filter"),
    }
    assert {tag.source_kind for tag in invoice_tags} == {"auto_suggested"}
    work_order_tags = db_session.scalars(
        select(VehicleDocumentRecordTag).where(
            VehicleDocumentRecordTag.record_id == work_order.id
        )
    ).all()
    assert {(tag.category, tag.value) for tag in work_order_tags} == {
        ("tyres", "puncture")
    }


def test_preclassify_preserves_validated_and_existing_human_classifications(db_session):
    vehicle = _create_vehicle(db_session)
    validated_invoice = Document(
        title="Revisão",
        document_type="workshop_supplier_invoice",
        classification="invoice",
        source="v2_clean_manual",
        original_name="validated.pdf",
        file_name="validated.pdf",
        storage_provider="local",
        storage_path="validated.pdf",
        status="classified",
        vehicle_id=vehicle.id,
    )
    work_order = VehicleDocumentRecord(
        vehicle_id=vehicle.id,
        source_record_type="structured",
        main_group="work_orders",
        title="Calços e reparação de furo",
        status="structured",
    )
    db_session.add_all([validated_invoice, work_order])
    db_session.flush()
    db_session.add(
        VehicleDocumentRecordTag(
            vehicle_id=vehicle.id,
            record_id=work_order.id,
            category="pads",
            value="rear",
            source_kind="manual",
        )
    )
    db_session.commit()

    result = preclassify_invoices_and_work_orders(db_session)
    db_session.commit()

    db_session.refresh(validated_invoice)
    assert validated_invoice.status == "classified"
    assert result["already_validated"] == 1
    tags = db_session.scalars(
        select(VehicleDocumentRecordTag).where(
            VehicleDocumentRecordTag.record_id == work_order.id
        )
    ).all()
    assert ("pads", "rear", "manual") in {
        (tag.category, tag.value, tag.source_kind) for tag in tags
    }
    assert ("tyres", "puncture", "auto_suggested") in {
        (tag.category, tag.value, tag.source_kind) for tag in tags
    }


def test_clean_vehicle_documents_can_explicitly_validate_row_classification(
    authenticated_client,
    db_session,
):
    vehicle = _create_vehicle(db_session)
    record = VehicleDocumentRecord(
        vehicle_id=vehicle.id,
        source_record_type="structured",
        main_group="work_orders",
        title="FO-1702",
        external_reference="1702",
        plate=vehicle.plate,
        status="structured",
    )
    db_session.add(record)
    db_session.commit()
    db_session.refresh(record)

    response = authenticated_client.post(
        f"/v2-clean/fleet/{vehicle.id}/documents/classify-row",
        data={
            "record_id": str(record.id),
            "return_group": "work_orders",
            "maintenance": "revision",
            "classification_action": "validate",
        },
        follow_redirects=False,
    )

    assert response.status_code == 303
    db_session.refresh(record)
    assert record.status == "classified"
    assert record.comparison_state == "validado"


def test_clean_vehicle_documents_saves_multiple_services_and_custom_values(authenticated_client, db_session):
    vehicle = _create_vehicle(db_session)
    record = VehicleDocumentRecord(
        vehicle_id=vehicle.id,
        source_record_type="structured",
        main_group="work_orders",
        title="FO-1700",
        external_reference="1700",
        plate=vehicle.plate,
        status="structured",
    )
    db_session.add(record)
    db_session.commit()
    db_session.refresh(record)

    response = authenticated_client.post(
        f"/v2-clean/fleet/{vehicle.id}/documents/classify-row",
        data={
            "record_id": str(record.id),
            "return_group": "work_orders",
            "maintenance": ["revision", "degradation"],
            "pads": ["front", "rear"],
            "discs": ["undefined"],
            "other_custom": "Bateria; Correia auxiliar",
        },
        follow_redirects=False,
    )

    assert response.status_code == 303
    db_session.refresh(record)
    assert record.status == "pending_validation"
    assert record.comparison_state == "por_validar"
    tags = db_session.scalars(
        select(VehicleDocumentRecordTag).where(VehicleDocumentRecordTag.record_id == record.id)
    ).all()
    assert {(tag.category, tag.value) for tag in tags if tag.value} >= {
        ("maintenance", "revision"),
        ("maintenance", "degradation"),
        ("pads", "front"),
        ("pads", "rear"),
        ("discs", "undefined"),
    }
    assert {tag.free_text for tag in tags if tag.free_text} == {"Bateria", "Correia auxiliar"}


def test_clean_vehicle_documents_shows_existing_invoice_ocr_lines(authenticated_client, db_session):
    vehicle = _create_vehicle(db_session)
    invoice = Document(
        title="Fatura 4458",
        document_type="workshop_supplier_invoice",
        source="v2_clean_manual",
        entry_channel="v2_clean_manual",
        original_name="fatura-4458.pdf",
        file_name="fatura-4458.pdf",
        storage_provider="local",
        storage_path="Frota/fatura-4458.pdf",
        folder_path="Frota/Faturas",
        status="received",
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
        supplier_name="Oficina Porto",
        contract_number="4458",
        document_date=date(2026, 6, 1),
    )
    db_session.add(invoice)
    db_session.flush()
    db_session.add(
        DocumentEvent(
            document_id=invoice.id,
            action="invoice.ocr.extracted",
            new_value=json.dumps(
                {
                    "invoice_lines": [
                        {"description": "Óleo motor 5W30", "quantity": 5, "amount": "45,00", "service": "Revisão"},
                        {"description": "Filtro de óleo", "quantity": 1, "amount": "12,00"},
                    ]
                },
                ensure_ascii=False,
            ),
        )
    )
    db_session.commit()

    response = authenticated_client.get(
        f"/v2-clean/fleet/{vehicle.id}/documents?main_group=invoices"
    )

    assert response.status_code == 200
    assert "Óleo motor 5W30" in response.text
    assert "Filtro de óleo" in response.text
    assert "Associar FO" in response.text


def test_clean_vehicle_documents_treats_cruz_allen_fs_reference_as_invoice(
    authenticated_client,
    db_session,
):
    vehicle = _create_vehicle(db_session)
    work_order = VehicleDocumentRecord(
        vehicle_id=vehicle.id,
        source_record_type="structured",
        main_group="work_orders",
        title="FO 2877",
        external_reference="2877",
        plate=vehicle.plate,
        status="structured",
    )
    invoice = Document(
        title="FS0002877",
        document_type="archive_document",
        source="v2_clean_manual",
        entry_channel="batch_import",
        original_name="FS0002877.pdf",
        file_name="FS0002877.pdf",
        storage_provider="local",
        storage_path="Frota/FS0002877.pdf",
        status="received",
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
        supplier_name="Cruz & Allen - B.B.C Oficina de Automóveis, Lda",
        contract_number="FS0002877",
    )
    db_session.add_all([work_order, invoice])
    db_session.commit()
    db_session.refresh(work_order)
    db_session.refresh(invoice)

    context = vehicle_document_module_context(db_session, vehicle)
    row = next(item for item in context["archive_rows"] if item["id"] == invoice.id)
    assert row["archive_group"] == "invoices"
    assert row["main_group"] == "invoices"

    page = authenticated_client.get(
        f"/v2-clean/fleet/{vehicle.id}/documents?main_group=invoices"
    )
    assert page.status_code == 200
    assert "FS0002877" in page.text
    assert "Serviços relevantes da fatura" in page.text
    assert "Associar FO" in page.text
    assert f"/v2-clean/documents/{invoice.id}/remove" in page.text

    linked = authenticated_client.post(
        f"/v2-clean/fleet/{vehicle.id}/documents/link-work-order",
        data={
            "document_id": str(invoice.id),
            "work_order_record_id": str(work_order.id),
            "return_group": "invoices",
            "open_item": f"document:{invoice.id}",
        },
        follow_redirects=False,
    )
    assert linked.status_code == 303
    assert f"open_item=document%3A{invoice.id}" in linked.headers["location"]
    assert db_session.scalar(
        select(DocumentLink).where(
            DocumentLink.document_id == invoice.id,
            DocumentLink.entity_type == "vehicle_document_record",
            DocumentLink.entity_id == str(work_order.id),
            DocumentLink.category == "invoice_work_order",
        )
    )


def test_clean_vehicle_documents_links_invoice_to_work_order_once(authenticated_client, db_session):
    vehicle = _create_vehicle(db_session)
    work_order = VehicleDocumentRecord(
        vehicle_id=vehicle.id,
        source_record_type="structured",
        main_group="work_orders",
        title="FO 1608",
        external_reference="1608",
        plate=vehicle.plate,
        document_date=date(2026, 5, 25),
        status="classified",
    )
    invoice = Document(
        title="Fatura 4458",
        document_type="workshop_supplier_invoice",
        source="v2_clean_manual",
        entry_channel="v2_clean_manual",
        original_name="fatura-4458.pdf",
        file_name="fatura-4458.pdf",
        storage_provider="local",
        storage_path="Frota/fatura-4458.pdf",
        folder_path="Frota/Faturas",
        status="received",
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
    )
    db_session.add_all([work_order, invoice])
    db_session.commit()
    db_session.refresh(work_order)
    db_session.refresh(invoice)

    for _ in range(2):
        response = authenticated_client.post(
            f"/v2-clean/fleet/{vehicle.id}/documents/link-work-order",
            data={
                "document_id": str(invoice.id),
                "work_order_record_id": str(work_order.id),
                "return_group": "invoices",
            },
            follow_redirects=False,
        )
        assert response.status_code == 303

    links = db_session.scalars(
        select(DocumentLink).where(
            DocumentLink.document_id == invoice.id,
            DocumentLink.entity_type == "vehicle_document_record",
            DocumentLink.entity_id == str(work_order.id),
            DocumentLink.category == "invoice_work_order",
        )
    ).all()
    assert len(links) == 1
    events = db_session.scalars(
        select(DocumentEvent).where(
            DocumentEvent.document_id == invoice.id,
            DocumentEvent.action == "invoice.work_order_linked",
        )
    ).all()
    assert len(events) == 1
    page = authenticated_client.get(f"/v2-clean/fleet/{vehicle.id}/documents?main_group=invoices")
    assert page.status_code == 200
    assert "FO 1608" in page.text


def test_clean_vehicle_documents_keeps_detail_open_after_classification(authenticated_client, db_session):
    vehicle = _create_vehicle(db_session)
    document = Document(
        title="Service Box 396",
        document_type="servicebox_tsb",
        source="v2_clean_manual",
        entry_channel="v2_clean_manual",
        original_name="service-box-396.pdf",
        file_name="service-box-396.pdf",
        storage_provider="local",
        storage_path="Frota/service-box-396.pdf",
        folder_path="Frota/Service Box",
        status="received",
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
        supplier_name="service box",
        contract_number="396",
        document_date=date(2026, 7, 21),
    )
    db_session.add(document)
    db_session.commit()
    db_session.refresh(document)

    response = authenticated_client.post(
        f"/v2-clean/fleet/{vehicle.id}/documents/classify-row",
        data={
            "document_id": str(document.id),
            "return_group": "servicebox_tsb",
            "open_item": f"document:{document.id}",
            "maintenance": "revision",
        },
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert f"open_item=document%3A{document.id}" in response.headers["location"]

    page = authenticated_client.get(response.headers["location"])
    assert page.status_code == 200
    assert ">Faturas<" in page.text
    assert "Outros documentos" in page.text
    assert "Service Box / TSB" in page.text
    assert ">396<" in page.text
    assert f'id="archive-other-detail-SB-{document.id}"' in page.text
    assert "clean-doc-detail-open" in page.text


def test_clean_vehicle_documents_imports_multiple_detail_lines_by_work_order_number(
    authenticated_client,
    db_session,
):
    vehicle = _create_vehicle(db_session)
    record = VehicleDocumentRecord(
        vehicle_id=vehicle.id,
        source_record_type="structured",
        main_group="work_orders",
        title="1608",
        external_reference="1608",
        plate=vehicle.plate,
        status="structured",
    )
    db_session.add(record)
    db_session.commit()
    db_session.refresh(record)
    workbook = _make_workbook(
        ["Folha de obra nº", "Descrição", "Referência", "Quantidade", "Preço unitário", "Total", "Kms"],
        [
            ["1608", "Jogo de calços traseiros", "CAL-01", 1, 75, 75, 32100],
            ["1608", "Mão de obra", "MO-01", 1.5, 40, 60, 32100],
        ],
    )

    response = authenticated_client.post(
        "/v2-clean/documents/import/work-order-details",
        files={
            "file": (
                "detalhe fo.xlsx",
                workbook.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert "imported_count=2" in response.headers["location"]
    db_session.refresh(record)
    assert record.km == 32100
    assert len(record.metadata_json["work_order_lines"]) == 2
    assert record.metadata_json["work_order_lines"][0]["description"] == "Jogo de calços traseiros"
    module_ctx = vehicle_document_module_context(db_session, vehicle)
    row = next(item for item in module_ctx["structured_rows"] if item["id"] == record.id)
    assert len(row["work_order_lines"]) == 2


def test_clean_vehicle_documents_import_work_orders_deduplicates_by_number(authenticated_client, db_session):
    vehicle = _create_vehicle(db_session)
    workbook = _make_workbook(
        ["Número", "Data", "Matrícula", "Nome fornecedor", "Observações"],
        [
            ["1608", "2026-05-25", "CC-11-AA", "Oficina Porto", "Calços atrás gastos"],
            ["1608", "2026-05-25", "CC-11-AA", "Oficina Porto", "Calços atrás gastos repetido"],
        ],
    )

    response = authenticated_client.post(
        f"/v2-clean/fleet/{vehicle.id}/documents/import/work-orders",
        files={"file": ("fo.xlsx", workbook.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert "imported_count=1" in response.headers["location"]
    records = db_session.scalars(
        select(VehicleDocumentRecord).where(
            VehicleDocumentRecord.main_group == "work_orders",
            VehicleDocumentRecord.external_reference == "1608",
        )
    ).all()
    assert len(records) == 1
    assert records[0].vehicle_id == vehicle.id


def test_clean_document_detail_page_renders_in_v2(authenticated_client, db_session):
    vehicle = _create_vehicle(db_session)
    document = Document(
        title="Fatura oficina",
        document_type="workshop_supplier_invoice",
        classification="invoice",
        source="v2_clean_manual",
        entry_channel="v2_clean",
        original_name="fatura.pdf",
        file_name="fatura.pdf",
        file_type="pdf",
        file_size=2048,
        storage_provider="local",
        storage_path="/tmp/fatura.pdf",
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
        status="archived",
    )
    db_session.add(document)
    db_session.commit()

    response = authenticated_client.get(f"/v2-clean/documents/{document.id}")

    assert response.status_code == 200
    assert "Fatura oficina" in response.text
    assert vehicle.plate in response.text
    assert "Voltar à documentação" in response.text


def test_clean_vehicle_documents_import_work_orders_stays_on_current_vehicle(authenticated_client, db_session):
    current_vehicle = _create_vehicle(db_session)
    target_vehicle = Vehicle(
        plate="BB-69-TE",
        vin="VINBB69TE123456789",
        brand="CITROEN",
        model="BERLINGO",
        version="XL",
        rentway_unit_nr="244",
        lifecycle_status="active",
        operational_status="free",
        active=True,
    )
    db_session.add(target_vehicle)
    db_session.commit()
    db_session.refresh(target_vehicle)
    workbook = _make_workbook(
        ["Número", "Data", "Matrícula", "Nome fornecedor", "Observações"],
        [["1682", "2026-06-22", "BB-69-TE", "CARFAST RENT-A-CAR LDA (OFICINA)", "IPO"]],
    )

    response = authenticated_client.post(
        f"/v2-clean/fleet/{current_vehicle.id}/documents/import/work-orders",
        files={"file": ("fo.xlsx", workbook.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert response.headers["location"].startswith(f"/v2-clean/fleet/{current_vehicle.id}/documents")
    assert "imported_count=1" in response.headers["location"]
    target_record = db_session.scalar(
        select(VehicleDocumentRecord).where(
            VehicleDocumentRecord.vehicle_id == target_vehicle.id,
            VehicleDocumentRecord.main_group == "work_orders",
        )
    )
    current_record = db_session.scalar(
        select(VehicleDocumentRecord).where(
            VehicleDocumentRecord.vehicle_id == current_vehicle.id,
            VehicleDocumentRecord.main_group == "work_orders",
        )
    )
    assert target_record is None
    assert current_record is not None
    assert current_record.title == "1682"
    assert current_record.plate == current_vehicle.plate


def test_clean_document_global_import_resolves_vehicle_identifier(authenticated_client, db_session):
    vehicle = Vehicle(
        plate=None,
        vin="VINBB69TE123456789",
        brand="CITROEN",
        model="BERLINGO",
        version="XL",
        rentway_unit_nr="244",
        lifecycle_status="active",
        operational_status="free",
        active=True,
    )
    db_session.add(vehicle)
    db_session.commit()
    db_session.refresh(vehicle)
    db_session.add(
        VehicleIdentifier(
            vehicle_id=vehicle.id,
            identifier_type="plate",
            identifier_value="BB-69-TE",
            source_system="test",
            active=True,
        )
    )
    db_session.commit()
    workbook = _make_workbook(
        ["Número", "Data", "Matrícula", "Nome fornecedor", "Observações"],
        [["1682", "2026-06-22", "BB-69-TE", "CARFAST RENT-A-CAR LDA (OFICINA)", "IPO"]],
    )

    response = authenticated_client.post(
        "/v2-clean/documents/import/work-orders",
        files={"file": ("fo.xlsx", workbook.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert "imported_count=1" in response.headers["location"]
    record = db_session.scalar(
        select(VehicleDocumentRecord).where(
            VehicleDocumentRecord.vehicle_id == vehicle.id,
            VehicleDocumentRecord.main_group == "work_orders",
        )
    )
    assert record is not None
    assert record.title == "1682"

    page = authenticated_client.get(f"/v2-clean/fleet/{vehicle.id}/documents")
    assert page.status_code == 200
    assert "1682" in page.text
    assert "Folhas de obra" in page.text


def test_clean_document_center_shows_global_structured_import_rows(authenticated_client, db_session):
    vehicle = Vehicle(
        plate="BB-69-TE",
        vin="VR7EFYHT2PJ697244",
        brand="CITROEN",
        model="BERLINGO",
        version="XL 1.5 BH 100 S&S CVM6",
        rentway_unit_nr="244",
        lifecycle_status="active",
        operational_status="free",
        active=True,
    )
    db_session.add(vehicle)
    db_session.commit()
    db_session.refresh(vehicle)
    workbook = _make_workbook(
        ["Número", "Data", "Matrícula", "Nome fornecedor", "Observações"],
        [["1682", "22/06/2026", "BB-69-TE", "CARFAST RENT-A-CAR LDA (OFICINA)", "IPO"]],
    )

    response = authenticated_client.post(
        "/v2-clean/documents/import/work-orders",
        files={
            "file": (
                "ordem_de_reparo (2).xlsx",
                workbook.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert "imported_count=1" in response.headers["location"]

    page = authenticated_client.get("/v2-clean/documents")

    assert page.status_code == 200
    assert "ordem_de_reparo (2).xlsx" in page.text
    assert "linhas estruturadas" in page.text
    assert "BB-69-TE" in page.text
    assert "1682" in page.text
    assert "CARFAST RENT-A-CAR LDA (OFICINA)" in page.text


def test_clean_document_center_renders_real_invoice_for_vehicle(authenticated_client, db_session):
    vehicle = _create_vehicle(db_session)
    document = Document(
        title="Fatura oficina",
        document_type="invoice",
        classification="fleet",
        source="v2_clean_manual",
        entry_channel=None,
        original_name="fatura_11168770.pdf",
        file_name="fatura_11168770.pdf",
        storage_path="Frota/CC-11-AA/fatura_11168770.pdf",
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
        supplier_name="Filinto Mota",
        contract_number="11168770",
        status="received",
    )
    db_session.add(document)
    db_session.commit()

    page = authenticated_client.get("/v2-clean/documents")

    assert page.status_code == 200
    assert "Faturas" in page.text
    assert "Filinto Mota" in page.text
    assert "PEUGEOT 2008" in page.text
    assert "11168770" in page.text


def test_clean_document_reprocess_structured_source_materializes_rows(authenticated_client, db_session, tmp_path):
    vehicle = _create_vehicle(db_session)
    workbook = _make_workbook(
        ["Número", "Data", "Matrícula", "Nome fornecedor", "Observações"],
        [["1682", "2026-06-22", "CC-11-AA", "CARFAST RENT-A-CAR LDA (OFICINA)", "IPO"]],
    )
    source_path = tmp_path / "fo.xlsx"
    source_path.write_bytes(workbook.getvalue())
    source = Document(
        title="Importação Folhas de obra - CC-11-AA",
        document_type="general_fleet",
        classification="fleet",
        source="v2_clean_manual",
        entry_channel="structured_import",
        source_subject="work_orders:0",
        original_name="fo.xlsx",
        file_name="fo.xlsx",
        file_type="xlsx",
        file_size=source_path.stat().st_size,
        storage_provider="local",
        storage_path=str(source_path),
        status="archived",
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
        archived=True,
    )
    db_session.add(source)
    db_session.commit()
    db_session.refresh(source)

    response = authenticated_client.post(
        f"/v2-clean/documents/{source.id}/reprocess-structured-import",
        data={"return_url": f"/v2-clean/fleet/{vehicle.id}/documents"},
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert "reprocessed=1" in response.headers["location"]
    assert "reprocessed_count=1" in response.headers["location"]
    record = db_session.scalar(
        select(VehicleDocumentRecord).where(
            VehicleDocumentRecord.vehicle_id == vehicle.id,
            VehicleDocumentRecord.main_group == "work_orders",
            VehicleDocumentRecord.external_reference == "1682",
        )
    )
    assert record is not None
    assert record.document_id == source.id
    db_session.refresh(source)
    assert source.source_subject == "work_orders:1"

    page = authenticated_client.get(f"/v2-clean/fleet/{vehicle.id}/documents")
    assert page.status_code == 200
    assert "Documentação estruturada" in page.text
    assert "Reprocessar linhas" not in page.text
    assert "1682" in page.text


def test_clean_document_reprocess_legacy_source_kind_uses_filename(authenticated_client, db_session, tmp_path):
    vehicle = _create_vehicle(db_session)
    workbook = _make_workbook(
        ["Número", "Data", "Matrícula", "Nome fornecedor", "Observações"],
        [["1682", "22/06/2026", "CC-11-AA", "CARFAST RENT-A-CAR LDA (OFICINA)", "IPO"]],
    )
    source_path = tmp_path / "ordem_de_reparo (2).xlsx"
    source_path.write_bytes(workbook.getvalue())
    source = Document(
        title="Importação Folhas de obra - CC-11-AA",
        document_type="general_fleet",
        classification="fleet",
        source="v2_clean_manual",
        entry_channel="structured_import",
        source_subject="Folhas de obra:0",
        original_name="ordem_de_reparo (2).xlsx",
        file_name="ordem_de_reparo (2).xlsx",
        file_type="xlsx",
        file_size=source_path.stat().st_size,
        storage_provider="local",
        storage_path=str(source_path),
        status="archived",
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
        archived=True,
    )
    db_session.add(source)
    db_session.commit()
    db_session.refresh(source)

    response = authenticated_client.post(
        f"/v2-clean/documents/{source.id}/reprocess-structured-import",
        data={"return_url": f"/v2-clean/fleet/{vehicle.id}/documents"},
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert "reprocessed_count=1" in response.headers["location"]
    record = db_session.scalar(
        select(VehicleDocumentRecord).where(
            VehicleDocumentRecord.vehicle_id == vehicle.id,
            VehicleDocumentRecord.main_group == "work_orders",
            VehicleDocumentRecord.external_reference == "1682",
        )
    )
    assert record is not None
    assert record.document_id == source.id
    db_session.refresh(source)
    assert source.source_subject == "work_orders:1"

    module_ctx = vehicle_document_module_context(db_session, vehicle)
    assert module_ctx["group_counts"]["work_orders"] == 1
    assert any(row["title"] == "1682" for row in module_ctx["structured_rows"])
    assert any(
        any(card["group"] == "work_orders" and card["title"] == "1682" for card in event["right"])
        for event in module_ctx["timeline_events"]
    )


def test_clean_vehicle_documents_treats_legacy_import_source_as_structured(db_session):
    vehicle = _create_vehicle(db_session)
    source = Document(
        title="Importação Folhas de obra - CC-11-AA - 16/07/2026 00:03",
        document_type="general_fleet",
        classification="fleet",
        source="v2_clean_manual",
        entry_channel="v2_clean",
        source_subject="Folhas de obra:18",
        original_name="ordem_de_reparo (2).xlsx",
        file_name="ordem_de_reparo (2).xlsx",
        file_type="xlsx",
        file_size=2048,
        storage_provider="local",
        storage_path="/tmp/ordem_de_reparo.xlsx",
        status="archived",
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
        archived=True,
    )
    db_session.add(source)
    db_session.commit()

    module_ctx = vehicle_document_module_context(db_session, vehicle)

    assert module_ctx["archive_rows"] == []
    assert len(module_ctx["import_rows"]) == 1
    assert module_ctx["import_rows"][0]["import_kind"] == "work_orders"
    assert module_ctx["import_rows"][0]["imported_count"] == "18"


def test_clean_vehicle_documents_materializes_existing_import_source(db_session, tmp_path):
    vehicle = _create_vehicle(db_session)
    workbook = _make_workbook(
        ["Número", "Data", "Matrícula", "Nome fornecedor", "Observações"],
        [
            ["1682", "22/06/2026", "CC-11-AA", "CARFAST RENT-A-CAR LDA (OFICINA)", "IPO"],
            ["1608", "25/05/2026", "CC-11-AA", "CARFAST RENT-A-CAR LDA (OFICINA)", "CALÇOS ATRAS GASTOS"],
        ],
    )
    source_path = tmp_path / "ordem_de_reparo (2).xlsx"
    source_path.write_bytes(workbook.getvalue())
    source = Document(
        title="Importação Folhas de obra - CC-11-AA - 16/07/2026 00:03",
        document_type="general_fleet",
        classification="fleet",
        source="v2_clean_manual",
        entry_channel="structured_import",
        source_subject="work_orders:2",
        original_name="ordem_de_reparo (2).xlsx",
        file_name="ordem_de_reparo (2).xlsx",
        file_type="xlsx",
        file_size=source_path.stat().st_size,
        storage_provider="local",
        storage_path=str(source_path),
        status="archived",
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
        archived=True,
    )
    db_session.add(source)
    db_session.commit()
    db_session.refresh(source)

    module_ctx = vehicle_document_module_context(db_session, vehicle)

    assert module_ctx["group_counts"]["work_orders"] == 2
    assert [row["title"] for row in module_ctx["structured_rows"] if row["main_group"] == "work_orders"][:2] == [
        "1682",
        "1608",
    ]
    records = db_session.scalars(
        select(VehicleDocumentRecord).where(
            VehicleDocumentRecord.vehicle_id == vehicle.id,
            VehicleDocumentRecord.main_group == "work_orders",
        )
    ).all()
    assert len(records) == 2
    assert {record.document_id for record in records} == {source.id}


def test_clean_vehicle_documents_import_real_work_order_headers_updates_context(authenticated_client, db_session):
    vehicle = Vehicle(
        plate="BB-69-TE",
        vin="VR7EFYHT2PJ697244",
        brand="CITROEN",
        model="BERLINGO",
        version="XL 1.5 BH 100 S&S CVM6",
        rentway_unit_nr="244",
        lifecycle_status="active",
        operational_status="free",
        active=True,
    )
    db_session.add(vehicle)
    db_session.commit()
    db_session.refresh(vehicle)
    workbook = _make_workbook(
        ["Número", "Data", "Matrícula", "Nome fornecedor", "Observações"],
        [
            ["1682", "22/06/2026", "BB-69-TE", "CARFAST RENT-A-CAR LDA (OFICINA)", "IPO"],
            ["1608", "25/05/2026", "BB-69-TE", "CARFAST RENT-A-CAR LDA (OFICINA)", "CALÇOS ATRAS GASTOS"],
            ["1606", "25/05/2026", "BB-69-TE", "CARFAST RENT-A-CAR LDA (OFICINA)", "REVISAO"],
        ],
    )

    response = authenticated_client.post(
        f"/v2-clean/fleet/{vehicle.id}/documents/import/work-orders",
        files={
            "file": (
                "ordem_de_reparo (2).xlsx",
                workbook.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert "imported_count=3" in response.headers["location"]
    module_ctx = vehicle_document_module_context(db_session, vehicle)
    assert module_ctx["group_counts"]["work_orders"] == 3
    assert [row["title"] for row in module_ctx["structured_rows"] if row["main_group"] == "work_orders"][:3] == [
        "1682",
        "1608",
        "1606",
    ]
    assert any(
        any(card["group"] == "work_orders" for card in event["right"])
        for event in module_ctx["timeline_events"]
    )


def test_clean_vehicle_documents_import_impros(authenticated_client, db_session):
    vehicle = _create_vehicle(db_session)
    workbook = _make_workbook(
        [
            "Status",
            "Impro",
            "PlateNr",
            "Date_In",
            "Date_Out",
            "Garage",
            "Driven_Kms",
            "Impro_Type_Code",
            "Impro_Type_Description",
            "Driver_Name",
        ],
        [[
            "Open",
            "IMP-9281",
            "CC-11-AA",
            "2026-04-12",
            "2026-04-18",
            "Oficina Norte",
            42110,
            "MEC",
            "Avaria mecânica",
            "André",
        ]],
    )

    response = authenticated_client.post(
        f"/v2-clean/fleet/{vehicle.id}/documents/import/impros",
        files={"file": ("impros.xlsx", workbook.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        follow_redirects=False,
    )

    assert response.status_code == 303
    record = db_session.scalar(
        select(VehicleDocumentRecord).where(
            VehicleDocumentRecord.vehicle_id == vehicle.id,
            VehicleDocumentRecord.main_group == "impros",
        )
    )
    assert record is not None
    assert record.title == "IMP-9281"
    assert record.km == 42110
    assert record.supplier_name == "Oficina Norte"
    assert record.document_date == date(2026, 4, 12)
    assert record.comparison_state == "imported_rentway"
    assert record.metadata_json["_start_date"] == "2026-04-12"
    assert record.metadata_json["_date_out"] == "2026-04-18"
    module_ctx = vehicle_document_module_context(db_session, vehicle)
    impro_row = next(row for row in module_ctx["structured_rows"] if row["main_group"] == "impros")
    assert impro_row["comparison_label"] == "Importado RW"
    assert impro_row["period_display"] == "12/04/2026 a 18/04/2026"
    import_source = db_session.scalar(
        select(Document).where(
            Document.vehicle_id == vehicle.id,
            Document.entry_channel == "structured_import",
            Document.source_subject == "impros:1",
        )
    )
    assert import_source is not None


def test_clean_vehicle_documents_import_contracts(authenticated_client, db_session):
    vehicle = _create_vehicle(db_session)
    workbook = _make_workbook(
        ["Contrato", "Matrícula", "Fornecedor", "Data início", "Data fim", "Estado", "Valor mensal", "Observações"],
        [["CTR-2026-001", "CC-11-AA", "Locadora X", "2026-01-01", "2029-01-01", "Ativo", "425.50", "Contrato de aluguer operacional"]],
    )

    response = authenticated_client.post(
        f"/v2-clean/fleet/{vehicle.id}/documents/import/contracts",
        files={"file": ("contracts.xlsx", workbook.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        follow_redirects=False,
    )

    assert response.status_code == 303
    record = db_session.scalar(
        select(VehicleDocumentRecord).where(
            VehicleDocumentRecord.vehicle_id == vehicle.id,
            VehicleDocumentRecord.main_group == "contracts",
        )
    )
    assert record is not None
    assert record.title == "CTR-2026-001"
    assert record.supplier_name == "Locadora X"
    assert record.subtype == "Ativo"
    assert record.document_date == date(2026, 1, 1)
    assert record.comparison_state == "imported_rentway"
    assert record.metadata_json["_start_date"] == "2026-01-01"
    assert record.metadata_json["_end_date"] == "2029-01-01"
    assert "Fim:" not in (record.raw_description or "")
    module_ctx = vehicle_document_module_context(db_session, vehicle)
    contract_row = next(row for row in module_ctx["structured_rows"] if row["main_group"] == "contracts")
    assert contract_row["comparison_label"] == "Importado RW"
    assert contract_row["period_display"] == "01/01/2026 a 01/01/2029"
    import_source = db_session.scalar(
        select(Document).where(
            Document.vehicle_id == vehicle.id,
            Document.entry_channel == "structured_import",
            Document.source_subject == "contracts:1",
        )
    )
    assert import_source is not None


def test_clean_vehicle_documents_import_contracts_rentway_checkout_dates(authenticated_client, db_session):
    vehicle = _create_vehicle(db_session)
    workbook = _make_workbook(
        [
            "ra",
            "plate",
            "customer_name",
            "checkout_date",
            "checkin_date",
            "station",
            "origin",
            "category",
            "ndays",
            "invoiced_amount",
        ],
        [["15394", "CC-11-AA", "Rentway Cliente", "2025-08-01", "2025-08-31", "Aeroporto Porto", "Diretos", "Peugeot 208", 31, "422.18"]],
    )

    response = authenticated_client.post(
        f"/v2-clean/fleet/{vehicle.id}/documents/import/contracts",
        files={"file": ("contracts.xlsx", workbook.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        follow_redirects=False,
    )

    assert response.status_code == 303
    record = db_session.scalar(
        select(VehicleDocumentRecord).where(
            VehicleDocumentRecord.vehicle_id == vehicle.id,
            VehicleDocumentRecord.main_group == "contracts",
        )
    )
    assert record is not None
    assert record.title == "RA 15394"
    assert record.document_date == date(2025, 8, 1)
    assert record.metadata_json["_start_date"] == "2025-08-01"
    assert record.metadata_json["_end_date"] == "2025-08-31"
    assert "Estação: Aeroporto Porto" in (record.raw_description or "")
    assert "Fim:" not in (record.raw_description or "")
    center_ctx = document_center_module_context(db_session)
    contract_row = next(row for row in center_ctx["structured_rows"] if row["main_group"] == "contracts")
    assert center_ctx["pending_structured_count"] == 0
    assert contract_row["comparison_label"] == "Importado RW"
    assert contract_row["period_display"] == "01/08/2025 a 31/08/2025"


def test_clean_vehicle_documents_recovers_legacy_timeline_dates_from_raw_metadata(db_session):
    vehicle = _create_vehicle(db_session)
    work_order = VehicleDocumentRecord(
        vehicle_id=vehicle.id,
        source_record_type="structured",
        main_group="work_orders",
        title="FO 1773",
        external_reference="1773",
        document_date=None,
        source_system="work_order_import",
        metadata_json={"Data": "24/07/2026 13:42:10"},
    )
    contract = VehicleDocumentRecord(
        vehicle_id=vehicle.id,
        source_record_type="structured",
        main_group="contracts",
        title="RA 15394",
        external_reference="15394",
        document_date=None,
        source_system="contract_import",
        metadata_json={"date_out": "20250801", "date_in": "2025-08-31T18:30:00"},
    )
    portuguese_contract = VehicleDocumentRecord(
        vehicle_id=vehicle.id,
        source_record_type="structured",
        main_group="contracts",
        title="RA 49",
        external_reference="49",
        document_date=None,
        source_system="contract_import",
        metadata_json={"Data de saída": "2024-01-01T00:01:59", "Data de entrada": "2024-01-31T23:59:00"},
    )
    portuguese_impro = VehicleDocumentRecord(
        vehicle_id=vehicle.id,
        source_record_type="structured",
        main_group="impros",
        title="Impro 8461",
        external_reference="8461",
        document_date=None,
        source_system="impro_import",
        metadata_json={"Data de saída": "2024-01-28T00:00:00", "Data de entrada": "2024-02-05T00:00:00"},
    )
    db_session.add_all([work_order, contract, portuguese_contract, portuguese_impro])
    db_session.commit()

    module_ctx = vehicle_document_module_context(db_session, vehicle)
    work_order_row = next(row for row in module_ctx["structured_rows"] if row["main_group"] == "work_orders")
    contract_row = next(row for row in module_ctx["structured_rows"] if row["main_group"] == "contracts")
    portuguese_contract_row = next(
        row
        for row in module_ctx["structured_rows"]
        if row["main_group"] == "contracts" and row["external_reference"] == "49"
    )
    portuguese_impro_row = next(row for row in module_ctx["structured_rows"] if row["main_group"] == "impros")

    assert work_order_row["date_display"] == "24/07/2026"
    assert contract_row["period_display"] == "01/08/2025 a 31/08/2025"
    assert portuguese_contract_row["period_display"] == "01/01/2024 a 31/01/2024"
    assert portuguese_impro_row["period_display"] == "28/01/2024 a 05/02/2024"
    assert any(
        any(card["group"] == "work_orders" for card in event["right"])
        for event in module_ctx["timeline_events"]
    )
    assert any(
        any(card["group"] == "contracts" and card["period"] == "01/08/2025 a 31/08/2025" for card in event["center"])
        for event in module_ctx["timeline_events"]
    )


def test_clean_vehicle_documents_derives_contract_end_date_from_duration(db_session):
    vehicle = _create_vehicle(db_session)
    record = VehicleDocumentRecord(
        vehicle_id=vehicle.id,
        source_record_type="structured",
        main_group="contracts",
        title="RA 49",
        external_reference="49",
        document_date=None,
        source_system="contract_import",
        metadata_json={"checkout_date": "01.08.2025", "ndays": 31},
    )
    db_session.add(record)
    db_session.commit()

    module_ctx = vehicle_document_module_context(db_session, vehicle)
    contract_row = next(row for row in module_ctx["structured_rows"] if row["main_group"] == "contracts")

    assert contract_row["period_display"] == "01/08/2025 a 31/08/2025"


def test_clean_document_center_reimport_refreshes_existing_structured_source(authenticated_client, db_session, tmp_path):
    vehicle = _create_vehicle(db_session)
    workbook = _make_workbook(
        ["Contrato", "Matrícula", "Fornecedor", "Data início", "Data fim", "Estado", "Valor mensal", "Observações"],
        [["CTR-2026-001", "CC-11-AA", "Locadora X", "2026-01-01", "2029-01-01", "Ativo", "425.50", "Contrato de aluguer operacional"]],
    )
    workbook_bytes = workbook.getvalue()
    digest = hashlib.sha256(workbook_bytes).hexdigest()
    missing_source_path = tmp_path / "fonte-antiga-apagada.xlsx"
    source = Document(
        title="Importação Contratos - global",
        document_type="general_fleet",
        classification="fleet",
        source="v2_clean_manual",
        entry_channel="structured_import",
        source_subject="contracts:1",
        original_name="contracts.xlsx",
        file_name="contracts.xlsx",
        file_type="xlsx",
        file_hash=digest,
        file_size=len(workbook_bytes),
        storage_provider="local",
        storage_path=str(missing_source_path),
        status="archived",
        archived=True,
    )
    db_session.add(source)
    db_session.commit()
    source_id = source.id

    response = authenticated_client.post(
        "/v2-clean/documents/import/contracts",
        files={"file": ("contracts.xlsx", workbook_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        follow_redirects=False,
    )

    assert response.status_code == 303
    db_session.expire_all()
    refreshed_source = db_session.get(Document, source_id)
    assert refreshed_source is not None
    assert refreshed_source.source_subject == "contracts:1"
    assert refreshed_source.storage_path
    assert Path(refreshed_source.storage_path).exists()
    assert refreshed_source.plate is None
    record = db_session.scalar(
        select(VehicleDocumentRecord).where(
            VehicleDocumentRecord.vehicle_id == vehicle.id,
            VehicleDocumentRecord.main_group == "contracts",
            VehicleDocumentRecord.document_id == source_id,
        )
    )
    assert record is not None
    assert record.title == "CTR-2026-001"


def test_clean_vehicle_documents_materializes_zero_count_contract_source(db_session, tmp_path):
    vehicle = _create_vehicle(db_session)
    workbook = _make_workbook(
        ["Contrato", "Matrícula", "Fornecedor", "Data início", "Data fim", "Estado", "Valor mensal", "Observações"],
        [["CTR-2026-001", "CC-11-AA", "Locadora X", "2026-01-01", "2029-01-01", "Ativo", "425.50", "Contrato de aluguer operacional"]],
    )
    source_path = tmp_path / "contratos.xlsx"
    source_path.write_bytes(workbook.getvalue())
    source = Document(
        title="Importação Contratos - CC-11-AA",
        document_type="general_fleet",
        classification="fleet",
        source="v2_clean_manual",
        entry_channel="structured_import",
        source_subject="contracts:0",
        original_name="contratos.xlsx",
        file_name="contratos.xlsx",
        file_type="xlsx",
        file_size=source_path.stat().st_size,
        storage_provider="local",
        storage_path=str(source_path),
        status="archived",
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
        archived=True,
    )
    db_session.add(source)
    db_session.commit()

    module_ctx = vehicle_document_module_context(db_session, vehicle)

    assert module_ctx["group_counts"]["contracts"] == 1
    assert any(row["main_group"] == "contracts" and row["title"] == "CTR-2026-001" for row in module_ctx["structured_rows"])
    db_session.refresh(source)
    assert source.source_subject == "contracts:1"


def test_clean_document_center_materializes_zero_count_global_impro_source(db_session, tmp_path):
    vehicle = _create_vehicle(db_session)
    workbook = _make_workbook(
        [
            "Status",
            "Impro",
            "PlateNr",
            "Date_In",
            "Date_Out",
            "Garage",
            "Driven_Kms",
            "Impro_Type_Code",
            "Impro_Type_Description",
        ],
        [["Closed", "IMP-9281", "CC-11-AA", "2026-04-12", "2026-04-18", "Oficina Norte", 42110, "MEC", "Avaria mecânica"]],
    )
    source_path = tmp_path / "impros.xlsx"
    source_path.write_bytes(workbook.getvalue())
    source = Document(
        title="Importação Impros - global",
        document_type="general_fleet",
        classification="fleet",
        source="v2_clean_manual",
        entry_channel="structured_import",
        source_subject="impros:0",
        original_name="impros.xlsx",
        file_name="impros.xlsx",
        file_type="xlsx",
        file_size=source_path.stat().st_size,
        storage_provider="local",
        storage_path=str(source_path),
        status="archived",
        archived=True,
    )
    db_session.add(source)
    db_session.commit()

    center_ctx = document_center_module_context(db_session)
    vehicle_ctx = vehicle_document_module_context(db_session, vehicle)

    assert center_ctx["structured_counts"]["impros"] == 1
    assert vehicle_ctx["group_counts"]["impros"] == 1
    assert any(row["main_group"] == "impros" and row["title"] == "IMP-9281" for row in vehicle_ctx["structured_rows"])
    assert any(
        any(card["group"] == "impros" for card in event["center"])
        for event in vehicle_ctx["timeline_events"]
    )
    db_session.refresh(source)
    assert source.source_subject == "impros:1"


def test_clean_documents_counts_legacy_structured_records(db_session):
    vehicle = _create_vehicle(db_session)
    source = Document(
        title="Importação Contratos - global",
        document_type="general_fleet",
        classification="fleet",
        source="v2_clean_manual",
        entry_channel="structured_import",
        source_subject="contracts:1",
        original_name="contracts.xlsx",
        file_name="contracts.xlsx",
        file_type="xlsx",
        file_size=1024,
        storage_provider="local",
        storage_path="/tmp/contracts.xlsx",
        status="archived",
        archived=True,
    )
    db_session.add(source)
    db_session.flush()
    db_session.add(
        VehicleDocumentRecord(
            vehicle_id=vehicle.id,
            document_id=source.id,
            source_record_type="legacy_structured",
            main_group="contracts",
            status="structured",
            comparison_state="por_validar",
            external_reference="15394",
            title="RA 15394",
            plate=vehicle.plate,
            document_date=date(2026, 7, 1),
            source_system="contract_import",
        )
    )
    db_session.commit()

    center_ctx = document_center_module_context(db_session)
    vehicle_ctx = vehicle_document_module_context(db_session, vehicle)

    assert center_ctx["structured_counts"]["contracts"] == 1
    assert vehicle_ctx["group_counts"]["contracts"] == 1
    assert any(row["main_group"] == "contracts" and row["title"] == "RA 15394" for row in center_ctx["structured_rows"])
    assert any(row["main_group"] == "contracts" and row["title"] == "RA 15394" for row in vehicle_ctx["structured_rows"])


def test_clean_vehicle_documents_import_rental_agreements_format(authenticated_client, db_session):
    vehicle = _create_vehicle(db_session)
    workbook = _make_workbook(
        [
            "ra",
            "station",
            "creation_date",
            "ndays",
            "date_out",
            "date_in",
            "rate_code",
            "salesperson",
            "origin",
            "plate",
            "category",
            "category_requested",
            "invoiced_amount",
            "customer_name",
            "cashier_amount",
        ],
        [[
            15519,
            "VILA DAS AVES",
            "2026-07-10",
            28,
            "2026-07-10",
            "2026-08-07",
            "GR01",
            "DIRECTOS",
            "TO",
            "CC-11-AA",
            "PEUGEOT 208 OU SIMILAR (G)",
            "PEUGEOT 208 OU SIMILAR (G)",
            685.86,
            "NEGRELCAR",
            685.86,
        ]],
    )

    response = authenticated_client.post(
        f"/v2-clean/fleet/{vehicle.id}/documents/import/contracts",
        files={"file": ("rental_agreements.xlsx", workbook.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        follow_redirects=False,
    )

    assert response.status_code == 303
    record = db_session.scalar(
        select(VehicleDocumentRecord).where(
            VehicleDocumentRecord.vehicle_id == vehicle.id,
            VehicleDocumentRecord.main_group == "contracts",
            VehicleDocumentRecord.external_reference == "15519",
        )
    )
    assert record is not None
    assert record.title == "RA 15519"
    assert record.supplier_name == "NEGRELCAR"
    assert record.document_date is not None


def test_clean_vehicle_documents_import_rentway_exports_with_preamble(authenticated_client, db_session):
    vehicle = _create_vehicle(db_session)
    impros = _make_rentway_export_workbook(
        "Impros - 13/07/2026 15:20",
        [
            "Status",
            "Impro",
            "Station_In",
            "Date_In",
            "PlateNr",
            "BrandID",
            "ModelID",
            "Driver_Name",
            "GroupID",
            "Station_Out",
            "Date_Out",
            "Garage",
            "Driven_Kms",
            "Impro_Type_Code",
            "Impro_Type_Description",
        ],
        [[
            "Closed",
            6400,
            "AEROPORTO PORTO",
            "2026-01-12",
            "CC-11-AA",
            "CITROEN",
            "BERLINGO",
            "Filinto Mota",
            "2",
            "OFICINA",
            "2026-01-06",
            "",
            51,
            "0010",
            "OFICINA",
        ]],
    )
    contracts = _make_rentway_export_workbook(
        "Informações de Contratos - 13/07/2026 15:21",
        [
            "ra",
            "station",
            "creation_date",
            "ndays",
            "date_out",
            "date_in",
            "rate_code",
            "salesperson",
            "origin",
            "plate",
            "category",
            "category_requested",
            "invoiced_amount",
            "customer_name",
            "cashier_amount",
        ],
        [[
            48,
            "AEROPORTO PORTO",
            "2024-01-11",
            31,
            "2024-01-01",
            "2024-01-31",
            "CORP MENSAL",
            "DIRECTOS",
            "DIRECTOS",
            "CC-11-AA",
            "CITROEN BERLINGO OU SIMILAR",
            "CITROEN BERLINGO OU SIMILAR",
            711.91,
            "ROTA LATINA, LDA.",
            0,
        ]],
    )

    impro_response = authenticated_client.post(
        f"/v2-clean/fleet/{vehicle.id}/documents/import/impros",
        files={"file": ("impros.xlsx", impros.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        follow_redirects=False,
    )
    contract_response = authenticated_client.post(
        f"/v2-clean/fleet/{vehicle.id}/documents/import/contracts",
        files={"file": ("contracts.xlsx", contracts.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        follow_redirects=False,
    )

    assert impro_response.status_code == 303
    assert "imported_count=1" in impro_response.headers["location"]
    assert contract_response.status_code == 303
    assert "imported_count=1" in contract_response.headers["location"]
    assert db_session.scalar(
        select(VehicleDocumentRecord).where(
            VehicleDocumentRecord.vehicle_id == vehicle.id,
            VehicleDocumentRecord.main_group == "impros",
            VehicleDocumentRecord.external_reference == "6400",
        )
    )
    assert db_session.scalar(
        select(VehicleDocumentRecord).where(
            VehicleDocumentRecord.vehicle_id == vehicle.id,
            VehicleDocumentRecord.main_group == "contracts",
            VehicleDocumentRecord.external_reference == "48",
        )
    )
    module_ctx = vehicle_document_module_context(db_session, vehicle)
    assert module_ctx["group_counts"]["impros"] == 1
    assert module_ctx["group_counts"]["contracts"] == 1
    assert len(module_ctx["import_rows"]) == 2
    impro_rows = [row for row in module_ctx["structured_rows"] if row["main_group"] == "impros"]
    assert impro_rows
    assert impro_rows[0]["period_display"] == "06/01/2026 a 12/01/2026"
    contract_rows = [row for row in module_ctx["structured_rows"] if row["main_group"] == "contracts"]
    assert contract_rows
    assert contract_rows[0]["period_display"] == "01/01/2024 a 31/01/2024"
    assert any(
        any(card["group"] == "contracts" and card["period"] == "01/01/2024 a 31/01/2024" for card in event["center"])
        for event in module_ctx["timeline_events"]
    )
    assert any(
        any(card["group"] == "impros" for card in event["center"])
        for event in module_ctx["timeline_events"]
    )


def test_clean_vehicle_documents_import_real_rentway_portuguese_date_headers(
    authenticated_client,
    db_session,
):
    vehicle = _create_vehicle(db_session)
    impros = _make_rentway_export_workbook(
        "Impros - 25/07/2026 11:22",
        [
            "Estado",
            "Impro",
            "Estação de devolução",
            "Data de entrada",
            "Matrícula",
            "Marca",
            "Modelo do Veículo",
            "Condutor",
            "ID Grupo",
            "Estação de saída",
            "Data de saída",
            "Oficina",
            "Km percorridos",
            "Código de tipo de Impro",
            "Descrição do tipo de Impro",
        ],
        [[
            "Closed",
            8461,
            "AEROPORTO PORTO",
            datetime(2024, 2, 5),
            "CC-11-AA",
            "CITROEN",
            "BERLINGO",
            "Filinto Mota",
            "2",
            "AEROPORTO PORTO",
            datetime(2024, 1, 28),
            "",
            51,
            "SERVICO",
            "SERVIÇO STAFF",
        ]],
    )
    contracts = _make_rentway_export_workbook(
        "Informações de Contratos - 25/07/2026 13:23",
        [
            "Contrato",
            "Estação",
            "Data de criação",
            "Dias",
            "Data de saída",
            "Data de entrada",
            "Código da tarifa",
            "Vendedor",
            "Origem",
            "Matrícula",
            "Categoria",
            "Categoria solicitada",
            "Qtd. faturada",
            "Nome do cliente",
            "Valor de caixa",
        ],
        [[
            49,
            "AEROPORTO PORTO",
            datetime(2024, 1, 5, 16, 28, 19),
            31,
            datetime(2024, 1, 1, 0, 1, 59),
            datetime(2024, 1, 31, 23, 59),
            "CORP MENSAL",
            "DIRECTOS",
            "DIRECTOS",
            "CC-11-AA",
            "CITROEN BERLINGO OU SIMILAR",
            "CITROEN BERLINGO OU SIMILAR",
            711.91,
            "ROTA LATINA, LDA.",
            0,
        ]],
    )

    impro_response = authenticated_client.post(
        f"/v2-clean/fleet/{vehicle.id}/documents/import/impros",
        files={
            "file": (
                "impros.xlsx",
                impros.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        follow_redirects=False,
    )
    contract_response = authenticated_client.post(
        f"/v2-clean/fleet/{vehicle.id}/documents/import/contracts",
        files={
            "file": (
                "contracts.xlsx",
                contracts.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        follow_redirects=False,
    )

    assert impro_response.status_code == 303
    assert contract_response.status_code == 303
    module_ctx = vehicle_document_module_context(db_session, vehicle)
    impro_row = next(row for row in module_ctx["structured_rows"] if row["main_group"] == "impros")
    contract_row = next(row for row in module_ctx["structured_rows"] if row["main_group"] == "contracts")

    assert impro_row["period_display"] == "28/01/2024 a 05/02/2024"
    assert contract_row["period_display"] == "01/01/2024 a 31/01/2024"
    assert any(
        any(card["group"] == "contracts" and card["period"] == "01/01/2024 a 31/01/2024" for card in event["center"])
        for event in module_ctx["timeline_events"]
    )
    assert any(
        any(card["group"] == "impros" and card["period"] == "28/01/2024 a 05/02/2024" for card in event["center"])
        for event in module_ctx["timeline_events"]
    )
