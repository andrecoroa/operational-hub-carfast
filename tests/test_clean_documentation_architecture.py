from io import BytesIO
import json
from pathlib import Path
from datetime import date

from openpyxl import Workbook
import pytest
from sqlalchemy import func, select

import app.web.router as web_router
from app.core.config import settings
from app.models.audit import AuditLog
from app.models.documents import (
    DiagnosticDocument,
    Document,
    DocumentEvent,
    DocumentLink,
    DocumentWorkflowState,
    VehicleDocumentRecord,
    VehicleDocumentRecordTag,
)
from app.models.imports import ImportBatch, ImportFile
from app.models.tasks import Task, TaskDocument
from app.models.vehicles import Vehicle
from app.models.workshop import WorkshopTechnicalReading


def _rentway_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Vehicles"
    sheet.append(
        [
            "platenr",
            "chassinr",
            "unitnr",
            "brandid",
            "modelid",
            "CurrentStatus",
        ]
    )
    sheet.append(
        [
            "AA-10-BB",
            "VF3TEST0000000001",
            "9901",
            "PEUGEOT",
            "208",
            "FREE",
        ]
    )
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _work_orders_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        ["Número", "Data", "Matrícula", "Nome fornecedor", "Observações"]
    )
    sheet.append(
        ["FO-991", "2026-07-20", "AA-10-BB", "CarFast Oficina", "Revisão"]
    )
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _triage_document(index: int) -> Document:
    return Document(
        title=f"Entrada universal {index:02d}",
        document_type="unknown_document",
        classification="triage",
        source="email" if index % 2 else "v2_clean_manual",
        entry_channel="email" if index % 2 else "document_inbox",
        original_name=f"entrada_{index:02d}.pdf",
        file_name=f"entrada_{index:02d}.pdf",
        storage_provider="local",
        storage_path=f"Frota/_POR_ASSOCIAR/entrada_{index:02d}.pdf",
        status="unclassified" if index % 2 else "pending_triage",
    )


def test_new_documentation_navigation_and_legacy_center_coexist(
    authenticated_client,
):
    center = authenticated_client.get("/v2-clean/documentation")
    legacy = authenticated_client.get("/v2-clean/documents")

    assert center.status_code == 200
    assert "Centro de documentação" in center.text
    assert "Triagem" in center.text
    assert "Modelos de extração" in center.text
    assert legacy.status_code == 200
    assert "gestão v2" in legacy.text


@pytest.mark.parametrize(
    "path",
    [
        "/v2-clean/documentation/imports",
        "/v2-clean/documentation/imports/rentway",
        "/v2-clean/documentation/imports/reports",
        "/v2-clean/documentation/imports/invoices",
        "/v2-clean/documentation/imports/fleet",
        "/v2-clean/documentation/imports/other",
        "/v2-clean/documentation/invoices",
        "/v2-clean/documentation/archive",
        "/v2-clean/documentation/extraction-models",
    ],
)
def test_documentation_workspaces_render(authenticated_client, path):
    response = authenticated_client.get(path)

    assert response.status_code == 200
    assert "Centro de documentação" in response.text


def test_documentation_primary_flow_and_family_workspaces(authenticated_client):
    treatment = authenticated_client.get(
        "/v2-clean/documentation/treatment?family=diagnostics"
    )
    imports = authenticated_client.get(
        "/v2-clean/documentation/imports/invoices"
    )
    archive = authenticated_client.get(
        "/v2-clean/documentation/archive?family=fleet"
    )

    assert treatment.status_code == 200
    assert "Tratamento" in treatment.text
    assert "Diagnósticos por tratar" in treatment.text
    assert "Doc. Frota" in treatment.text
    assert imports.status_code == 200
    assert "Documentos importados" in imports.text
    assert "Lotes de importação" in imports.text
    assert "Modelos de extração" in imports.text
    assert archive.status_code == 200
    assert "Histórico documental concluído" in archive.text


def test_fleet_document_workspace_exposes_financial_plan_importer(
    authenticated_client,
):
    page = authenticated_client.get("/v2-clean/documentation/imports/fleet")

    assert page.status_code == 200
    assert "Doc. Frota" in page.text
    assert "Planos financeiros" in page.text
    assert 'href="/v2-clean/documentation/financial-plans"' in page.text
    assert "/v2-clean/documentation/imports/rentway?tab=fleet" in page.text
    assert "Estruturado com pré-visualização" in page.text

    importer = authenticated_client.get(
        "/v2-clean/documentation/financial-plans"
    )
    assert importer.status_code == 200
    assert "Voltar a Doc. Frota" in importer.text


def test_invoice_workspace_owns_import_actions_and_hides_legacy_center(
    authenticated_client,
):
    page = authenticated_client.get("/v2-clean/documentation/imports/invoices")

    assert page.status_code == 200
    assert 'action="/v2-clean/documents/import/archive-batch"' in page.text
    assert 'action="/v2-clean/documents/import/pending-invoices"' in page.text
    assert 'action="/v2-clean/documents/import/invoice-ocr-manifest"' in page.text
    assert 'name="return_to" value="documentation"' in page.text
    assert "/v2-clean/documentation/invoices#pending-invoices" in page.text
    assert "/v2-clean/documents#invoice-import" not in page.text
    assert "Centro anterior" not in page.text


def test_invoice_import_errors_return_to_new_workspace(authenticated_client):
    response = authenticated_client.post(
        "/v2-clean/documents/import/archive-batch",
        data={"return_to": "documentation"},
        files={"file": ("faturas.txt", b"invalid", "text/plain")},
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert response.headers["location"].startswith(
        "/v2-clean/documentation/imports/invoices?batch_error="
    )


def test_clean_document_navigation_does_not_link_back_to_legacy_center(
    authenticated_client,
):
    home = authenticated_client.get("/v2-clean")
    ocr_validation = authenticated_client.get(
        "/v2-clean/documents/ocr-validation"
    )

    assert home.status_code == 200
    assert 'href="/v2-clean/documentation"' in home.text
    assert 'href="/v2-clean/documents"><span>Documentação' not in home.text
    assert ocr_validation.status_code == 200
    assert 'href="/v2-clean/documentation/extraction-models"' in ocr_validation.text


def test_document_inbox_forms_return_to_their_clean_workspaces(
    authenticated_client,
):
    center = authenticated_client.get("/v2-clean/documentation")
    other = authenticated_client.get("/v2-clean/documentation/imports/other")

    assert center.status_code == 200
    assert 'name="return_to" value="documentation_center"' in center.text
    assert other.status_code == 200
    assert 'name="return_to" value="documentation_other"' in other.text

    invalid = authenticated_client.post(
        "/v2-clean/documents/import/inbox",
        data={"return_to": "documentation_other"},
        files={"files": ("ficheiro.txt", b"unsupported", "text/plain")},
        follow_redirects=False,
    )
    assert invalid.status_code == 303
    assert invalid.headers["location"].startswith(
        "/v2-clean/documentation/imports/other?inbox_error="
    )


def test_clean_invoices_monitor_routes_expected_records_to_unified_treatment(
    authenticated_client,
    db_session,
):
    vehicle = Vehicle(
        plate="AA-10-BB",
        vin="VF3EXPECTED0000001",
        rentway_unit_nr="9901",
    )
    pending = VehicleDocumentRecord(
        source_record_type="pending_import",
        main_group="invoices",
        status="pending",
        external_reference="HFO/3000/2026",
        supplier_name="Fornecedor Teste",
        source_system="pending_document_import",
        metadata_json={"supplier_nif": "500000000", "expected_total": "123,45"},
    )
    db_session.add_all([vehicle, pending])
    db_session.commit()

    page = authenticated_client.get("/v2-clean/documentation/invoices")
    treatment = authenticated_client.get(
        "/v2-clean/documentation/treatment?family=invoices"
    )
    associated = authenticated_client.post(
        f"/v2-clean/documents/pending/{pending.id}/associate",
        data={"identifier": "9901", "return_to": "clean"},
        follow_redirects=False,
    )
    db_session.expire_all()

    assert page.status_code == 200
    assert "Faturas esperadas" in page.text
    assert "HFO/3000/2026" not in page.text
    assert "<h2>Faturas recebidas</h2>" not in page.text
    assert "<table" not in page.text
    assert "Abrir Tratamento" in page.text
    assert treatment.status_code == 200
    assert "HFO/3000/2026" in treatment.text
    assert "500000000" in treatment.text
    assert "Fatura esperada" in treatment.text
    assert "Sem documento real" in treatment.text
    assert associated.status_code == 303
    assert associated.headers["location"].startswith(
        "/v2-clean/documentation/invoices"
    )
    assert db_session.get(VehicleDocumentRecord, pending.id).vehicle_id == vehicle.id


def test_reports_workspace_exposes_batch_reprocessing(
    authenticated_client,
    db_session,
):
    document = Document(
        title="Diagnóstico pendente",
        document_type="diagnostic_report",
        original_name="diagnostico.pdf",
        file_name="diagnostico.pdf",
        storage_provider="local",
        storage_path="diagnostico.pdf",
        status="received",
    )
    db_session.add(document)
    db_session.flush()
    db_session.add(
        DiagnosticDocument(
            document_id=document.id,
            diagnostic_type="vehicle_diagnostic",
            diagnostic_status="received",
            association_status="unassociated",
            ocr_status="not_requested",
            validation_status="pending",
        )
    )
    db_session.commit()

    page = authenticated_client.get(
        "/v2-clean/documentation/imports/reports?tab=diagnostics"
    )

    assert page.status_code == 200
    assert "Lotes de diagnósticos" in page.text
    assert "Extração pendente" in page.text
    assert "/v2-clean/diagnostics/batches/reprocess" in page.text
    assert 'id="reports-import-dialog"' in page.text
    assert 'action="/v2-clean/documents/import/historical-reports"' in page.text
    assert 'name="return_to" value="documentation"' in page.text


def test_extraction_models_lists_builtin_extractors_without_database_mappings(
    authenticated_client,
):
    page = authenticated_client.get("/v2-clean/documentation/extraction-models")

    assert page.status_code == 200
    assert "Extratores incorporados" in page.text
    assert "Diagnósticos Autel" in page.text
    assert "Diagnósticos Stellantis" in page.text
    assert "OCR de faturas" in page.text
    assert "Famílias de diagnóstico reconhecidas" in page.text
    assert "Plano de manutenção do construtor" in page.text


def test_diagnostic_batch_reprocessing_synchronizes_clean_and_reading_states(
    authenticated_client,
    db_session,
    tmp_path,
    monkeypatch,
):
    source_path = tmp_path / "A_ILM_VF3YBBPFBPG057051_260117_0952.pdf"
    source_path.write_bytes(b"%PDF-1.4 diagnostic")
    vehicle = Vehicle(
        plate="AA-10-BB",
        vin="VF3YBBPFBPG057051",
        brand="PEUGEOT",
        model="208",
    )
    db_session.add(vehicle)
    db_session.flush()
    document = Document(
        title="Diagnóstico pendente",
        document_type="diagnostic_report",
        original_name=source_path.name,
        file_name=source_path.name,
        storage_provider="local",
        storage_path=str(source_path),
        status="received",
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
    )
    db_session.add(document)
    db_session.flush()
    profile = DiagnosticDocument(
        document_id=document.id,
        diagnostic_type="other_diagnostic",
        diagnostic_status="processing",
        association_status="unassociated",
        ocr_status="pending",
        validation_status="pending",
    )
    db_session.add(profile)
    db_session.flush()
    reading = WorkshopTechnicalReading(
        process_id=None,
        vehicle_id=vehicle.id,
        user_id=1,
        reading_type="ILM",
        summary="Informação de lubrificação",
        data_json={"document_id": document.id},
        storage_provider="local",
        external_url=f"/v2-clean/documents/{document.id}/file",
        status="pending_extraction",
        updated_by_id=1,
    )
    db_session.add(reading)
    db_session.flush()
    db_session.add(
        DocumentLink(
            document_id=document.id,
            entity_type="workshop_technical_reading",
            entity_id=str(reading.id),
            category="ILM",
        )
    )
    db_session.add(
        DocumentEvent(
            document_id=document.id,
            action="historical_report.imported",
            new_value=json.dumps({"archive_name": "Lote diagnóstico/relatorio.pdf"}),
            user_id=1,
        )
    )
    db_session.commit()

    monkeypatch.setattr(
        web_router,
        "extract_diagnostic_pdf",
        lambda _path: {
            "source_sha256": "a" * 64,
            "extractor_name": "carfast_diagnostic_pdf",
            "extractor_version": web_router.DIAGNOSTIC_EXTRACTOR_VERSION,
            "parser_name": "diagnostic_parser",
            "parser_version": web_router.DIAGNOSTIC_PARSER_VERSION,
            "source_machine": "Autel",
            "source_family": "ILM",
            "source_filename": source_path.name,
            "source_page_count": 1,
            "extraction_method": "native_text",
            "extraction_status": "extracted",
            "confidence": 0.98,
            "native_text": "Informações de lubrificação",
            "ocr_text": None,
            "raw_metadata": {},
            "pages": [],
            "normalized": {
                "vin": "VF3YBBPFBPG057051",
                "report_datetime": "2026-01-17 09:52:00",
                "diagnostic_type": "engine_lubrication_information",
            },
            "dynamic_fields": {"observations": [], "label_values": [], "dtcs": []},
            "warnings": [],
        },
    )
    monkeypatch.setattr(web_router, "_document_resolved_file", lambda _document: source_path)

    response = authenticated_client.post(
        "/v2-clean/diagnostics/batches/reprocess",
        data={"batch_name": "Lote diagnóstico"},
        follow_redirects=False,
    )

    assert response.status_code == 303
    db_session.expire_all()
    refreshed_document = db_session.get(Document, document.id)
    refreshed_profile = db_session.get(DiagnosticDocument, profile.id)
    refreshed_reading = db_session.get(WorkshopTechnicalReading, reading.id)
    workflow = db_session.scalar(
        select(DocumentWorkflowState).where(
            DocumentWorkflowState.document_id == document.id
        )
    )
    assert refreshed_document.status == "pending_validation"
    assert refreshed_profile.ocr_status == "extracted"
    assert refreshed_profile.diagnostic_status == "ready_for_review"
    assert refreshed_reading.status == "pending_validation"
    assert workflow.extraction_status == "extracted"
    assert workflow.destination_status == "diagnostics"


@pytest.mark.parametrize(
    "path",
    [
        "/v2-clean/documentation",
        "/v2-clean/documentation/triage",
        "/v2-clean/documentation/imports",
        "/v2-clean/documentation/imports/rentway",
        "/v2-clean/documentation/imports/reports",
        "/v2-clean/documentation/imports/invoices",
        "/v2-clean/documentation/imports/other",
        "/v2-clean/documentation/invoices",
        "/v2-clean/diagnostics",
        "/v2-clean/documentation/archive",
        "/v2-clean/documentation/extraction-models",
    ],
)
def test_documentation_pages_use_compact_module_header(authenticated_client, path):
    response = authenticated_client.get(path)

    assert response.status_code == 200
    assert '<header class="doc-arch-header">' in response.text


def test_documentation_headers_keep_only_essential_metrics(authenticated_client):
    center = authenticated_client.get("/v2-clean/documentation")
    invoices = authenticated_client.get("/v2-clean/documentation/invoices")

    assert 'aria-label="Indicadores essenciais"' in center.text
    assert "Importações recentes" not in center.text
    assert "Diagnósticos pendentes" not in center.text
    assert 'class="doc-arch-kpis doc-invoice-monitor-kpis"' in invoices.text
    assert "Serviços pendentes" in invoices.text
    assert "Por extrair" in invoices.text
    assert "Sem associação" not in invoices.text


def test_invoice_extraction_queue_includes_every_non_extracted_state(
    authenticated_client,
    db_session,
):
    vehicle = Vehicle(plate="SE-10-XT", active=True)
    db_session.add(vehicle)
    db_session.flush()
    statuses = [None, "not_requested", "queued", "processing", "failed", "extracted"]
    documents = []
    for index, extraction_status in enumerate(statuses):
        document = Document(
            title=f"Fatura estado {extraction_status or 'sem-workflow'}",
            document_type="workshop_supplier_invoice",
            classification="invoice",
            source="document_inbox",
            original_name=f"state-{index}.pdf",
            file_name=f"state-{index}.pdf",
            storage_provider="local",
            storage_path=f"Faturas/state-{index}.pdf",
            status="pending_validation",
            vehicle_id=vehicle.id,
            plate=vehicle.plate,
        )
        db_session.add(document)
        db_session.flush()
        documents.append(document)
        if extraction_status is not None:
            db_session.add(
                DocumentWorkflowState(
                    document_id=document.id,
                    ingestion_status="completed",
                    association_status="associated",
                    extraction_status=extraction_status,
                    validation_status="pending",
                    destination_status="invoices",
                    invoice_nature="operacional",
                )
            )
    db_session.commit()

    response = authenticated_client.get(
        "/v2-clean/documentation/treatment?family=invoices&stage=extract&q=SE-10-XT"
    )

    assert response.status_code == 200
    for extraction_status in statuses[:-1]:
        assert f"Fatura estado {extraction_status or 'sem-workflow'}" in response.text
    assert "Fatura estado extracted" not in response.text
    assert "Sem extração / reprocessar" in response.text


def test_vehicle_document_page_links_to_canonical_invoice_treatment(
    authenticated_client,
    db_session,
):
    vehicle = Vehicle(plate="FV-20-XT", active=True)
    db_session.add(vehicle)
    db_session.commit()

    response = authenticated_client.get(f"/v2-clean/fleet/{vehicle.id}/documents")

    assert response.status_code == 200
    assert "/v2-clean/documentation/treatment?family=invoices&q=FV-20-XT" in response.text
    assert (
        "/v2-clean/documentation/treatment?family=invoices&stage=extract&q=FV-20-XT"
        in response.text
    )


def test_rentway_import_actions_are_scoped_to_active_page(authenticated_client):
    work_orders = authenticated_client.get(
        "/v2-clean/documentation/imports/rentway?tab=work_orders"
    )
    fleet = authenticated_client.get(
        "/v2-clean/documentation/imports/rentway?tab=fleet"
    )

    assert ">+ Importar<" not in work_orders.text
    assert 'data-import-kind="work_order_details"' in work_orders.text
    assert 'data-import-kind="work_orders"' in work_orders.text
    assert '<select name="import_kind"' not in work_orders.text
    assert 'data-import-kind="fleet"' in fleet.text
    assert 'data-import-kind="work_orders"' not in fleet.text


@pytest.mark.parametrize(
    ("path", "expected_action", "excluded_action"),
    [
        (
            "/v2-clean/documentation/imports/reports?tab=service_box",
            "Importar relatórios",
            "Importar faturas",
        ),
        (
            "/v2-clean/documentation/imports/invoices",
            "Importar faturas",
            "Importar relatórios",
        ),
        (
            "/v2-clean/documentation/imports/other",
            "Importar outros documentos",
            "Importar faturas",
        ),
    ],
)
def test_each_import_workspace_exposes_only_its_own_import_action(
    authenticated_client,
    path,
    expected_action,
    excluded_action,
):
    response = authenticated_client.get(path)

    assert response.status_code == 200
    assert expected_action in response.text
    assert excluded_action not in response.text


def test_triage_is_universal_and_server_paginated(
    authenticated_client,
    db_session,
):
    db_session.add_all([_triage_document(index) for index in range(31)])
    db_session.commit()

    first = authenticated_client.get(
        "/v2-clean/documentation/triage?page=1&page_size=10"
    )
    second = authenticated_client.get(
        "/v2-clean/documentation/triage?page=2&page_size=10"
    )

    assert first.status_code == 200
    assert second.status_code == 200
    assert "31 registos" in first.text
    assert "Entrada universal 30" in first.text
    assert "Entrada universal 20" in second.text
    assert "Entrada universal 30" not in second.text
    assert "Email" in first.text


def test_triage_decision_preserves_filters_and_queues_diagnostic_extraction(
    authenticated_client,
    db_session,
):
    document = _triage_document(80)
    db_session.add(document)
    db_session.commit()

    response = authenticated_client.post(
        f"/v2-clean/documentation/triage/{document.id}",
        data={
            "destination": "diagnostics",
            "decision_reason": "Relatório técnico confirmado",
            "q": "entrada",
            "origin": "email",
            "confidence": "low",
            "page": "2",
            "page_size": "10",
        },
        follow_redirects=False,
    )
    db_session.expire_all()

    assert response.status_code == 303
    assert response.headers["location"].startswith(
        "/v2-clean/documentation/triage?"
    )
    assert "q=entrada" in response.headers["location"]
    assert "origin=email" in response.headers["location"]
    assert "confidence=low" in response.headers["location"]
    assert "page=2" in response.headers["location"]
    assert "page_size=10" in response.headers["location"]
    assert "saved=1" in response.headers["location"]

    stored = db_session.get(Document, document.id)
    profile = db_session.scalar(
        select(DiagnosticDocument).where(
            DiagnosticDocument.document_id == document.id
        )
    )
    workflow = db_session.scalar(
        select(DocumentWorkflowState).where(
            DocumentWorkflowState.document_id == document.id
        )
    )
    assert stored.document_type == "workshop_diagnostic"
    assert stored.status == "pending_extraction"
    assert profile.ocr_status == "pending"
    assert profile.validation_status == "pending"
    assert workflow.extraction_status == "queued"
    assert workflow.validation_status == "pending"
    assert workflow.destination_status == "diagnostics"


def test_triage_rejects_invalid_invoice_nature_without_server_error(
    authenticated_client,
    db_session,
):
    document = _triage_document(81)
    db_session.add(document)
    db_session.commit()

    response = authenticated_client.post(
        f"/v2-clean/documentation/triage/{document.id}",
        data={
            "destination": "invoices",
            "invoice_nature": "operacional,financeira",
            "page": "3",
            "page_size": "25",
        },
        follow_redirects=False,
    )
    db_session.expire_all()

    assert response.status_code == 303
    assert "error=invoice_nature" in response.headers["location"]
    assert "page=3" in response.headers["location"]
    stored = db_session.get(Document, document.id)
    assert stored.document_type == "unknown_document"
    assert stored.status == "unclassified"


def test_invoice_nature_endpoint_requires_single_valid_nature(
    authenticated_client,
    db_session,
):
    document = Document(
        title="Fatura por decidir",
        document_type="workshop_supplier_invoice",
        classification="invoice",
        source="v2_clean_manual",
        original_name="fatura.pdf",
        file_name="fatura.pdf",
        storage_provider="local",
        storage_path="Frota/_POR_ASSOCIAR/fatura.pdf",
        status="received",
    )
    db_session.add(document)
    db_session.commit()

    invalid = authenticated_client.post(
        f"/v2-clean/documentation/invoices/{document.id}/nature",
        data={"nature": "operacional,financeira"},
        follow_redirects=False,
    )
    valid = authenticated_client.post(
        f"/v2-clean/documentation/invoices/{document.id}/nature",
        data={
            "nature": "financeira",
            "suggested_nature": "financeira",
            "suggestion_confidence": "0,91",
        },
        follow_redirects=False,
    )
    db_session.expire_all()

    assert invalid.status_code == 303
    assert "error=" in invalid.headers["location"]
    assert valid.status_code == 303
    state = db_session.scalar(
        select(DocumentWorkflowState).where(
            DocumentWorkflowState.document_id == document.id
        )
    )
    stored = db_session.get(Document, document.id)
    assert state.invoice_nature == "financeira"
    assert state.suggestion_confidence == 0.91
    assert stored.document_type == "finance_supplier_invoice"
    assert stored.archived is True


def test_rentway_preview_does_not_write_and_confirm_uses_durable_source(
    authenticated_client,
    db_session,
    monkeypatch,
    tmp_path,
):
    archive_root = tmp_path / "documents"
    monkeypatch.setattr(settings, "document_archive_root", str(archive_root))
    content = _rentway_workbook()

    preview_response = authenticated_client.post(
        "/v2-clean/documentation/imports/rentway/preview",
        data={"import_kind": "fleet"},
        files={
            "file": (
                "rentway_fleet.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        follow_redirects=False,
    )
    db_session.expire_all()

    assert preview_response.status_code == 303
    assert "/preview/" in preview_response.headers["location"]
    assert db_session.scalar(select(func.count()).select_from(Vehicle)) == 0
    assert db_session.scalar(select(func.count()).select_from(ImportBatch)) == 0

    preview_page = authenticated_client.get(preview_response.headers["location"])
    assert preview_page.status_code == 200
    assert "nenhum dado foi gravado" in preview_page.text
    token = preview_response.headers["location"].rsplit("/", 1)[-1]
    confirmed = authenticated_client.post(
        "/v2-clean/documentation/imports/rentway/confirm",
        data={"preview_token": token},
        follow_redirects=False,
    )
    db_session.expire_all()

    vehicle = db_session.scalar(select(Vehicle).where(Vehicle.plate == "AA-10-BB"))
    batch = db_session.scalar(
        select(ImportBatch).where(ImportBatch.import_type == "rentway_fleet")
    )
    import_file = db_session.scalar(
        select(ImportFile).where(ImportFile.batch_id == batch.id)
    )
    assert confirmed.status_code == 303
    assert vehicle is not None
    assert batch.status == "completed"
    assert Path(import_file.storage_path).is_file()
    assert str(archive_root) in import_file.storage_path


def test_rentway_workspace_keeps_separate_work_order_buttons(
    authenticated_client,
):
    page = authenticated_client.get(
        "/v2-clean/documentation/imports/rentway?tab=work_orders"
    )

    assert page.status_code == 200
    assert "Importar folhas" in page.text
    assert "Importar detalhes" in page.text
    assert "Importar conjunto" not in page.text


def test_structured_rentway_confirmation_creates_reprocessable_batch(
    authenticated_client,
    db_session,
    monkeypatch,
    tmp_path,
):
    archive_root = tmp_path / "documents"
    monkeypatch.setattr(settings, "document_archive_root", str(archive_root))
    vehicle = Vehicle(
        plate="AA-10-BB",
        vin="VF3WORKORDER000001",
        rentway_unit_nr="9901",
        brand="PEUGEOT",
        model="208",
        lifecycle_status="active",
        operational_status="free",
    )
    db_session.add(vehicle)
    db_session.commit()

    preview = authenticated_client.post(
        "/v2-clean/documentation/imports/rentway/preview",
        data={"import_kind": "work_orders"},
        files={
            "file": (
                "folhas_obra.xlsx",
                _work_orders_workbook(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        follow_redirects=False,
    )
    token = preview.headers["location"].rsplit("/", 1)[-1]
    confirmed = authenticated_client.post(
        "/v2-clean/documentation/imports/rentway/confirm",
        data={"preview_token": token},
        follow_redirects=False,
    )
    db_session.expire_all()

    batch = db_session.scalar(
        select(ImportBatch).where(
            ImportBatch.import_type == "rentway_work_orders"
        )
    )
    import_file = db_session.scalar(
        select(ImportFile).where(ImportFile.batch_id == batch.id)
    )
    record = db_session.scalar(
        select(VehicleDocumentRecord).where(
            VehicleDocumentRecord.external_reference == "FO-991"
        )
    )
    assert preview.status_code == 303
    assert confirmed.status_code == 303
    assert batch.status == "completed"
    assert record is not None
    assert Path(import_file.storage_path).is_file()

    reprocessed = authenticated_client.post(
        f"/v2-clean/documentation/imports/rentway/batches/{batch.id}/reprocess",
        follow_redirects=False,
    )
    db_session.expire_all()

    assert reprocessed.status_code == 303
    assert (
        db_session.scalar(
            select(func.count())
            .select_from(ImportBatch)
            .where(ImportBatch.import_type == "rentway_work_orders")
        )
        == 2
    )
    assert (
        db_session.scalar(
            select(func.count())
            .select_from(VehicleDocumentRecord)
            .where(VehicleDocumentRecord.external_reference == "FO-991")
        )
        == 1
    )


def test_treatment_groups_invoices_by_supplier_and_keeps_preview_links(
    authenticated_client,
    db_session,
):
    documents = []
    for index, supplier in enumerate(("Dispnal Pneus, S.A.", "Dispnal Pneus, S.A.", "Outro Fornecedor")):
        document = Document(
            title=f"Fatura agrupada {index}",
            document_type="workshop_supplier_invoice",
            classification="invoice",
            source="document_inbox",
            original_name=f"fatura_{index}.pdf",
            file_name=f"fatura_{index}.pdf",
            storage_provider="local",
            storage_path=f"Faturas/fatura_{index}.pdf",
            status="pending_validation",
            supplier_name=supplier,
        )
        db_session.add(document)
        db_session.flush()
        db_session.add(
            DocumentWorkflowState(
                document_id=document.id,
                ingestion_status="completed",
                association_status="unassociated",
                extraction_status="extracted",
                validation_status="pending",
                destination_status="invoices",
                invoice_nature="por_classificar",
            )
        )
        documents.append(document)
    db_session.commit()

    response = authenticated_client.get(
        "/v2-clean/documentation/treatment?family=invoices&group=Dispnal%20Pneus%2C%20S.A."
    )

    assert response.status_code == 200
    assert "Dispnal Pneus, S.A." in response.text
    assert "2 pendentes" in response.text
    assert "Fatura agrupada 0" in response.text
    assert "Fatura agrupada 1" in response.text
    assert "Fatura agrupada 2" not in response.text
    assert f'data-preview-src="/v2-clean/documents/{documents[0].id}/file?inline=1"' in response.text
    assert 'action="/v2-clean/documentation/treatment/bulk"' in response.text


def test_treatment_file_preview_uses_durable_storage_for_central_sources(
    authenticated_client,
    db_session,
    monkeypatch,
    tmp_path,
):
    archive_root = tmp_path / "archive"
    invoice_path = archive_root / "Faturas" / "central.pdf"
    invoice_path.parent.mkdir(parents=True)
    invoice_path.write_bytes(b"%PDF-1.4\n% central preview\n")
    monkeypatch.setattr(settings, "document_archive_root", str(archive_root))
    document = Document(
        title="Fatura central com PDF",
        document_type="workshop_supplier_invoice",
        classification="invoice",
        source="document_inbox",
        original_name="central.pdf",
        file_name="central.pdf",
        file_type="application/pdf",
        storage_provider="local",
        storage_path="Faturas/central.pdf",
        status="pending_validation",
    )
    db_session.add(document)
    db_session.commit()

    response = authenticated_client.get(
        f"/v2-clean/documents/{document.id}/file?inline=1"
    )

    assert response.status_code == 200
    assert response.content.startswith(b"%PDF-1.4")
    assert response.headers["content-type"].startswith("application/pdf")


def test_document_removal_requires_reason_and_records_global_audit(
    authenticated_client,
    db_session,
):
    document = Document(
        title="Documento a remover com auditoria",
        document_type="vehicle_document",
        classification="fleet",
        source="v2_clean_manual",
        original_name="remover.pdf",
        file_name="remover.pdf",
        storage_provider="local",
        storage_path="Frota/remover.pdf",
        status="received",
    )
    db_session.add(document)
    db_session.commit()

    missing_reason = authenticated_client.post(
        f"/v2-clean/documents/{document.id}/remove",
        data={"return_url": "/v2-clean/documentation/treatment?family=fleet"},
        follow_redirects=False,
    )
    removed = authenticated_client.post(
        f"/v2-clean/documents/{document.id}/remove",
        data={
            "reason": "Duplicado confirmado no documento original",
            "return_url": "/v2-clean/documentation/treatment?family=fleet",
        },
        follow_redirects=False,
    )
    db_session.expire_all()

    assert "remove_error=missing_reason" in missing_reason.headers["location"]
    assert removed.status_code == 303
    assert db_session.get(Document, document.id).status == "removed"
    assert db_session.scalar(
        select(DocumentEvent).where(
            DocumentEvent.document_id == document.id,
            DocumentEvent.action == "document.removed",
        )
    )
    assert db_session.scalar(
        select(AuditLog).where(
            AuditLog.action == "document.removed",
            AuditLog.entity_type == "document",
            AuditLog.entity_id == str(document.id),
        )
    )


def test_treatment_groups_diagnostics_by_extracted_type(
    authenticated_client,
    db_session,
):
    document = Document(
        title="Relatório técnico",
        document_type="workshop_diagnostic",
        classification="technical_report",
        source="historical_report_import",
        original_name="relatorio_acer.pdf",
        file_name="relatorio_acer.pdf",
        storage_provider="local",
        storage_path="Diagnosticos/relatorio_acer.pdf",
        status="pending_validation",
    )
    db_session.add(document)
    db_session.flush()
    db_session.add(
        DiagnosticDocument(
            document_id=document.id,
            diagnostic_type="ACER",
            ocr_status="completed",
            validation_status="pending",
        )
    )
    db_session.add(
        DocumentWorkflowState(
            document_id=document.id,
            ingestion_status="completed",
            association_status="unassociated",
            extraction_status="extracted",
            validation_status="pending",
            destination_status="diagnostics",
        )
    )
    db_session.commit()

    response = authenticated_client.get(
        "/v2-clean/documentation/treatment?family=diagnostics&group=ACER"
    )

    assert response.status_code == 200
    assert "ACER" in response.text
    assert "Relatório técnico" in response.text


def test_treatment_bulk_validates_compatible_documents_individually(
    authenticated_client,
    db_session,
):
    documents = []
    for index in range(2):
        document = Document(
            title=f"Fatura lote {index}",
            document_type="workshop_supplier_invoice",
            classification="invoice",
            source="document_inbox",
            original_name=f"lote_{index}.pdf",
            file_name=f"lote_{index}.pdf",
            storage_provider="local",
            storage_path=f"Faturas/lote_{index}.pdf",
            status="pending_validation",
            supplier_name="Dispnal Pneus, S.A.",
        )
        db_session.add(document)
        db_session.flush()
        db_session.add(
            DocumentWorkflowState(
                document_id=document.id,
                ingestion_status="completed",
                association_status="unassociated",
                extraction_status="extracted",
                validation_status="pending",
                destination_status="invoices",
                invoice_nature="stock",
            )
        )
        documents.append(document)
    db_session.commit()

    response = authenticated_client.post(
        "/v2-clean/documentation/treatment/bulk",
        data={
            "document_ids": [str(documents[0].id), str(documents[1].id)],
            "action": "validate",
            "reason": "Validação homogénea do lote",
            "return_url": "/v2-clean/documentation/treatment?family=invoices",
        },
        follow_redirects=False,
    )
    db_session.expire_all()

    assert response.status_code == 303
    assert "bulk_processed=2" in response.headers["location"]
    assert "bulk_failed=0" in response.headers["location"]
    states = db_session.scalars(
        select(DocumentWorkflowState).where(
            DocumentWorkflowState.document_id.in_([document.id for document in documents])
        )
    ).all()
    assert {state.validation_status for state in states} == {"human_validated"}
    assert all(state.human_confirmed for state in states)
    assert (
        db_session.scalar(
            select(func.count())
            .select_from(DocumentEvent)
            .where(DocumentEvent.action == "document.treatment.validate")
        )
        == 2
    )


def test_treatment_preview_saves_services_and_creates_linked_audit_task(
    authenticated_client,
    db_session,
):
    vehicle = Vehicle(plate="AA-44-ZZ", rentway_unit_nr="4400", active=True)
    db_session.add(vehicle)
    db_session.flush()
    document = Document(
        title="Fatura para auditoria",
        document_type="workshop_supplier_invoice",
        classification="invoice",
        source="document_inbox",
        original_name="fatura_auditoria.pdf",
        file_name="fatura_auditoria.pdf",
        storage_provider="local",
        storage_path="Faturas/fatura_auditoria.pdf",
        status="pending_validation",
        supplier_name="Fornecedor Teste",
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
    )
    db_session.add(document)
    db_session.flush()
    db_session.add(
        DocumentWorkflowState(
            document_id=document.id,
            ingestion_status="completed",
            association_status="associated",
            extraction_status="extracted",
            validation_status="pending",
            destination_status="invoices",
            invoice_nature="operacional",
        )
    )
    db_session.add(
        DocumentEvent(
            document_id=document.id,
            action="invoice.ocr.extracted",
            new_value=json.dumps(
                {
                    "document_number": "FT 44",
                    "odometer_km": "8180",
                    "total_with_vat": "123.45",
                    "invoice_lines": [
                        {"reference": "OLEO", "description": "Mudança de óleo", "quantity": "1"}
                    ],
                }
            ),
        )
    )
    db_session.commit()

    page = authenticated_client.get("/v2-clean/documentation/treatment?family=invoices")
    assert page.status_code == 200
    assert "FT 44" in page.text
    assert "8180" in page.text
    assert "Mudança de óleo" in page.text
    assert "Serviços da fatura" in page.text
    assert "Validar documento" in page.text
    assert "Ordem recomendada" in page.text
    assert "Falta validar o documento." in page.text
    assert "Criar tarefa de auditoria:</strong> cria uma tarefa na fila Auditoria" in page.text
    assert page.text.index("1. Guardar natureza/pendente") < page.text.index(
        "2. Associar viatura"
    ) < page.text.index("3. Extrair / reprocessar") < page.text.index(
        "4. Guardar serviços"
    ) < page.text.index("5. Validar documento") < page.text.index(
        "6. Concluir tratamento"
    )
    assert "doc-treatment-inline-status" in page.text
    assert "submitter.hasAttribute('formaction')" in page.text
    assert "FR + TR" not in page.text
    for label in ("Chaves", "Bateria", "Iluminação", "Lavagem"):
        assert label in page.text

    saved = authenticated_client.post(
        "/v2-clean/documentation/treatment/bulk",
        data={
            "document_ids": str(document.id),
            "action": "classify",
            "service_classification_present": "1",
            "maintenance": "revision",
            "tyres": "front",
            "return_url": (
                "/v2-clean/documentation/treatment?family=invoices"
                f"&open_item=document:{document.id}"
            ),
        },
        follow_redirects=False,
    )
    assert saved.status_code == 303
    assert f"open_item=document:{document.id}" in saved.headers["location"]
    tags = db_session.scalars(
        select(VehicleDocumentRecordTag).where(
            VehicleDocumentRecordTag.document_id == document.id
        )
    ).all()
    assert {(tag.category, tag.value) for tag in tags} == {
        ("maintenance", "revision"),
        ("tyres", "front"),
    }

    created = authenticated_client.post(
        f"/v2-clean/documentation/treatment/{document.id}/audit-task",
        data={
            "reason": "Confirmar serviços faturados",
            "return_url": (
                "/v2-clean/documentation/treatment?family=invoices"
                f"&open_item=document:{document.id}"
            ),
        },
        follow_redirects=False,
    )
    db_session.expire_all()
    assert created.status_code == 303
    assert f"open_item=document:{document.id}" in created.headers["location"]
    task = db_session.scalar(
        select(Task).where(Task.entity_type == "document", Task.entity_id == str(document.id))
    )
    assert task is not None
    assert task.task_type == "audit_task"
    assert task.description == "Confirmar serviços faturados"
    assert db_session.scalar(
        select(TaskDocument).where(
            TaskDocument.task_id == task.id,
            TaskDocument.document_id == document.id,
        )
    ) is not None


@pytest.mark.parametrize(
    "path, expected_label",
    [
        ("/v2-clean/documentation/imports/invoices", "Arquivo de faturas"),
        ("/v2-clean/documentation/imports/reports", "Relatórios"),
        ("/v2-clean/documentation/imports/other", "Outros documentos"),
    ],
)
def test_legacy_document_importers_require_preflight_confirmation(
    authenticated_client,
    path,
    expected_label,
):
    response = authenticated_client.get(path)

    assert response.status_code == 200
    assert 'id="doc-import-preflight-dialog"' in response.text
    assert 'class="doc-import-preflight"' in response.text
    assert f'data-import-label="{expected_label}"' in response.text
    assert "Nenhum registo foi gravado" in response.text


def test_expected_invoice_is_blocked_from_real_document_bulk_actions(
    authenticated_client,
    db_session,
):
    expected = VehicleDocumentRecord(
        source_record_type="pending_import",
        main_group="invoices",
        status="pending",
        title="Fatura esperada HFO-77",
        external_reference="HFO-77",
        supplier_name="Fornecedor Esperado",
        metadata_json={"supplier_nif": "509999999"},
    )
    db_session.add(expected)
    db_session.commit()

    review = authenticated_client.post(
        "/v2-clean/documentation/treatment/bulk/review",
        data={"item_refs": f"expected:{expected.id}", "action": "validate"},
    )

    assert review.status_code == 200
    payload = review.json()
    assert payload["compatible"] == []
    assert payload["incompatible"][0]["reason"] == (
        "expected_invoice_requires_real_document"
    )


def test_save_services_does_not_change_nature_validation_or_completion(
    authenticated_client,
    db_session,
):
    vehicle = Vehicle(plate="SV-10-CE", active=True)
    db_session.add(vehicle)
    db_session.flush()
    document = Document(
        title="Fatura serviços isolados",
        document_type="workshop_supplier_invoice",
        classification="invoice",
        source="document_inbox",
        original_name="servicos.pdf",
        file_name="servicos.pdf",
        storage_provider="local",
        storage_path="Faturas/servicos.pdf",
        status="pending_validation",
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
    )
    db_session.add(document)
    db_session.flush()
    state = DocumentWorkflowState(
        document_id=document.id,
        ingestion_status="completed",
        association_status="associated",
        extraction_status="extracted",
        validation_status="pending",
        destination_status="invoices",
        invoice_nature="operacional",
    )
    db_session.add(state)
    db_session.commit()

    response = authenticated_client.post(
        "/v2-clean/documentation/treatment/bulk",
        data={
            "item_refs": f"document:{document.id}",
            "action": "save_services",
            "service_classification_present": "1",
            "maintenance": ["revision", "degradation"],
            "pads": "both",
            "return_url": "/v2-clean/documentation/treatment?family=invoices",
        },
        follow_redirects=False,
    )
    db_session.expire_all()

    assert response.status_code == 303
    stored = db_session.get(Document, document.id)
    stored_state = db_session.get(DocumentWorkflowState, state.id)
    assert stored.status == "pending_validation"
    assert stored_state.invoice_nature == "operacional"
    assert stored_state.validation_status == "pending"
    tags = db_session.scalars(
        select(VehicleDocumentRecordTag).where(
            VehicleDocumentRecordTag.document_id == document.id
        )
    ).all()
    assert {(tag.category, tag.value) for tag in tags} == {
        ("maintenance", "revision"),
        ("maintenance", "degradation"),
        ("pads", "front"),
        ("pads", "rear"),
    }


def test_validate_action_does_not_reclassify_or_associate_visible_form_values(
    authenticated_client,
    db_session,
):
    vehicle = Vehicle(plate="SM-21-NT", active=True)
    document = Document(
        title="Fatura financeira com ações separadas",
        document_type="finance_supplier_invoice",
        classification="finance",
        source="document_inbox",
        original_name="semantica.pdf",
        file_name="semantica.pdf",
        storage_provider="local",
        storage_path="Faturas/semantica.pdf",
        status="pending_validation",
    )
    db_session.add_all([vehicle, document])
    db_session.flush()
    state = DocumentWorkflowState(
        document_id=document.id,
        ingestion_status="completed",
        association_status="unassociated",
        extraction_status="extracted",
        validation_status="pending",
        destination_status="archive",
        invoice_nature="financeira",
    )
    db_session.add(state)
    db_session.commit()

    response = authenticated_client.post(
        "/v2-clean/documentation/treatment/bulk",
        data={
            "item_refs": f"document:{document.id}",
            "action": "validate",
            "invoice_nature": "stock",
            "plate": vehicle.plate,
            "destination": "invoices",
            "return_url": "/v2-clean/documentation/treatment?family=invoices",
        },
        follow_redirects=False,
    )
    db_session.expire_all()

    assert response.status_code == 303
    assert "bulk_processed=1" in response.headers["location"]
    assert db_session.get(Document, document.id).vehicle_id is None
    stored_state = db_session.get(DocumentWorkflowState, state.id)
    assert stored_state.invoice_nature == "financeira"
    assert stored_state.destination_status == "archive"
    assert stored_state.validation_status == "human_validated"


def test_bulk_validation_isolated_per_document_and_enforces_prerequisites(
    authenticated_client,
    db_session,
):
    stock_document = Document(
        title="Fatura stock pronta",
        document_type="finance_supplier_invoice",
        classification="finance",
        source="document_inbox",
        original_name="stock.pdf",
        file_name="stock.pdf",
        storage_provider="local",
        storage_path="Faturas/stock.pdf",
        status="pending_validation",
    )
    operational_document = Document(
        title="Fatura operacional incompleta",
        document_type="workshop_supplier_invoice",
        classification="invoice",
        source="document_inbox",
        original_name="operacional.pdf",
        file_name="operacional.pdf",
        storage_provider="local",
        storage_path="Faturas/operacional.pdf",
        status="pending_validation",
    )
    db_session.add_all([stock_document, operational_document])
    db_session.flush()
    stock_state = DocumentWorkflowState(
        document_id=stock_document.id,
        ingestion_status="completed",
        association_status="unassociated",
        extraction_status="extracted",
        validation_status="pending",
        destination_status="archive",
        invoice_nature="stock",
    )
    operational_state = DocumentWorkflowState(
        document_id=operational_document.id,
        ingestion_status="completed",
        association_status="unassociated",
        extraction_status="extracted",
        validation_status="pending",
        destination_status="invoices",
        invoice_nature="operacional",
    )
    db_session.add_all([stock_state, operational_state])
    db_session.commit()

    response = authenticated_client.post(
        "/v2-clean/documentation/treatment/bulk",
        data={
            "item_refs": [
                f"document:{stock_document.id}",
                f"document:{operational_document.id}",
            ],
            "action": "validate",
            "return_url": "/v2-clean/documentation/treatment?family=invoices",
        },
        follow_redirects=False,
    )
    db_session.expire_all()

    assert response.status_code == 303
    assert "bulk_processed=1" in response.headers["location"]
    assert "bulk_failed=1" in response.headers["location"]
    assert db_session.get(DocumentWorkflowState, stock_state.id).validation_status == (
        "human_validated"
    )
    assert db_session.get(DocumentWorkflowState, operational_state.id).validation_status == (
        "pending"
    )


def test_bulk_extraction_failure_is_persisted_without_losing_item_result(
    authenticated_client,
    db_session,
    monkeypatch,
):
    vehicle = Vehicle(plate="ER-25-RO", active=True)
    db_session.add(vehicle)
    db_session.flush()
    document = Document(
        title="Fatura com erro OCR",
        document_type="workshop_supplier_invoice",
        classification="invoice",
        source="document_inbox",
        original_name="erro-ocr.pdf",
        file_name="erro-ocr.pdf",
        storage_provider="local",
        storage_path="Faturas/erro-ocr.pdf",
        status="pending_extraction",
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
    )
    db_session.add(document)
    db_session.flush()
    state = DocumentWorkflowState(
        document_id=document.id,
        ingestion_status="completed",
        association_status="associated",
        extraction_status="not_requested",
        validation_status="pending",
        destination_status="invoices",
        invoice_nature="operacional",
    )
    db_session.add(state)
    db_session.commit()
    monkeypatch.setattr(
        web_router,
        "_reprocess_invoice_document",
        lambda *_args, **_kwargs: {"error": "ocr_failed", "extracted_text": ""},
    )

    response = authenticated_client.post(
        "/v2-clean/documentation/treatment/bulk",
        data={
            "item_refs": f"document:{document.id}",
            "action": "reprocess",
            "return_url": "/v2-clean/documentation/treatment?family=invoices",
        },
        follow_redirects=False,
    )
    db_session.expire_all()

    assert response.status_code == 303
    assert "bulk_processed=0" in response.headers["location"]
    assert "bulk_failed=1" in response.headers["location"]
    assert db_session.get(Document, document.id).status == "ocr_issue"
    assert db_session.get(DocumentWorkflowState, state.id).extraction_status == "failed"
    assert db_session.scalar(
        select(DocumentEvent).where(
            DocumentEvent.document_id == document.id,
            DocumentEvent.action == "document.treatment.reprocess",
        )
    )


def test_bulk_common_plate_requires_explicit_confirmation(
    authenticated_client,
    db_session,
):
    vehicle = Vehicle(plate="LT-20-AA", active=True)
    documents = [
        Document(
            title=f"Fatura associação {index}",
            document_type="workshop_supplier_invoice",
            classification="invoice",
            source="document_inbox",
            original_name=f"associate-{index}.pdf",
            file_name=f"associate-{index}.pdf",
            storage_provider="local",
            storage_path=f"Faturas/associate-{index}.pdf",
            status="pending_validation",
        )
        for index in range(2)
    ]
    db_session.add(vehicle)
    db_session.add_all(documents)
    db_session.flush()
    for document in documents:
        db_session.add(
            DocumentWorkflowState(
                document_id=document.id,
                ingestion_status="completed",
                association_status="unassociated",
                extraction_status="extracted",
                validation_status="pending",
                destination_status="invoices",
                invoice_nature="operacional",
            )
        )
    db_session.commit()
    data = {
        "item_refs": [f"document:{document.id}" for document in documents],
        "action": "associate",
        "plate": vehicle.plate,
        "return_url": "/v2-clean/documentation/treatment?family=invoices",
    }

    rejected = authenticated_client.post(
        "/v2-clean/documentation/treatment/bulk",
        data=data,
        follow_redirects=False,
    )
    assert "common_plate_confirmation_required" in rejected.headers["location"]
    assert all(db_session.get(Document, document.id).vehicle_id is None for document in documents)

    accepted = authenticated_client.post(
        "/v2-clean/documentation/treatment/bulk",
        data={**data, "confirm_common_plate": "1"},
        follow_redirects=False,
    )
    db_session.expire_all()
    assert "bulk_processed=2" in accepted.headers["location"]
    assert {
        db_session.get(Document, document.id).vehicle_id for document in documents
    } == {vehicle.id}


def test_vehicle_invoice_nature_uses_common_classification_audit(
    authenticated_client,
    db_session,
):
    vehicle = Vehicle(plate="NT-30-RE", active=True)
    db_session.add(vehicle)
    db_session.flush()
    document = Document(
        title="Fatura natureza ficha",
        document_type="workshop_supplier_invoice",
        classification="invoice",
        source="document_inbox",
        original_name="natureza.pdf",
        file_name="natureza.pdf",
        storage_provider="local",
        storage_path="Faturas/natureza.pdf",
        status="pending_validation",
        vehicle_id=vehicle.id,
        plate=vehicle.plate,
    )
    db_session.add(document)
    db_session.commit()

    response = authenticated_client.post(
        f"/v2-clean/fleet/{vehicle.id}/documents/{document.id}/nature",
        data={"nature": "operacional", "decision_reason": "Confirmado na ficha"},
        follow_redirects=False,
    )
    db_session.expire_all()

    assert response.status_code == 303
    state = db_session.scalar(
        select(DocumentWorkflowState).where(
            DocumentWorkflowState.document_id == document.id
        )
    )
    assert state.invoice_nature == "operacional"
    assert db_session.scalar(
        select(DocumentEvent).where(
            DocumentEvent.document_id == document.id,
            DocumentEvent.action == "invoice.nature.classified",
        )
    )


def test_by_vehicle_overview_exposes_explainable_aggregates_and_central_preview(
    authenticated_client,
    db_session,
):
    vehicle = Vehicle(plate="PV-40-UA", brand="PEUGEOT", model="208", active=True)
    db_session.add(vehicle)
    db_session.flush()
    work_order = VehicleDocumentRecord(
        vehicle_id=vehicle.id,
        source_record_type="structured",
        main_group="work_orders",
        title="FO-400",
        external_reference="FO-400",
        document_date=date(2026, 7, 1),
        status="structured",
    )
    expected = VehicleDocumentRecord(
        vehicle_id=vehicle.id,
        source_record_type="pending_import",
        main_group="invoices",
        title="Fatura esperada 401",
        external_reference="401",
        document_date=date(2026, 7, 3),
        status="pending",
    )
    invoices = [
        Document(
            title=f"Fatura repetição {index}",
            document_type="workshop_supplier_invoice",
            classification="invoice",
            source="document_inbox",
            original_name=f"repeat-{index}.pdf",
            file_name=f"repeat-{index}.pdf",
            storage_provider="local",
            storage_path=f"Faturas/repeat-{index}.pdf",
            status="pending_validation",
            vehicle_id=vehicle.id,
            plate=vehicle.plate,
            document_date=date(2026, 7, 5 + index * 20),
        )
        for index in range(2)
    ]
    db_session.add_all([work_order, expected, *invoices])
    db_session.flush()
    db_session.add(
        DocumentLink(
            document_id=invoices[0].id,
            entity_type="vehicle_document_record",
            entity_id=str(work_order.id),
            category="invoice_work_order",
        )
    )
    for document in invoices:
        db_session.add(
            VehicleDocumentRecordTag(
                vehicle_id=vehicle.id,
                document_id=document.id,
                category="maintenance",
                value="revision",
                source_kind="manual",
            )
        )
    db_session.commit()

    page = authenticated_client.get("/v2-clean/documentation/by-vehicle")
    preview = authenticated_client.get(
        f"/v2-clean/documentation/by-vehicle/{vehicle.id}"
    )

    assert page.status_code == 200
    assert vehicle.plate in page.text
    assert "Pares confirmados" in page.text
    assert "FO sem fatura" in page.text
    assert "≤ 30 dias" in page.text
    assert "sem algoritmo opaco" in page.text.lower()
    assert preview.status_code == 200
    assert "Preview central por viatura" in preview.text
    assert "Faturas reais" in preview.text
    assert "Fatura repetição 0" in preview.text
    assert "Faturas esperadas" in preview.text
    assert "Correspondência FO–fatura" in preview.text
    assert "Histórico cronológico" in preview.text
