from openpyxl import Workbook
from sqlalchemy import select

import json

from app.models.documents import Document, DocumentEvent, VehicleDocumentRecord
from app.models.vehicles import Vehicle
from app.services.pending_document_importer import (
    create_pending_documents_from_preview,
    import_pending_documents,
    preview_pending_documents,
    reconcile_pending_invoices,
)


def _workbook(path, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Faturas"
    sheet.append(
        [
            "Nº fatura",
            "Data",
            "Fornecedor",
            "NIF fornecedor",
            "Chassi",
            "Matrícula",
            "Unit",
            "Total",
        ]
    )
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def test_pending_invoice_import_associates_by_vin_plate_and_unit(db_session, tmp_path):
    vehicles = [
        Vehicle(plate="AA-11-AA", vin="VF3YBBPFC12W31462", rentway_unit_nr="101"),
        Vehicle(plate="BB-22-BB", vin="VR3EDYHT9RJ968630", rentway_unit_nr="102"),
        Vehicle(plate="CC-33-CC", vin="VR3EDYHT2RJ968629", rentway_unit_nr="103"),
    ]
    db_session.add_all(vehicles)
    db_session.flush()
    path = tmp_path / "pendentes.xlsx"
    _workbook(
        path,
        [
            [
                "HFO/1/2025",
                "01/01/2025",
                "Fornecedor A",
                "500000001",
                "774234/VF3YBBPFC12W31462",
                "",
                "",
                "100,20",
            ],
            ["HFO/2/2025", "02/01/2025", "Fornecedor A", "500000001", "", "BB22BB", "", "200,30"],
            ["HFO/3/2025", "03/01/2025", "Fornecedor A", "500000001", "", "", 103, "300,40"],
            ["HFO/4/2025", "04/01/2025", "Fornecedor A", "500000001", "", "", "", "400,50"],
        ],
    )

    result = import_pending_documents(
        db_session,
        path=path,
        original_name=path.name,
        user_id=None,
    )

    assert result == {"created": 4, "associated": 3, "unmatched": 1, "duplicates": 0, "invalid": 0}
    records = db_session.scalars(
        select(VehicleDocumentRecord).order_by(VehicleDocumentRecord.external_reference)
    ).all()
    assert [record.vehicle_id for record in records] == [
        vehicles[0].id,
        vehicles[1].id,
        vehicles[2].id,
        None,
    ]
    assert records[0].metadata_json["association_method"] == "vin"
    assert records[3].metadata_json["expected_total"] == "400.50"


def test_pending_invoice_import_is_idempotent(db_session, tmp_path):
    path = tmp_path / "pendentes.xlsx"
    _workbook(
        path,
        [["FAC 10", "01/01/2025", "Fornecedor", "500000001", "", "", "", "10,00"]],
    )
    first = import_pending_documents(db_session, path=path, original_name=path.name, user_id=None)
    second = import_pending_documents(db_session, path=path, original_name=path.name, user_id=None)

    assert first["created"] == 1
    assert second["created"] == 0
    assert second["duplicates"] == 1


def test_pending_invoice_preview_does_not_write_and_confirmation_is_selective(
    db_session,
    tmp_path,
):
    vehicle = Vehicle(
        plate="AA-11-AA",
        vin="VF3YBBPFC12W31462",
        rentway_unit_nr="101",
    )
    db_session.add(vehicle)
    db_session.flush()
    path = tmp_path / "rever.xlsx"
    _workbook(
        path,
        [
            ["FAC 20", "01/01/2025", "Fornecedor", "500000001", vehicle.vin, "", "", "10,00"],
            ["FAC 21", "02/01/2025", "Fornecedor", "500000001", "", "", "", "20,00"],
        ],
    )

    preview = preview_pending_documents(
        db_session,
        path=path,
        original_name=path.name,
    )

    assert preview["summary"] == {
        "total": 2,
        "ready": 2,
        "associated": 1,
        "unmatched": 1,
        "duplicates": 0,
        "invalid": 0,
    }
    assert db_session.scalar(select(VehicleDocumentRecord)) is None

    result = create_pending_documents_from_preview(
        db_session,
        preview=preview,
        selected_row_ids={"2"},
        user_id=None,
    )

    assert result["created"] == 1
    record = db_session.scalar(select(VehicleDocumentRecord))
    assert record.external_reference == "FAC 21"
    assert record.vehicle_id is None


def test_pending_invoice_duplicate_uses_supplier_nif_and_number(db_session, tmp_path):
    first_path = tmp_path / "primeira.xlsx"
    second_path = tmp_path / "segunda.xlsx"
    _workbook(
        first_path,
        [["FAC 30", "01/01/2025", "Fornecedor A", "500000001", "", "", "", "10,00"]],
    )
    _workbook(
        second_path,
        [
            ["FAC 30", "02/01/2025", "Fornecedor B", "500000002", "", "", "", "20,00"],
            ["FAC 30", "03/01/2025", "Fornecedor A", "500000001", "", "", "", "30,00"],
        ],
    )
    import_pending_documents(
        db_session,
        path=first_path,
        original_name=first_path.name,
        user_id=None,
    )

    preview = preview_pending_documents(
        db_session,
        path=second_path,
        original_name=second_path.name,
    )

    assert [row["status"] for row in preview["rows"]] == ["ready", "duplicate"]
    assert preview["rows"][1]["status_detail"] == "Já existe para este NIF e número"


def test_reconcile_pending_invoice_uses_ocr_identity_and_links_vehicle(db_session):
    vehicle = Vehicle(
        plate="AA-11-AA",
        vin="VF3YBBPFC12W31462",
        rentway_unit_nr="101",
    )
    pending = VehicleDocumentRecord(
        source_record_type="pending_import",
        main_group="invoices",
        status="pending",
        external_reference="HFO/3081/2025",
        supplier_name="Caetano Gamobar",
        metadata_json={
            "supplier_nif": "500112967",
            "unit_number": "101",
        },
    )
    document = Document(
        title="Fatura oficina",
        document_type="workshop_supplier_invoice",
        classification="invoice",
        source="v2_clean_manual",
        entry_channel="v2_clean_batch",
        original_name="hfo_3081.pdf",
        file_name="hfo_3081.pdf",
        storage_provider="local",
        storage_path="Frota/AA11AA/hfo_3081.pdf",
        status="extracted",
        archived=True,
    )
    db_session.add_all([vehicle, pending, document])
    db_session.flush()
    db_session.add(
        DocumentEvent(
            document_id=document.id,
            action="invoice.ocr.extracted",
            old_value=None,
            new_value=json.dumps(
                {
                    "document_number": "HFO/3081/2025",
                    "supplier_nif": "500112967",
                }
            ),
            user_id=None,
        )
    )
    db_session.flush()

    result = reconcile_pending_invoices(db_session, user_id=None)

    assert result["associated"] == 1
    assert result["fulfilled"] == 1
    assert pending.vehicle_id == vehicle.id
    assert pending.document_id == document.id
    assert pending.status == "fulfilled"
    assert document.vehicle_id == vehicle.id


def test_reconcile_pending_invoice_keeps_ambiguous_number_for_manual_review(db_session):
    pending = VehicleDocumentRecord(
        source_record_type="pending_import",
        main_group="invoices",
        status="pending",
        external_reference="FAC 50",
        metadata_json={},
    )
    documents = [
        Document(
            title=f"Fatura {index}",
            document_type="workshop_supplier_invoice",
            classification="invoice",
            source="v2_clean_manual",
            entry_channel="v2_clean_batch",
            original_name=f"fac_50_{index}.pdf",
            file_name=f"fac_50_{index}.pdf",
            storage_provider="local",
            storage_path=f"Frota/fac_50_{index}.pdf",
            contract_number="FAC 50",
            status="extracted",
            archived=True,
        )
        for index in (1, 2)
    ]
    db_session.add_all([pending, *documents])
    db_session.flush()

    result = reconcile_pending_invoices(db_session, user_id=None)

    assert result["ambiguous"] == 1
    assert result["fulfilled"] == 0
    assert pending.status == "pending"
    assert pending.document_id is None
    assert pending.metadata_json["reconciliation_state"] == "ambiguous"
